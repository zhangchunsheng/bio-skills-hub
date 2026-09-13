#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
generate_docx.py — 罗定市人民医院药物临床试验机构体系文件格式修正文档生成器

依据《SOP-SOP》V1.0 强制格式标准，将结构化 JSON 渲染为符合规范的 .docx 文件。
仅生成格式，不修改业务内容；业务内容由调用方（智能体）从原文解析后原样传入。

用法：
    python generate_docx.py <input_json_path> [<output_docx_path>]

输入 JSON 结构示例：
{
  "doc_type": "SOP",                  // JOB | MGT | SOP | DSN | EMR
  "sop_sop_version": "V1.0",
  "output_path": "/abs/path/out.docx", // 可选；若未提供则用第二个命令行参数或同目录同名 .docx
  "file_code": "SOP-GCP-001",          // 页眉右侧文件编码
  "title": "标准操作规程",              // 第 1 页主标题（不含"/修订记录"）
  "title_suffix": "/修订记录",          // 可选，追加到主标题
  "header_tables": [                   // 第 1 页表头表格（按出现顺序）
    {
      "rows": [
        ["制/修订人", "张三", "审核人", "李四"],
        ["批准人", "王五", "发布日期", "2024-01-01", "生效日期", "2024-02-01"]
      ]
    }
  ],
  "body_blocks": [
    {"type": "h1", "text": "一、目的"},
    {"type": "p",  "text": "为规范本机构……"},
    {"type": "h2", "text": "（一）适用范围"},
    {"type": "h3", "text": "1. 临床试验项目"},
    {"type": "h4", "text": "（1）立项阶段"},
    {"type": "h5", "text": "① 递交申请"},
    {"type": "h6", "text": "1) 形式审查"},
    {"type": "body_table", "header": ["序号", "文件名称"], "rows": [["1", "知情同意书"]]},
    {"type": "attachment_title", "text": "附件1：知情同意书模板"},
    {"type": "attachment_body", "text": "……"}
  ]
}

退出码：
    0 成功
    1 输入文件不存在 / JSON 解析失败
    2 缺少 python-docx 依赖
    3 渲染过程异常
"""

import json
import os
import sys
import traceback

# ---------- 依赖检查 ----------
try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_ROW_HEIGHT, WD_TABLE_ALIGNMENT
    from docx.enum.section import WD_SECTION
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.stderr.write(
        "[generate_docx] 缺少依赖 python-docx。\n"
        "请运行：C:\\Users\\qqq\\.workbuddy\\binaries\\python\\envs\\default\\Scripts\\pip install python-docx\n"
    )
    sys.exit(2)


# ---------- 常量 ----------
FONT_SONG = "宋体"
FONT_HTSONG = "华文中宋"
FONT_HEI = "黑体"

PT_SANHAO = Pt(16)        # 三号
PT_XIAOSI = Pt(12)        # 小四
PT_11 = Pt(11)            # 11 号
PT_LINE_24 = Pt(24)       # 固定行距 24 磅
PT_SPACE_BEFORE_15 = Pt(24)  # 段前 1.5 行（≈ 1.5 × 16pt）
PT_SPACE_AFTER_05 = Pt(8)    # 段后 0.5 行（≈ 0.5 × 16pt）
INDENT_2CHAR = Pt(24)     # 首行缩进 2 字符（≈ 2 × 12pt）
ROW_HEIGHT_12CM = Cm(1.2) # 表头表格行高 1.2 厘米

MARGIN_TOP = Cm(2.5)
MARGIN_BOTTOM = Cm(2.5)
MARGIN_LEFT = Cm(2.6)
MARGIN_RIGHT = Cm(2.6)


# ---------- 通用工具 ----------
def set_run_font(run, font_name, size, bold=False):
    """设置 Run 的中西文字体、字号、加粗。"""
    run.font.name = font_name
    # 东亚字体必须通过 rPr 显式指定，否则中文会回退到默认字体
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    run.font.size = size
    run.font.bold = bold


def set_para_format(para, *, alignment=None, line_24=False,
                    space_before=None, space_after=None,
                    first_line_indent=None, line_rule_exact=True):
    """统一设置段落格式。"""
    pf = para.paragraph_format
    if alignment is not None:
        para.alignment = alignment
    if line_24:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = PT_LINE_24
    if space_before is not None:
        pf.space_before = space_before
    if space_after is not None:
        pf.space_after = space_after
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent


def add_field(paragraph, instr):
    """在段落中插入 Word 域（如 PAGE / NUMPAGES）。"""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    run._r.append(instr_text)

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)

    # 占位文本（打开 Word 后会自动刷新）
    t = OxmlElement("w:t")
    t.text = "1"
    run._r.append(t)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)

    return run


def set_cell_text(cell, text, font_name, size, bold=False,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, line_24=False,
                  space_before=None, space_after=None,
                  vertical_center=True):
    """重置单元格文本并设置格式。vertical_center 控制单元格垂直居中。"""
    # 清空原内容
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = alignment
    pf = para.paragraph_format
    if line_24:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = PT_LINE_24
    if space_before is not None:
        pf.space_before = space_before
    if space_after is not None:
        pf.space_after = space_after
    pf.first_line_indent = None
    run = para.add_run(text or "")
    set_run_font(run, font_name, size, bold=bold)
    # 设置单元格垂直居中
    if vertical_center:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        vAlign = tcPr.find(qn("w:vAlign"))
        if vAlign is None:
            vAlign = OxmlElement("w:vAlign")
            tcPr.append(vAlign)
        vAlign.set(qn("w:val"), "center")


def set_row_height(row, height, exactly=True):
    """设置表格行高。"""
    row.height = height
    rule = WD_ROW_HEIGHT.EXACTLY if exactly else WD_ROW_HEIGHT.AT_LEAST
    row.height_rule = rule


def set_table_borders(table):
    """为表格添加单线边框（细黑线）。"""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tbl_pr.append(borders)


# ---------- 页面与页眉页脚 ----------
def setup_page(section):
    """设置 A4 纸张与页边距。"""
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = MARGIN_TOP
    section.bottom_margin = MARGIN_BOTTOM
    section.left_margin = MARGIN_LEFT
    section.right_margin = MARGIN_RIGHT
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)


def setup_header(section, file_code):
    """页眉：左=机构名，右=文件编码（不带前缀），宋体小四。
    使用 1×2 无边框表格实现左右分布，确保在所有 Word/WPS 版本中可靠右对齐。"""
    from docx.enum.table import WD_TABLE_ALIGNMENT

    header = section.header
    header.is_linked_to_previous = False

    # 清空默认段落
    default_para = header.paragraphs[0]
    for child in list(default_para._element):
        default_para._element.remove(child)

    # 创建 1×2 无边框表格
    tbl = OxmlElement("w:tbl")

    # 表格属性：100% 宽度，无边框，固定布局
    tblPr = OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")  # 5000 = 100% (在五十分之一单位中)
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)

    # 无边框
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:{}".format(edge))
        b.set(qn("w:val"), "none")
        b.set(qn("w:sz"), "0")
        b.set(qn("w:space"), "0")
        tblBorders.append(b)
    tblPr.append(tblBorders)

    # 表格布局：固定
    tblLayout = OxmlElement("w:tblLayout")
    tblLayout.set(qn("w:type"), "fixed")
    tblPr.append(tblLayout)
    tbl.append(tblPr)

    # 表格网格：两列各占 50%
    tblGrid = OxmlElement("w:tblGrid")
    gridCol1 = OxmlElement("w:gridCol")
    gridCol1.set(qn("w:w"), "4479")  # 半宽（twips）
    tblGrid.append(gridCol1)
    gridCol2 = OxmlElement("w:gridCol")
    gridCol2.set(qn("w:w"), "4480")
    tblGrid.append(gridCol2)
    tbl.append(tblGrid)

    # 行
    tr = OxmlElement("w:tr")

    # 左单元格：机构名，左对齐
    tc1 = OxmlElement("w:tc")
    tcPr1 = OxmlElement("w:tcPr")
    tcW1 = OxmlElement("w:tcW")
    tcW1.set(qn("w:w"), "4479")
    tcW1.set(qn("w:type"), "dxa")
    tcPr1.append(tcW1)
    # 垂直居中
    vAlign1 = OxmlElement("w:vAlign")
    vAlign1.set(qn("w:val"), "center")
    tcPr1.append(vAlign1)
    tc1.append(tcPr1)
    # 段落
    p1 = OxmlElement("w:p")
    pPr1 = OxmlElement("w:pPr")
    pStyle1 = OxmlElement("w:pStyle")
    pStyle1.set(qn("w:val"), "Header")
    pPr1.append(pStyle1)
    # 左对齐
    jc1 = OxmlElement("w:jc")
    jc1.set(qn("w:val"), "left")
    pPr1.append(jc1)
    p1.append(pPr1)
    # Run: 机构名
    r1 = OxmlElement("w:r")
    rPr1 = OxmlElement("w:rPr")
    rFonts1 = OxmlElement("w:rFonts")
    rFonts1.set(qn("w:ascii"), FONT_SONG)
    rFonts1.set(qn("w:hAnsi"), FONT_SONG)
    rFonts1.set(qn("w:eastAsia"), FONT_SONG)
    rPr1.append(rFonts1)
    sz1 = OxmlElement("w:sz")
    sz1.set(qn("w:val"), "24")
    rPr1.append(sz1)
    r1.append(rPr1)
    t1 = OxmlElement("w:t")
    t1.text = "罗定市人民医院药物临床试验机构"
    r1.append(t1)
    p1.append(r1)
    tc1.append(p1)
    tr.append(tc1)

    # 右单元格：文件编码，右对齐
    tc2 = OxmlElement("w:tc")
    tcPr2 = OxmlElement("w:tcPr")
    tcW2 = OxmlElement("w:tcW")
    tcW2.set(qn("w:w"), "4480")
    tcW2.set(qn("w:type"), "dxa")
    tcPr2.append(tcW2)
    vAlign2 = OxmlElement("w:vAlign")
    vAlign2.set(qn("w:val"), "center")
    tcPr2.append(vAlign2)
    tc2.append(tcPr2)
    # 段落
    p2 = OxmlElement("w:p")
    pPr2 = OxmlElement("w:pPr")
    pStyle2 = OxmlElement("w:pStyle")
    pStyle2.set(qn("w:val"), "Header")
    pPr2.append(pStyle2)
    # 右对齐
    jc2 = OxmlElement("w:jc")
    jc2.set(qn("w:val"), "right")
    pPr2.append(jc2)
    p2.append(pPr2)
    # Run: 文件编码
    r2 = OxmlElement("w:r")
    rPr2 = OxmlElement("w:rPr")
    rFonts2 = OxmlElement("w:rFonts")
    rFonts2.set(qn("w:ascii"), FONT_SONG)
    rFonts2.set(qn("w:hAnsi"), FONT_SONG)
    rFonts2.set(qn("w:eastAsia"), FONT_SONG)
    rPr2.append(rFonts2)
    sz2 = OxmlElement("w:sz")
    sz2.set(qn("w:val"), "24")
    rPr2.append(sz2)
    r2.append(rPr2)
    t2 = OxmlElement("w:t")
    t2.text = file_code or ""
    r2.append(t2)
    p2.append(r2)
    tc2.append(p2)
    tr.append(tc2)

    tbl.append(tr)

    # 将表格插入到页眉中（在默认段落之前）
    header._element.insert(0, tbl)


def setup_footer(section):
    """页脚：第 X 页 共 Y 页，居中，宋体小四。"""
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r1 = para.add_run("第 ")
    set_run_font(r1, FONT_SONG, PT_XIAOSI, bold=False)
    add_field(para, "PAGE")
    r2 = para.add_run(" 页 共 ")
    set_run_font(r2, FONT_SONG, PT_XIAOSI, bold=False)
    add_field(para, "NUMPAGES")
    r3 = para.add_run(" 页")
    set_run_font(r3, FONT_SONG, PT_XIAOSI, bold=False)
    # 设置所有 run 字体（包括域内的）
    for r in para.runs:
        set_run_font(r, FONT_SONG, PT_XIAOSI, bold=False)


# ---------- 渲染各类块 ----------
def render_title(doc, text):
    """主标题：华文中宋 三号 加粗 居中 段前1.5行 段后0.5行 固定24磅。"""
    para = doc.add_paragraph()
    set_para_format(
        para,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_24=True,
        space_before=PT_SPACE_BEFORE_15,
        space_after=PT_SPACE_AFTER_05,
        first_line_indent=None,
    )
    run = para.add_run(text or "")
    set_run_font(run, FONT_HTSONG, PT_SANHAO, bold=True)


def render_header_table(doc, rows):
    """第 1 页表头表格：宋体 11 号 加粗 居中 行高 1.2cm。"""
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        set_row_height(row, ROW_HEIGHT_12CM, exactly=True)
        for j in range(n_cols):
            cell = row.cells[j]
            text = row_data[j] if j < len(row_data) else ""
            set_cell_text(
                cell, text, FONT_SONG, PT_11, bold=True,
                alignment=WD_ALIGN_PARAGRAPH.CENTER, line_24=False,
                space_before=Pt(0), space_after=Pt(0),
            )
    # 表格后空一行
    doc.add_paragraph()


def render_body_table(doc, header, rows):
    """正文表格：表头 宋体小四加粗；内容 宋体小四。"""
    n_cols = len(header) if header else (max(len(r) for r in rows) if rows else 0)
    if n_cols == 0:
        return
    total_rows = (1 if header else 0) + len(rows)
    table = doc.add_table(rows=total_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    r_idx = 0
    if header:
        row = table.rows[r_idx]
        for j, text in enumerate(header):
            cell = row.cells[j]
            set_cell_text(
                cell, text, FONT_SONG, PT_XIAOSI, bold=True,
                alignment=WD_ALIGN_PARAGRAPH.CENTER, line_24=False,
                space_before=Pt(0), space_after=Pt(0),
            )
        r_idx += 1
    for body_row in rows:
        row = table.rows[r_idx]
        for j in range(n_cols):
            cell = row.cells[j]
            text = body_row[j] if j < len(body_row) else ""
            set_cell_text(
                cell, text, FONT_SONG, PT_XIAOSI, bold=False,
                alignment=WD_ALIGN_PARAGRAPH.LEFT, line_24=False,
                space_before=Pt(0), space_after=Pt(0),
            )
        r_idx += 1
    doc.add_paragraph()


def render_heading(doc, text, level, bold):
    """一~六级标题：宋体小四，首行缩进2字符，两端对齐，固定24磅。"""
    para = doc.add_paragraph()
    set_para_format(
        para,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_24=True,
        space_before=Pt(0),
        space_after=Pt(0),
        first_line_indent=INDENT_2CHAR,
    )
    run = para.add_run(text or "")
    set_run_font(run, FONT_SONG, PT_XIAOSI, bold=bold)


def render_paragraph(doc, text):
    """正文段落：宋体小四，固定24磅，段前段后0，首行缩进2字符，两端对齐。"""
    para = doc.add_paragraph()
    set_para_format(
        para,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_24=True,
        space_before=Pt(0),
        space_after=Pt(0),
        first_line_indent=INDENT_2CHAR,
    )
    run = para.add_run(text or "")
    set_run_font(run, FONT_SONG, PT_XIAOSI, bold=False)


def render_attachment_title(doc, text):
    """附件标题：黑体 三号 居中。"""
    para = doc.add_paragraph()
    set_para_format(
        para,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_24=True,
        space_before=PT_SPACE_BEFORE_15,
        space_after=PT_SPACE_AFTER_05,
        first_line_indent=None,
    )
    run = para.add_run(text or "")
    set_run_font(run, FONT_HEI, PT_SANHAO, bold=False)


def render_attachment_body(doc, text):
    """附件正文：宋体小四。"""
    para = doc.add_paragraph()
    set_para_format(
        para,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_24=True,
        space_before=Pt(0),
        space_after=Pt(0),
        first_line_indent=INDENT_2CHAR,
    )
    run = para.add_run(text or "")
    set_run_font(run, FONT_SONG, PT_XIAOSI, bold=False)


def render_page_break(doc):
    """插入分页符。"""
    para = doc.add_paragraph()
    run = para.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._element.append(br)


def render_block(doc, block):
    """根据 block.type 分发到对应渲染函数。"""
    btype = block.get("type", "").lower()
    text = block.get("text", "")
    if btype == "title":
        render_title(doc, text)
    elif btype == "header_table":
        render_header_table(doc, block.get("rows", []))
    elif btype == "body_table":
        render_body_table(doc, block.get("header", []), block.get("rows", []))
    elif btype == "h1":
        # 一级标题：加粗
        render_heading(doc, text, 1, bold=True)
    elif btype in ("h2", "h3", "h4", "h5", "h6"):
        # 二~六级标题：不加粗
        render_heading(doc, text, int(btype[1]), bold=False)
    elif btype == "p":
        render_paragraph(doc, text)
    elif btype == "attachment_title":
        render_attachment_title(doc, text)
    elif btype == "attachment_body":
        render_attachment_body(doc, text)
    elif btype == "blank":
        doc.add_paragraph()
    elif btype == "page_break":
        render_page_break(doc)
    else:
        # 未知类型当作普通段落处理
        render_paragraph(doc, text)


# ---------- 主入口 ----------
def render(data, output_path):
    doc = Document()

    # 默认正文样式设为宋体小四（防止 Normal 样式回退到 Calibri）
    normal = doc.styles["Normal"]
    normal.font.name = FONT_SONG
    normal.font.size = PT_XIAOSI
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_SONG)
    rfonts.set(qn("w:ascii"), FONT_SONG)
    rfonts.set(qn("w:hAnsi"), FONT_SONG)

    section = doc.sections[0]
    setup_page(section)
    setup_header(section, data.get("file_code", ""))
    setup_footer(section)

    # 主标题：仅当 body_blocks 中没有 title 类型块时，才用顶层 title/title_suffix 渲染。
    # 如果 body_blocks 中包含 title 块，则完全由 body_blocks 控制标题结构
    # （支持原文多段标题：如"岗位职责"一段、"修订记录"一段、重复标题一段）。
    body_blocks = data.get("body_blocks", [])
    has_title_in_blocks = any(b.get("type", "").lower() == "title" for b in body_blocks)
    if not has_title_in_blocks:
        title = data.get("title", "")
        suffix = data.get("title_suffix", "")
        if title:
            render_title(doc, title + (suffix or ""))

    # 第 1 页表头表格：仅当 body_blocks 中没有 header_table 类型块时，才用顶层 header_tables 渲染。
    has_header_table_in_blocks = any(b.get("type", "").lower() == "header_table" for b in body_blocks)
    if not has_header_table_in_blocks:
        for ht in data.get("header_tables", []):
            render_header_table(doc, ht.get("rows", []))

    # 正文块（可包含 title / header_table / h1~h6 / p / body_table / attachment_* / blank）
    for block in body_blocks:
        render_block(doc, block)

    # 确保输出目录存在
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    doc.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("用法：python generate_docx.py <input.json> [output.docx]\n")
        sys.exit(1)

    input_path = sys.argv[1]
    if not os.path.exists(input_path):
        sys.stderr.write(f"[generate_docx] 输入文件不存在：{input_path}\n")
        sys.exit(1)

    with open(input_path, "r", encoding="utf-8") as f:
        try:
            data = json.load(f)
        except json.JSONDecodeError as e:
            sys.stderr.write(f"[generate_docx] JSON 解析失败：{e}\n")
            sys.exit(1)

    output_path = (
        data.get("output_path")
        or (sys.argv[2] if len(sys.argv) > 2 else None)
        or os.path.splitext(input_path)[0] + ".docx"
    )

    try:
        result = render(data, output_path)
        # 输出绝对路径，便于智能体定位
        print(os.path.abspath(result))
    except Exception as e:
        sys.stderr.write("[generate_docx] 渲染异常：\n")
        traceback.print_exc()
        sys.exit(3)


if __name__ == "__main__":
    main()
