#!/usr/bin/env python3
"""Convert Excel/CSV/TSV data into readable Markdown documents."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence


EXCEL_OPENPYXL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
DELIMITED_SUFFIXES = {".csv", ".tsv"}
SUPPORTED_SUFFIXES = EXCEL_OPENPYXL_SUFFIXES | DELIMITED_SUFFIXES | {".xls"}
INVALID_FILENAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


@dataclass
class CellData:
    value: Any = None
    formula: str | None = None
    number_format: str = "General"
    hyperlink: str | None = None


@dataclass
class SheetData:
    title: str
    rows: list[list[CellData]]
    hidden: bool = False
    merged_ranges: int = 0
    charts: int = 0
    images: int = 0
    source_rows: int = 0
    source_columns: int = 0
    warnings: list[str] = field(default_factory=list)


@dataclass
class ExportResult:
    sheet: str
    layout: str
    records: int
    files: list[str]
    header_row: int | None
    header_rows: int
    warnings: list[str]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert Excel workbooks or CSV/TSV files into Markdown documents."
    )
    parser.add_argument("source", help="Source .xlsx/.xlsm/.xls/.csv/.tsv file")
    parser.add_argument(
        "--output-dir",
        help="Directory for generated Markdown files (default: <source>_markdown beside source)",
    )
    parser.add_argument(
        "--layout",
        choices=("auto", "table", "records"),
        default="auto",
        help="Markdown layout (default: auto)",
    )
    parser.add_argument(
        "--sheet",
        action="append",
        default=[],
        metavar="NAME",
        help="Export a named worksheet; repeat for multiple sheets",
    )
    parser.add_argument(
        "--header-row",
        default="auto",
        metavar="auto|none|N",
        help="First header row, one-based (default: auto)",
    )
    parser.add_argument(
        "--header-rows",
        type=int,
        default=1,
        metavar="N",
        help="Number of header rows to combine (default: 1)",
    )
    parser.add_argument(
        "--merged-cells",
        choices=("repeat", "anchor"),
        default="repeat",
        help="Repeat merged anchor values or retain anchors only (default: repeat)",
    )
    parser.add_argument(
        "--formulas",
        choices=("values", "formulas", "both"),
        default="values",
        help="Use cached values, formula text, or both (default: values)",
    )
    parser.add_argument(
        "--max-rows-per-file",
        type=int,
        default=0,
        metavar="N",
        help="Split each sheet after N data rows; 0 disables splitting",
    )
    parser.add_argument(
        "--include-hidden",
        action="store_true",
        help="Include hidden worksheets, rows, and columns",
    )
    parser.add_argument(
        "--keep-existing",
        action="store_true",
        help="Keep existing .md files in the output directory",
    )
    parser.add_argument(
        "--encoding",
        default="auto",
        help="CSV/TSV encoding (default: auto-detect UTF-8/UTF-8-BOM/GB18030)",
    )
    parser.add_argument(
        "--delimiter",
        default="auto",
        help=r"CSV delimiter, such as ',' or '\t' (default: auto)",
    )
    args = parser.parse_args()

    if args.header_rows < 1:
        parser.error("--header-rows must be at least 1")
    if args.max_rows_per_file < 0:
        parser.error("--max-rows-per-file cannot be negative")
    if args.max_rows_per_file and args.max_rows_per_file < 1:
        parser.error("--max-rows-per-file must be 0 or a positive integer")
    parse_header_row(args.header_row, parser)
    return args


def parse_header_row(value: str, parser: argparse.ArgumentParser | None = None) -> str | int | None:
    normalized = str(value).strip().lower()
    if normalized == "auto":
        return "auto"
    if normalized in {"none", "0"}:
        return None
    try:
        row = int(normalized)
    except ValueError:
        if parser:
            parser.error("--header-row must be auto, none, or a positive integer")
        raise
    if row < 1:
        if parser:
            parser.error("--header-row must be at least 1")
        raise ValueError("header row must be at least 1")
    return row


def safe_filename(name: str) -> str:
    cleaned = INVALID_FILENAME.sub("_", name.strip()).strip(" ._")
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned[:100] or "sheet"


def trim_matrix(rows: list[list[CellData]]) -> list[list[CellData]]:
    while rows and all(is_empty(cell.value) and not cell.formula for cell in rows[-1]):
        rows.pop()
    if not rows:
        return []
    last_col = 0
    for row in rows:
        for index, cell in enumerate(row, start=1):
            if not is_empty(cell.value) or cell.formula:
                last_col = max(last_col, index)
    return [row[:last_col] for row in rows] if last_col else []


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def load_source(
    source: Path,
    *,
    include_hidden: bool,
    merged_cells: str,
    formulas: str,
    encoding: str,
    delimiter: str,
) -> list[SheetData]:
    suffix = source.suffix.lower()
    if suffix in EXCEL_OPENPYXL_SUFFIXES:
        return load_openpyxl(source, include_hidden, merged_cells, formulas)
    if suffix in DELIMITED_SUFFIXES:
        return [load_delimited(source, encoding, delimiter)]
    if suffix == ".xls":
        return load_legacy_xls(source)
    raise ValueError(
        f"Unsupported source type {suffix!r}. Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
    )


def load_openpyxl(
    source: Path, include_hidden: bool, merged_cells: str, formulas: str
) -> list[SheetData]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for .xlsx/.xlsm conversion") from exc

    keep_vba = source.suffix.lower() in {".xlsm", ".xltm"}
    values_wb = openpyxl.load_workbook(
        source, data_only=True, read_only=False, keep_vba=keep_vba
    )
    # Open formula text as well as cached values. This lets "values" mode fall
    # back to the formula when Excel did not save a cached calculation result.
    formulas_wb = openpyxl.load_workbook(
        source, data_only=False, read_only=False, keep_vba=keep_vba
    )

    sheets: list[SheetData] = []
    try:
        for values_ws in values_wb.worksheets:
            hidden = values_ws.sheet_state != "visible"
            if hidden and not include_hidden:
                continue
            formula_ws = formulas_wb[values_ws.title]
            hidden_cols = {
                index
                for index in range(1, values_ws.max_column + 1)
                if values_ws.column_dimensions[openpyxl.utils.get_column_letter(index)].hidden
            }
            merge_lookup: dict[tuple[int, int], tuple[int, int]] = {}
            if merged_cells == "repeat":
                for merged in values_ws.merged_cells.ranges:
                    anchor = (merged.min_row, merged.min_col)
                    for row_index in range(merged.min_row, merged.max_row + 1):
                        for col_index in range(merged.min_col, merged.max_col + 1):
                            merge_lookup[(row_index, col_index)] = anchor

            rows: list[list[CellData]] = []
            for row_index in range(1, values_ws.max_row + 1):
                if not include_hidden and values_ws.row_dimensions[row_index].hidden:
                    continue
                row: list[CellData] = []
                for col_index in range(1, values_ws.max_column + 1):
                    if not include_hidden and col_index in hidden_cols:
                        continue
                    source_row, source_col = merge_lookup.get(
                        (row_index, col_index), (row_index, col_index)
                    )
                    value_cell = values_ws.cell(source_row, source_col)
                    formula_cell = formula_ws.cell(source_row, source_col)
                    formula = None
                    if formula_cell.data_type == "f":
                        raw_formula = str(formula_cell.value or "")
                        formula = raw_formula if raw_formula.startswith("=") else f"={raw_formula}"
                    row.append(
                        CellData(
                            value=value_cell.value,
                            formula=formula,
                            number_format=value_cell.number_format or "General",
                            hyperlink=(
                                value_cell.hyperlink.target
                                if getattr(value_cell, "hyperlink", None)
                                and value_cell.hyperlink.target
                                else None
                            ),
                        )
                    )
                rows.append(row)

            warnings: list[str] = []
            charts = len(getattr(values_ws, "_charts", []))
            images = len(getattr(values_ws, "_images", []))
            if charts:
                warnings.append(f"contains {charts} chart(s), which were not exported")
            if images:
                warnings.append(f"contains {images} image(s), which were not exported")
            if hidden and include_hidden:
                warnings.append("worksheet is hidden in the source workbook")
            sheets.append(
                SheetData(
                    title=values_ws.title,
                    rows=trim_matrix(rows),
                    hidden=hidden,
                    merged_ranges=len(values_ws.merged_cells.ranges),
                    charts=charts,
                    images=images,
                    source_rows=values_ws.max_row,
                    source_columns=values_ws.max_column,
                    warnings=warnings,
                )
            )
    finally:
        values_wb.close()
        formulas_wb.close()
    return sheets


def detect_text_encoding(path: Path, requested: str) -> str:
    if requested != "auto":
        return requested
    sample = path.read_bytes()[:65536]
    for candidate in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            sample.decode(candidate)
            return candidate
        except UnicodeDecodeError:
            continue
    return "latin-1"


def load_delimited(source: Path, encoding: str, delimiter: str) -> SheetData:
    selected_encoding = detect_text_encoding(source, encoding)
    text = source.read_text(encoding=selected_encoding)
    if delimiter == r"\t":
        delimiter = "\t"
    if delimiter == "auto":
        if source.suffix.lower() == ".tsv":
            delimiter = "\t"
        else:
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|").delimiter
            except csv.Error:
                delimiter = ","
    # StringIO preserves embedded newlines inside quoted CSV fields.
    parsed = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    width = max((len(row) for row in parsed), default=0)
    rows = [
        [CellData(value=value) for value in row + [""] * (width - len(row))]
        for row in parsed
    ]
    return SheetData(
        title=source.stem,
        rows=trim_matrix(rows),
        source_rows=len(parsed),
        source_columns=width,
        warnings=[f"decoded as {selected_encoding}; delimiter {delimiter!r}"],
    )


def load_legacy_xls(source: Path) -> list[SheetData]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise RuntimeError(
            "Legacy .xls conversion requires pandas and an Excel reader such as xlrd; "
            "resave the workbook as .xlsx if they are unavailable."
        ) from exc
    try:
        frames = pd.read_excel(source, sheet_name=None, header=None, dtype=object)
    except Exception as exc:
        raise RuntimeError(
            "Could not read the legacy .xls workbook. Install xlrd in the runtime or resave "
            "the workbook as .xlsx."
        ) from exc
    sheets: list[SheetData] = []
    for title, frame in frames.items():
        rows = []
        for values in frame.itertuples(index=False, name=None):
            rows.append(
                [
                    CellData(value=None if _is_pandas_na(value) else value)
                    for value in values
                ]
            )
        sheets.append(
            SheetData(
                title=str(title),
                rows=trim_matrix(rows),
                source_rows=len(frame.index),
                source_columns=len(frame.columns),
                warnings=["legacy .xls import does not preserve merged cells or hyperlinks"],
            )
        )
    return sheets


def _is_pandas_na(value: Any) -> bool:
    try:
        import pandas as pd

        result = pd.isna(value)
        return bool(result) if not hasattr(result, "__len__") else False
    except Exception:
        return False


def display_value(cell: CellData, formula_mode: str) -> str:
    value_text = format_scalar(cell.value, cell.number_format)
    if formula_mode == "formulas":
        text = cell.formula or value_text
    elif formula_mode == "both" and cell.formula:
        text = f"{value_text} ({cell.formula})" if value_text else cell.formula
    else:
        text = value_text or (cell.formula or "")
    text = clean_text(text)
    if cell.hyperlink and text:
        return f"[{escape_link_label(text)}]({escape_link_target(cell.hyperlink)})"
    if cell.hyperlink:
        return f"<{escape_link_target(cell.hyperlink)}>"
    return text


def format_scalar(value: Any, number_format: str = "General") -> str:
    if is_empty(value):
        return ""
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ", timespec="seconds").replace(" 00:00:00", "")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if math.isinf(value):
            return str(value)
        if "%" in (number_format or ""):
            decimals = decimal_places(number_format)
            return f"{value * 100:.{decimals}f}%"
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value)


def decimal_places(number_format: str) -> int:
    first_section = number_format.split(";", 1)[0]
    match = re.search(r"[0#]\.([0#]+)", first_section)
    return len(match.group(1)) if match else 0


def clean_text(text: str) -> str:
    text = str(text).replace("\r\n", "\n").replace("\r", "\n").strip()
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.replace("\n", "<br>")


def escape_link_label(text: str) -> str:
    return text.replace("\\", "\\\\").replace("[", "\\[").replace("]", "\\]")


def escape_link_target(text: str) -> str:
    return str(text).replace(" ", "%20").replace("(", "%28").replace(")", "%29")


def markdown_cell(text: str) -> str:
    return text.replace("\\", "\\\\").replace("|", "\\|")


def row_values(row: Sequence[CellData], formula_mode: str) -> list[str]:
    return [display_value(cell, formula_mode) for cell in row]


def nonempty_count(row: Sequence[CellData]) -> int:
    return sum(1 for cell in row if not is_empty(cell.value) or cell.formula)


def row_text_diversity(row: Sequence[CellData]) -> float:
    values = [
        str(cell.value).strip()
        for cell in row
        if not is_empty(cell.value) or cell.formula
    ]
    if not values:
        return 0.0
    return len(set(values)) / len(values)


def detect_header_index(rows: list[list[CellData]]) -> int | None:
    populated = [index for index, row in enumerate(rows[:30]) if nonempty_count(row)]
    if not populated:
        return None
    if len(populated) == 1:
        return populated[0]

    best_index = populated[0]
    best_score = float("-inf")
    for index in populated:
        row = rows[index]
        filled = nonempty_count(row)
        if not filled:
            continue
        strings = sum(
            1
            for cell in row
            if isinstance(cell.value, str) and bool(cell.value.strip())
        )
        next_rows = [candidate for candidate in rows[index + 1 : index + 5] if nonempty_count(candidate)]
        if not next_rows:
            next_density = 0.0
            similar_width = 0.0
            immediate_diversity = 0.0
        else:
            next_counts = [nonempty_count(candidate) for candidate in next_rows]
            next_density = sum(next_counts) / len(next_counts)
            similar_width = sum(1 for count in next_counts if count >= max(1, filled * 0.6)) / len(next_counts)
            immediate_diversity = row_text_diversity(next_rows[0])
        text_ratio = strings / filled
        diversity = row_text_diversity(row)
        singleton_penalty = 2.5 if filled == 1 and next_density > 1 else 0.0
        repeated_merge_penalty = 6.0 * (1.0 - diversity) if filled > 1 else 0.0
        score = (
            filled
            + 2.0 * text_ratio
            + 2.5 * similar_width
            + 3.0 * immediate_diversity
            - 0.03 * index
            - singleton_penalty
            - repeated_merge_penalty
        )
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def column_letters(count: int) -> list[str]:
    result = []
    for number in range(1, count + 1):
        label = ""
        current = number
        while current:
            current, remainder = divmod(current - 1, 26)
            label = chr(65 + remainder) + label
        result.append(label)
    return result


def build_headers(
    rows: list[list[CellData]], header_index: int | None, header_rows: int, formula_mode: str
) -> tuple[list[str], int]:
    width = max((len(row) for row in rows), default=0)
    fallback = [f"Column {letter}" for letter in column_letters(width)]
    if header_index is None:
        return fallback, 0

    header_end = min(len(rows), header_index + header_rows)
    headers: list[str] = []
    for col_index in range(width):
        parts: list[str] = []
        for row in rows[header_index:header_end]:
            if col_index >= len(row):
                continue
            text = display_value(row[col_index], formula_mode)
            if text and (not parts or parts[-1] != text):
                parts.append(text)
        headers.append(" / ".join(parts) or fallback[col_index])

    seen: dict[str, int] = {}
    unique: list[str] = []
    for header in headers:
        count = seen.get(header, 0) + 1
        seen[header] = count
        unique.append(header if count == 1 else f"{header} ({count})")
    return unique, header_end


def choose_layout(headers: Sequence[str], data_rows: Sequence[Sequence[CellData]], formula_mode: str) -> str:
    if len(headers) > 10:
        return "records"
    sample = data_rows[:100]
    values = [text for row in sample for text in row_values(row, formula_mode) if text]
    if not values:
        return "table"
    long_or_multiline = sum(1 for text in values if len(text) > 160 or "<br>" in text)
    return "records" if long_or_multiline / len(values) >= 0.08 else "table"


def notes_lines(
    rows: Sequence[Sequence[CellData]], header_index: int | None, formula_mode: str
) -> list[str]:
    if header_index is None or header_index <= 0:
        return []
    notes: list[str] = []
    for row in rows[:header_index]:
        values = [value for value in row_values(row, formula_mode) if value]
        if values:
            notes.append(" · ".join(values))
    return notes


def split_rows(rows: list[list[CellData]], limit: int) -> list[list[list[CellData]]]:
    if not limit or len(rows) <= limit:
        return [rows]
    return [rows[start : start + limit] for start in range(0, len(rows), limit)]


def make_document(
    *,
    sheet_title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[CellData]],
    notes: Sequence[str],
    layout: str,
    formula_mode: str,
    part: int,
    parts: int,
) -> str:
    title = sheet_title if parts == 1 else f"{sheet_title} (Part {part} of {parts})"
    lines = [f"# {title}", ""]
    if notes:
        lines.extend(["> " + note for note in notes])
        lines.append("")
    if layout == "table":
        lines.append("| " + " | ".join(markdown_cell(header) for header in headers) + " |")
        lines.append("| " + " | ".join("---" for _ in headers) + " |")
        for row in rows:
            values = row_values(row, formula_mode)
            values += [""] * (len(headers) - len(values))
            lines.append("| " + " | ".join(markdown_cell(value) for value in values[: len(headers)]) + " |")
    else:
        rendered_records = 0
        for index, row in enumerate(rows, start=1):
            values = row_values(row, formula_mode)
            pairs = [(header, values[i] if i < len(values) else "") for i, header in enumerate(headers)]
            pairs = [(header, value) for header, value in pairs if value]
            if not pairs:
                continue
            if rendered_records:
                lines.extend(["", "---", ""])
            title_value = next((value for _, value in pairs if value), f"Record {index}")
            title_value = re.sub(r"<br>.*", "", title_value)
            if len(title_value) > 100:
                title_value = title_value[:97].rstrip() + "..."
            lines.append(f"## {title_value}")
            lines.append("")
            for header, value in pairs:
                lines.append(f"- **{header}:** {value}")
            rendered_records += 1
    return "\n".join(lines).rstrip() + "\n"


def export_sheet(
    sheet: SheetData,
    *,
    output_dir: Path,
    file_index: int,
    layout_option: str,
    header_option: str | int | None,
    header_rows: int,
    formula_mode: str,
    max_rows_per_file: int,
) -> ExportResult | None:
    rows = sheet.rows
    if not rows or not any(nonempty_count(row) for row in rows):
        return None
    if header_option == "auto":
        header_index = detect_header_index(rows)
    elif header_option is None:
        header_index = None
    else:
        header_index = int(header_option) - 1
        if header_index >= len(rows):
            raise ValueError(
                f"Sheet {sheet.title!r} has only {len(rows)} visible row(s); "
                f"header row {header_index + 1} is out of range."
            )
    headers, data_start = build_headers(rows, header_index, header_rows, formula_mode)
    data_rows = [row for row in rows[data_start:] if nonempty_count(row)]
    if header_index is None:
        data_rows = [row for row in rows if nonempty_count(row)]
    layout = (
        choose_layout(headers, data_rows, formula_mode)
        if layout_option == "auto"
        else layout_option
    )
    notes = notes_lines(rows, header_index, formula_mode)
    chunks = split_rows(data_rows, max_rows_per_file)
    files: list[str] = []
    base = f"{file_index:02d}_{safe_filename(sheet.title)}"
    for part_index, chunk in enumerate(chunks, start=1):
        suffix = "" if len(chunks) == 1 else f"_part_{part_index:03d}"
        filename = f"{base}{suffix}.md"
        document = make_document(
            sheet_title=sheet.title,
            headers=headers,
            rows=chunk,
            notes=notes,
            layout=layout,
            formula_mode=formula_mode,
            part=part_index,
            parts=len(chunks),
        )
        (output_dir / filename).write_text(document, encoding="utf-8", newline="\n")
        files.append(filename)
    return ExportResult(
        sheet=sheet.title,
        layout=layout,
        records=len(data_rows),
        files=files,
        header_row=(header_index + 1 if header_index is not None else None),
        header_rows=(header_rows if header_index is not None else 0),
        warnings=sheet.warnings,
    )


def filter_sheets(sheets: list[SheetData], selected: Sequence[str]) -> list[SheetData]:
    if not selected:
        return sheets
    by_name = {sheet.title: sheet for sheet in sheets}
    missing = [name for name in selected if name not in by_name]
    if missing:
        available = ", ".join(by_name) or "(none)"
        raise ValueError(
            f"Worksheet(s) not found: {', '.join(missing)}. Available sheets: {available}"
        )
    return [by_name[name] for name in selected]


def remove_old_markdown(output_dir: Path) -> None:
    for old_file in output_dir.glob("*.md"):
        if old_file.is_file():
            old_file.unlink()


def main() -> int:
    args = parse_args()
    source = Path(args.source).expanduser().resolve()
    if not source.is_file():
        print(f"ERROR: source file not found: {source}", file=sys.stderr)
        return 2
    output_dir = (
        Path(args.output_dir).expanduser().resolve()
        if args.output_dir
        else source.with_name(f"{source.stem}_markdown")
    )
    if output_dir == source.parent and not args.keep_existing:
        print(
            "ERROR: refusing to clean Markdown files from the source directory; choose a dedicated --output-dir",
            file=sys.stderr,
        )
        return 2

    try:
        sheets = load_source(
            source,
            include_hidden=args.include_hidden,
            merged_cells=args.merged_cells,
            formulas=args.formulas,
            encoding=args.encoding,
            delimiter=args.delimiter,
        )
        sheets = filter_sheets(sheets, args.sheet)
        output_dir.mkdir(parents=True, exist_ok=True)
        if not args.keep_existing:
            remove_old_markdown(output_dir)

        header_option = parse_header_row(args.header_row)
        results: list[ExportResult] = []
        empty_sheets: list[str] = []
        for index, sheet in enumerate(sheets, start=1):
            result = export_sheet(
                sheet,
                output_dir=output_dir,
                file_index=index,
                layout_option=args.layout,
                header_option=header_option,
                header_rows=args.header_rows,
                formula_mode=args.formulas,
                max_rows_per_file=args.max_rows_per_file,
            )
            if result:
                results.append(result)
            else:
                empty_sheets.append(sheet.title)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    summary = {
        "source": str(source),
        "output_dir": str(output_dir),
        "documents": sum(len(result.files) for result in results),
        "sheets": [
            {
                "sheet": result.sheet,
                "layout": result.layout,
                "records": result.records,
                "header_row": result.header_row,
                "header_rows": result.header_rows,
                "files": result.files,
                "warnings": result.warnings,
            }
            for result in results
        ],
        "empty_sheets_skipped": empty_sheets,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
