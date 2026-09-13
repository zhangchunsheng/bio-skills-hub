# -*- coding: utf-8 -*-
"""
verify_output.py —— 班表生成结果校验（隐私安全版：不硬编码任何姓名）
校验三项：岗位硬约束、13:30 每周至多一次、每日在岗核心人数。

核心人员名单从输出文件自身的「名单」页自动识别；禁止岗位约束可选通过
--input 指向输入模板读取（输入模板的「禁止岗位」列）。本脚本不含任何真实姓名。

用法：
    python verify_output.py 2026年9月班表.xlsx
    python verify_output.py 2026年9月班表.xlsx --input 班表输入模板.xlsx
"""
import openpyxl, re, sys, os, argparse
from openpyxl import load_workbook
from collections import defaultdict
from datetime import date, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
from schedule_generator import parse_position_list, POSITIONS  # 复用岗位解析，避免漂移
ALL_POS = [p["name"] for p in POSITIONS]


def derive_core(wb):
    """从输出文件的「名单」页识别核心人员姓名（不含硬编码）。"""
    if "名单" not in wb.sheetnames:
        return []
    ws = wb["名单"]
    core = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        nm = str(row[0]).strip()
        typ = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if typ == "核心":
            core.append(nm)
    return core


def read_forbidden(input_path):
    """从输入模板的「人员名单」页读取 禁止岗位 约束 -> {name: set(positions)}。"""
    if not input_path or not os.path.exists(input_path):
        return {}
    wb = load_workbook(input_path, data_only=True)
    if "人员名单" not in wb.sheetnames:
        return {}
    ws = wb["人员名单"]
    forbidden = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        nm = str(row[0]).strip()
        raw = row[2] if len(row) > 2 else None
        if not raw:
            continue
        # 必须与生成器使用同一套解析：先按整串匹配岗位名（如“机动调配、麻精一调配”），
        # 再退化为按、拆分，避免把复合岗位名误拆成“机动调配”而误伤⑧机动调配岗。
        parsed = parse_position_list(raw, ALL_POS)
        if parsed:
            forbidden[nm] = set(parsed)
    return forbidden


def derive_own_month(wb):
    """从「总班表」页标题识别本班表归属的年/月（用于区分跨月拼接的 * 日）。"""
    if "总班表" not in wb.sheetnames:
        return None, None
    ws = wb["总班表"]
    title = ws.cell(1, 1).value or ""
    m = re.search(r'(\d{4})年(\d{1,2})月', title)
    if m:
        return int(m.group(1)), int(m.group(2))
    return None, None


def derive_holiday_days(wb):
    """从「总班表」自描述备注行解析假期日（如“假期日：25,1,2；盘点日：30；...”）。"""
    if "总班表" not in wb.sheetnames:
        return set()
    ws = wb["总班表"]
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and "假期日：" in v:
                m = re.search(r'假期日：([\d,\s]+)', v)
                if m:
                    return {int(x) for x in m.group(1).replace(" ", "").split(",") if x.strip()}
    return set()


def derive_inventory_days(wb):
    """从「总班表」自描述备注行解析盘点日（用于⑧机动调配排人日判定）。"""
    if "总班表" not in wb.sheetnames:
        return set()
    ws = wb["总班表"]
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and "盘点日：" in v:
                m = re.search(r'盘点日：([\d,\s]+)', v)
                if m:
                    return {int(x) for x in m.group(1).replace(" ", "").split(",") if x.strip()}
    return set()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="*", help="待校验的班表 xlsx 路径")
    ap.add_argument("--input", default=None, help="输入模板路径（用于读取禁止岗位约束）")
    args = ap.parse_args()

    files = args.files
    if not files:
        print("用法：python verify_output.py <班表.xlsx> [--input 模板.xlsx]")
        sys.exit(1)

    forbidden = read_forbidden(args.input)

    total_viol = 0
    total_v1330 = 0
    for fp in files:
        if not os.path.exists(fp):
            print(f"[跳过] 文件不存在：{fp}")
            continue
        print(f"\n校验文件：{fp}")
        wb = load_workbook(fp, data_only=True)
        core = derive_core(wb)
        if not core:
            print("  ⚠ 未能从「名单」页识别核心人员，跳过岗位/人数校验")
            continue
        print(f"  核心人员（自动识别 {len(core)} 人）：{core}")

        # 归属年月 + 假期日/盘点日（从总表自描述备注解析，避免重复计与误判）
        own_year, own_month = derive_own_month(wb)
        inventory_days = derive_inventory_days(wb)
        holiday_days = derive_holiday_days(wb)

        viol = 0
        week1330 = defaultdict(lambda: defaultdict(int))

        pos_rows = {
            4: ('片剂机配方', False, 'pos'),
            5: ('片剂审方、校对，出院药校对', False, 'pos'),
            6: ('片剂校对(核对机）、出院药调配', True, 'pos'),
            7: ('机动调配、麻精一调配', False, 'pos'),
            8: ('13:30岗位(机动调配)', False, 'pos'),
            9: ('窗口、汇总单、出院药调配', False, 'pos'),
            10: ('机动调配', True, 'p8b'),   # ⑧：仅周一与盘点日排人
            11: ('后勤', True, 'pos'),
            12: ('后勤', True, 'pos'),
            13: ('收单、收平板', False, 'derived'),
        }

        for sn in [s for s in wb.sheetnames if s.startswith('第') and s.endswith('周')]:
            w = wb[sn]
            dates = [w.cell(2, c).value for c in range(4, 11)]
            for c in range(4, 11):
                dv = dates[c - 4]
                if not dv or not isinstance(dv, str):
                    continue
                m = re.match(r'(\d{4})\.(\d{2})\.(\d{2})', dv)
                if not m:
                    continue
                y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
                dt = date(y, mo, d)

                # 跨月拼接日（*，归属邻月）非本月管辖范围，跳过校验（避免与邻月重复计 13:30 / 人数）
                if own_year and own_month and (dt.year != own_year or dt.month != own_month):
                    continue

                on_core = set()
                for r in range(1, 26):
                    v = w.cell(r, c).value
                    if not v:
                        continue
                    for nm in re.findall(r'[\u4e00-\u9fa5]{2,3}', str(v)):
                        if nm in core:
                            on_core.add(nm)

                if len(on_core) == len(core):
                    print(f"  · {dv} 全员在岗（{len(core)}人）→ 视为盘点日，跳过校验")
                    continue

                is_true_weekend = dt.weekday() >= 5 or d in holiday_days
                for r, (pname, wd_only, kind) in pos_rows.items():
                    v = w.cell(r, c).value
                    if not v:
                        # ⑧机动调配：仅周一与盘点日排人，且法定假期日不排（假期按周末模式）
                        if kind == 'p8b':
                            is_p8b_day = (dt.weekday() == 0 or d in inventory_days)
                            if is_p8b_day and d not in holiday_days:
                                print(f"  ✗ {dv} 岗位空缺：{pname} 无人（应为排人日）")
                                viol += 1
                            continue
                        if wd_only and is_true_weekend:
                            continue
                        print(f"  ✗ {dv} 岗位空缺：{pname} 无人")
                        viol += 1
                        continue
                    for nm in re.findall(r'[\u4e00-\u9fa5]{2,3}', str(v)):
                        if kind == 'derived':
                            if r == 8 and nm in core:
                                week1330[dt - timedelta(days=dt.weekday())][nm] += 1
                            continue
                        if nm in forbidden and pname in forbidden[nm]:
                            print(f"  ✗ {dv} {nm} 担任禁止岗位 {pname}")
                            viol += 1
                        if r == 8 and nm in core:
                            week1330[dt - timedelta(days=dt.weekday())][nm] += 1

                # 每日人数
                if is_true_weekend:
                    if len(on_core) != 5:
                        print(f"  ✗ {dv} 周末在岗核心 {len(on_core)} (应5，周末严格休4人)")
                        viol += 1
                else:
                    if len(on_core) not in (6, 7, 8):
                        print(f"  ✗ {dv} 工作日在岗核心 {len(on_core)} (应6~8)")
                        viol += 1

        v1330 = 0
        for monday, cnt in week1330.items():
            for nm, n in cnt.items():
                if n > 1:
                    print(f"  ✗ {monday}周 {nm} 13:30 排{n}次")
                    v1330 += 1

        print(f"13:30 违规周数: {v1330}")
        print(f"岗位约束/人数 违规总数: {viol}")
        total_viol += viol
        total_v1330 += v1330

    print(f"\n===== 合计：违规总数 {total_viol}，13:30 违规 {total_v1330} =====")
    print("校验完成" if (total_viol == 0 and total_v1330 == 0) else "存在违规，请检查")


if __name__ == "__main__":
    main()
