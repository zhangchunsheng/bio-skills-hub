# -*- coding: utf-8 -*-
"""严格自检：把生成文件原样 dump，交叉比对 总表 vs 周表，统计休息天数与岗位分布。"""
import os, re, importlib.util
from openpyxl import load_workbook
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
SG = os.path.join(HERE, "schedule_generator.py")
spec = importlib.util.spec_from_file_location("sg", SG)
sg = importlib.util.module_from_spec(spec); spec.loader.exec_module(sg)

CORE = []  # 每次运行由输出文件「名单」页动态识别，绝不硬编码姓名
NAME_RE = re.compile(r'[\u4e00-\u9fa5]{2,3}')

def _find_shelf_start(ws):
    """定位真实的「药架整理」区段起始行（即写有 SHELF_NOTE 的下一行）。
    注意：周表第 2 行 B 列也是“日期”，不能用它来定位药架区，否则会与表头混淆，
    导致真正的药架行（约 17~23 行）被漏扫 —— 这正是此前“总表/周表矛盾”的根因。"""
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value == sg.SHELF_NOTE:
            return r + 1
    # 兜底：退化为“最后一个”写有“日期”且同列是日期字符串的行
    found = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value == '日期' and ws.cell(r, 4).value:
            found = r
    return found

def derive_core(wb):
    """从输出文件「名单」页识别核心人员（类型==核心），避免硬编码任何姓名。"""
    if "名单" not in wb.sheetnames:
        return []
    ws = wb["名单"]
    out = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        nm = str(row[0]).strip()
        typ = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if typ == "核心":
            out.append(nm)
    return out

def core_names(v):
    if not v: return []
    return [n for n in NAME_RE.findall(str(v)) if n in CORE]

def load(fp):
    wb = load_workbook(fp, data_only=True)
    return wb

def parse_summary(wb):
    """新总表：人×日矩阵。返回 {date: set(当日标记“休”的核心姓名)} 。"""
    ws = wb["总班表"]
    title = ws.cell(1, 1).value or ""
    m = re.search(r'(\d{4})年(\d{1,2})月', title)
    yr, mo = int(m.group(1)), int(m.group(2))
    day_cols = {}
    for c in range(2, ws.max_column + 1):
        v = ws.cell(2, c).value
        if isinstance(v, int):
            day_cols[c] = v
    res = {}
    for r in range(4, ws.max_row + 1):
        nm = ws.cell(r, 1).value
        if nm not in CORE:
            continue
        for c, day in day_cols.items():
            v = ws.cell(r, c).value
            if v == '休':
                res.setdefault(date(yr, mo, day), set()).add(nm)
    return res

def parse_weekly(wb):
    """返回 {date: set(core working that day)} —— 来自周表（位置行+药架行）。"""
    work = {}
    for sn in wb.sheetnames:
        if not (sn.startswith('第') and sn.endswith('周')):
            continue
        ws = wb[sn]
        # 找日期行（row2）
        dates = {}
        for c in range(4, 12):
            v = ws.cell(2, c).value
            if not v: continue
            mm = re.match(r'(\d{4})\.(\d{2})\.(\d{2})', str(v))
            if mm:
                dates[c] = date(int(mm.group(1)), int(mm.group(2)), int(mm.group(3)))
        if not dates: continue
        # 药架整理区段起始行（用稳定的 SHELF_NOTE 锚点，避免误命中表头第 2 行）
        shelf_start = _find_shelf_start(ws)
        # 位置行 4..13
        for c, d in dates.items():
            names = set()
            for r in range(4, 14):
                names.update(core_names(ws.cell(r, c).value))
            # 药架行
            if shelf_start:
                for r in range(shelf_start + 1, shelf_start + 1 + len(sg.SHELF_TASKS)):
                    names.update(core_names(ws.cell(r, c).value))
            work.setdefault(d, set()).update(names)
    return work

def role_distribution(wb):
    """返回 {person: {role_label: count}}。"""
    roles = {}
    for sn in wb.sheetnames:
        if not (sn.startswith('第') and sn.endswith('周')):
            continue
        ws = wb[sn]
        # 位置行标签在 col B (rows 4..13)，对应 WEEKLY_ROWS
        row_labels = {}
        for r in range(4, 14):
            lab = ws.cell(r, 2).value
            row_labels[r] = lab or f"行{r}"
        dates = {}
        for c in range(4, 12):
            v = ws.cell(2, c).value
            if not v: continue
            mm = re.match(r'(\d{4})\.(\d{2})\.(\d{2})', str(v))
            if mm: dates[c] = date(int(mm.group(1)), int(mm.group(2)), int(mm.group(3)))
        for c in dates:
            for r in range(4, 14):
                for n in core_names(ws.cell(r, c).value):
                    roles.setdefault(n, {}).setdefault(row_labels[r], 0)
                    roles[n][row_labels[r]] += 1
            # 药架
            shelf_start = _find_shelf_start(ws)
            if shelf_start:
                for r in range(shelf_start + 1, shelf_start + 1 + len(sg.SHELF_TASKS)):
                    lab = ws.cell(r, 2).value or f"药架行{r}"
                    for n in core_names(ws.cell(r, c).value):
                        roles.setdefault(n, {}).setdefault(lab, 0)
                        roles[n][lab] += 1
    return roles

def run(fp):
    print("\n" + "=" * 70)
    print("文件:", fp)
    wb = load(fp)
    global CORE
    CORE = derive_core(wb)
    if not CORE:
        print("  ⚠ 未能从「名单」页识别核心人员，跳过自检")
        return 0, 0
    summary = parse_summary(wb)
    weekly = parse_weekly(wb)
    all_days = sorted(set(summary) | set(weekly))
    # 跨月拼接日（带*）在周表里也有，但 summary 只含本月；这里只查本月日
    # 取 summary 里的日期（即本月日）
    print("\n[A] 总表“休”标记 vs 周表真实休息 交叉比对（仅本月日，跨月拼接日不计入）")
    contradictions = 0
    for d in sorted(summary):   # summary 仅含本月日，跨月拼接日天然排除
        true_rest = set(CORE) - weekly.get(d, set())
        shown_rest = summary.get(d, set())
        bad = shown_rest.symmetric_difference(true_rest)
        if bad:
            contradictions += 1
            if contradictions <= 30:
                print(f"  ✗ {d} 总表标休={sorted(shown_rest)}  周表真休={sorted(true_rest)}  差异={sorted(bad)}")
    if contradictions == 0:
        print("  ✓ 总表“休”与周表真实休息 100% 一致（无任何一人既休又上 / 既上又休；亦无人永远不休息）")
    else:
        print(f"  >>> 共 {contradictions} 天存在总表/周表矛盾")

    print("\n[B] 每人休息天数（由周表推断：core - 当日上岗）")
    rest_count = {n: 0 for n in CORE}
    for d in sorted(weekly):
        for n in (set(CORE) - weekly[d]):
            rest_count[n] += 1
    for n in CORE:
        flag = "  <<< 异常：几乎不休息" if rest_count[n] <= 1 else ""
        print(f"  {n}: 休息 {rest_count[n]} 天{flag}")
    minr = min(rest_count.values())
    print(f"  >>> 最少休息天数 = {minr}" + ("  ⚠ 有人几乎不休息！" if minr <= 1 else ""))

    print("\n[C] 岗位分布（同一人是否总被钉在一个岗）")
    roles = role_distribution(wb)
    for n in CORE:
        dist = roles.get(n, {})
        total = sum(dist.values())
        if total == 0:
            print(f"  {n}: 未在周表任何岗位出现（异常？）"); continue
        top = max(dist, key=dist.get)
        top_pct = dist[top] / total * 100
        if top_pct >= 70 and total >= 5:
            others = [k for k in dist if k != top]
            print(f"  {n}: 总 {total} 次，{top} 占 {top_pct:.0f}%（其他：{others}）  <<< 疑似被钉死")
        else:
            print(f"  {n}: 总 {total} 次，主岗 {top} {top_pct:.0f}%，分布 {dist}")
    return contradictions, minr

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        files = sys.argv[1:]
    else:
        cwd = os.getcwd()
        files = sorted(
            os.path.join(cwd, fn) for fn in os.listdir(cwd)
            if fn.lower().endswith(".xlsx") and "班表" in fn and "输入模板" not in fn
        )
        if not files:
            print("未在当前目录发现班表文件；请传入文件路径，例如：")
            print("  python selfcheck.py 住院药房2026年9月班表.xlsx")
            sys.exit(1)
    for f in files:
        run(f)
