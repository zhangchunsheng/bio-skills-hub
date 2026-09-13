"""
PPT 解析服务
双通道：页图通道（soffice PDF + PyMuPDF）+ 文本通道（python-pptx）
"""
import asyncio
import json
import os
import subprocess
from pathlib import Path
from typing import Optional

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core.config import settings


class PPTParseResult:
    """PPT 解析结果"""
    def __init__(self):
        self.total_pages: int = 0
        self.page_images: list[str] = []       # 每页页图路径
        self.page_texts: list[str] = []        # 每页文本
        self.page_tables: list[Optional[str]] = []  # 每页表格 Markdown
        self.page_titles: list[str] = []       # 每页标题
        self.extracted_images: list[dict] = [] # 抠出的图片 [{path, page, context}]
        self.outline: list[dict] = []           # 大纲 [{index, title}]


async def ppt_to_pdf(ppt_path: str, output_dir: str) -> str:
    """soffice --headless --convert-to pdf"""
    pdf_path = os.path.join(output_dir, "converted.pdf")
    soffice = settings.soffice_path

    cmd = [
        soffice,
        "--headless",
        "--convert-to", "pdf",
        "--outdir", output_dir,
        ppt_path,
    ]
    try:
        proc = await asyncio.create_subprocess_exec(
            *cmd,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        stdout, stderr = await proc.communicate()
    except NotImplementedError:
        # Windows 事件循环不支持 subprocess，回退到同步方式
        import subprocess as sync_subprocess
        result = sync_subprocess.run(cmd, capture_output=True, timeout=120)
        stdout, stderr = result.stdout, result.stderr
        proc_returncode = result.returncode
    else:
        proc_returncode = proc.returncode

    if proc_returncode != 0:
        raise RuntimeError(f"LibreOffice 转换失败: {stderr.decode('utf-8', errors='replace')}")

    # soffice 输出的 PDF 文件名与原文件相同（扩展名变 .pdf）
    expected = os.path.join(output_dir, Path(ppt_path).stem + ".pdf")
    if os.path.exists(expected):
        return expected
    # 有时 soffice 输出到当前目录，尝试查找
    for f in os.listdir(output_dir):
        if f.endswith(".pdf"):
            return os.path.join(output_dir, f)
    raise FileNotFoundError(f"未找到转换后的 PDF 文件（输出目录: {output_dir}）")


def pdf_to_page_images(pdf_path: str, output_dir: str, dpi: int = 150) -> list[str]:
    """PyMuPDF 渲染 PDF 每页为 PNG"""
    doc = fitz.open(pdf_path)
    images = []
    for i, page in enumerate(doc):
        mat = fitz.Matrix(dpi / 72, dpi / 72)
        pix = page.get_pixmap(matrix=mat)
        img_path = os.path.join(output_dir, f"page_{i + 1:04d}.png")
        pix.save(img_path)
        images.append(img_path)
    doc.close()
    return images


def extract_text_from_pptx(ppt_path: str) -> tuple[list[str], list[Optional[str]], list[str], list[dict]]:
    """python-pptx 抽取文本、表格 Markdown、标题、抠图"""
    prs = Presentation(ppt_path)
    page_texts: list[str] = []
    page_tables: list[Optional[str]] = []
    page_titles: list[str] = []
    extracted_images: list[dict] = []

    for slide_idx, slide in enumerate(prs.slides):
        texts: list[str] = []
        tables: list[str] = []
        title_text = ""
        slide_images: list[dict] = []

        for shape in slide.shapes:
            # 某些PPT含 python-pptx 不认识的形状(如 SmartArt/内嵌OLE/新版图表)，
            # 访问其属性会抛 "unrecognized shape type"。逐形状兜底：坏形状跳过，不拖垮整页/整份解析。
            try:
                # 文本框
                if shape.has_text_frame:
                    para_texts: list[str] = []
                    for para in shape.text_frame.paragraphs:
                        run_text = "".join(run.text for run in para.runs)
                        if run_text.strip():
                            para_texts.append(run_text)
                    full_text = "\n".join(para_texts)
                    if full_text.strip():
                        texts.append(full_text)
                        # 第一个有内容的文本作为标题候选
                        if not title_text:
                            title_text = full_text[:80]

                # 表格
                if shape.has_table:
                    table = shape.table
                    md_rows: list[str] = []
                    # 先确定最大列数
                    max_cols = max(len(row.cells) for row in table.rows) if table.rows else 0
                    # 缓存每列当前有效的上下合并值
                    col_values: dict[int, str] = {}
                    for row_idx, row in enumerate(table.rows):
                        cells: list[str] = []
                        for cell_idx, cell in enumerate(row.cells):
                            if cell.is_spanned:
                                # 上下合并被覆盖的单元格，用缓存的值填充
                                cells.append(col_values.get(cell_idx, ""))
                                continue
                            val = cell.text.strip().replace("\n", " ")
                            # 读取 gridSpan，左右合并重复填充
                            span = int(cell._tc.get("gridSpan", 1))
                            for _ in range(span):
                                cells.append(val)
                            # 更新上下合并缓存（该列后续行被合并时用此值填充）
                            col_values[cell_idx] = val
                        # 补齐到 max_cols
                        while len(cells) < max_cols:
                            cells.append("")
                        if cells:
                            md_rows.append("| " + " | ".join(cells) + " |")
                            if row_idx == 0:
                                md_rows.append("| " + " | ".join(["---"] * len(cells)) + " |")
                    if md_rows:
                        tables.append("\n".join(md_rows))

                # 抠图
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    img = shape.image
                    img_bytes = img.blob
                    # 判断图片尺寸占比：shape 面积占 slide 比例
                    slide_width = prs.slide_width
                    slide_height = prs.slide_height
                    if slide_width and slide_height:
                        area_ratio = (shape.width * shape.height) / (slide_width * slide_height)
                    else:
                        area_ratio = 0
                    # 仅保留面积占比 > 5% 的图片（排除装饰性小图）
                    if area_ratio > 0.05:
                        ext = img.content_type.split("/")[-1] if "/" in img.content_type else "png"
                        img_filename = f"page_{slide_idx + 1:04d}_img_{len(slide_images) + 1}.{ext}"
                        slide_images.append({
                            "data": img_bytes,
                            "filename": img_filename,
                            "page": slide_idx + 1,
                            "context": title_text or f"第{slide_idx + 1}页",
                        })
            except Exception:
                # 无法识别/解析的形状：跳过，继续处理本页其余形状
                continue

        page_texts.append("\n".join(texts))
        page_tables.append("\n".join(tables) if tables else None)
        page_titles.append(title_text or f"第{slide_idx + 1}页")
        extracted_images.extend(slide_images)

    return page_texts, page_tables, page_titles, extracted_images


def build_outline(page_titles: list[str]) -> list[dict]:
    return [{"index": i + 1, "title": title} for i, title in enumerate(page_titles)]


async def parse_ppt(ppt_path: str, analysis_id: int) -> PPTParseResult:
    """完整 PPT 解析流程"""
    result = PPTParseResult()
    output_dir = os.path.join(settings.upload_dir, str(analysis_id))
    os.makedirs(output_dir, exist_ok=True)

    # 1. 文本通道
    page_texts, page_tables, page_titles, extracted_images = extract_text_from_pptx(ppt_path)
    result.page_texts = page_texts
    result.page_tables = page_tables
    result.page_titles = page_titles
    result.total_pages = len(page_texts)
    result.outline = build_outline(page_titles)

    # 2. 页图通道
    pdf_path = await ppt_to_pdf(ppt_path, output_dir)
    page_images = pdf_to_page_images(pdf_path, output_dir, dpi=150)
    result.page_images = page_images

    # 3. 保存抠图
    img_dir = os.path.join(output_dir, "images")
    os.makedirs(img_dir, exist_ok=True)
    for img_info in extracted_images:
        img_path = os.path.join(img_dir, img_info["filename"])
        with open(img_path, "wb") as f:
            f.write(img_info["data"])
        img_info["path"] = img_path
        del img_info["data"]
    result.extracted_images = extracted_images

    # 清理临时 PDF
    try:
        os.remove(pdf_path)
    except OSError:
        pass

    return result
