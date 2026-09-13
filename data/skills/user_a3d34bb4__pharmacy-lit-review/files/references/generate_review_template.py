# -*- coding: utf-8 -*-
"""
文献综述 Word 文档生成模板
==========================
使用说明：
1. 复制本脚本到工作目录
2. 替换 # ============ 文档内容 ============ 以下的内容
3. 运行：python generate_review.py
4. 生成的 .docx 文件在当前目录

依赖：python-docx (pip install python-docx)
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re

doc = Document()

# ============ 全局样式设置 ============
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(12)
font.color.rgb = RGBColor(0, 0, 0)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(0)

# 页面设置
sections = doc.sections
for section in sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

# ============ 辅助函数 ============

def set_run_font(run, font_name='宋体', size=12, bold=False, color=None):
    """设置 run 的字体（含东亚字体）"""
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if color:
        run.font.color.rgb = color

def add_title(text, level=0):
    """添加标题
    level=0: 文档标题（黑体16pt居中加粗）
    level=1: 一级标题（黑体14pt左对齐加粗）
    level=2: 二级标题（黑体12pt左对齐加粗）
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 0 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    if level == 0:
        set_run_font(run, '黑体', 16, True)
    elif level == 1:
        set_run_font(run, '黑体', 14, True)
    elif level == 2:
        set_run_font(run, '黑体', 12, True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_body_paragraph(text, first_line_indent=True):
    """添加正文段落，自动处理 [数字] 格式的上标引用"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line_indent:
        p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5

    # 使用正则表达式分割文本，处理 [数字] 格式的引用
    pattern = r'(\[\d+\])'
    parts = re.split(pattern, text)

    for part in parts:
        if re.match(r'^\[\d+\]$', part):
            # 引用标记，使用上标格式
            run = p.add_run(part)
            set_run_font(run, '宋体', 9)
            run.font.superscript = True
        else:
            run = p.add_run(part)
            set_run_font(run, '宋体', 12)

    return p

def add_ref_paragraph(num, text):
    """添加参考文献段落（悬挂缩进）"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-24)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(f'[{num}] ')
    set_run_font(run, '宋体', 10.5)
    run = p.add_run(text)
    set_run_font(run, '宋体', 10.5)
    return p

# ============ 文档内容 ============
# 以下内容请根据实际课题替换

# --- 标题 ---
add_title('【综述标题】', 0)

# --- 副标题/说明 ---
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('—基于PubMed与知网文献的系统梳理')
set_run_font(run, '楷体', 12, False)
p.paragraph_format.space_after = Pt(18)

# --- 摘要 ---
add_title('摘要', 1)
abstract_text = (
    '【在此填写摘要内容，200-300字。应包含研究背景、方法、主要发现和结论。】'
)
add_body_paragraph(abstract_text)

# --- 关键词 ---
p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Pt(24)
p.paragraph_format.line_spacing = 1.5
run = p.add_run('关键词：')
set_run_font(run, '宋体', 12, True)
run = p.add_run('关键词1；关键词2；关键词3；关键词4；关键词5')
set_run_font(run, '宋体', 12, False)

# --- 正文章节 ---
# 按照规划的结构，逐章节添加正文
# 每个段落使用 add_body_paragraph() 函数
# 引用标注直接写在文本中，如 '研究表明[1]...'

add_title('1 引言', 1)
add_body_paragraph('【引言正文...】')

add_title('2 【章节标题】', 1)
add_title('2.1 【小节标题】', 2)
add_body_paragraph('【正文...】')

# ... 继续添加其他章节 ...

add_title('8 总结与展望', 1)
add_body_paragraph('【总结正文...】')

# --- 参考文献 ---
add_title('参考文献', 1)

references = [
    '【作者1, 作者2, 作者3, 等. 中文文献题名[J]. 期刊名, 年, 卷(期): 起页-止页.】',
    '【Author AB, Author CD, Author EF, et al. English title[J]. Journal Name, Year, Volume(Issue): Pages. DOI: xxx.】',
    # ... 按引用顺序添加所有参考文献
]

for i, ref in enumerate(references, 1):
    add_ref_paragraph(i, ref)

# ============ 保存文档 ============
output_path = '【综述标题】.docx'  # 替换为实际输出路径
doc.save(output_path)
print(f'Document saved to: {output_path}')

# 统计字数
full_text = ''
for para in doc.paragraphs:
    full_text += para.text
chinese_chars = len(re.findall(r'[\u4e00-\u9fff]', full_text))
print(f'Chinese characters: {chinese_chars}')
print(f'Total characters: {len(full_text)}')
