import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from copy import copy
import os
import math
import re
import sys


def get_merged_range(cell, merged_ranges):
    for mr in merged_ranges:
        if cell.coordinate in mr:
            return mr
    return None


def get_effective_col_width(cell, ws):
    mr = get_merged_range(cell, list(ws.merged_cells.ranges))
    if mr:
        total_width = 0
        for c in range(mr.min_col, mr.max_col + 1):
            letter = get_column_letter(c)
            w = ws.column_dimensions[letter].width
            if w is None:
                w = 8.43
            total_width += w
        return total_width
    else:
        letter = get_column_letter(cell.column)
        w = ws.column_dimensions[letter].width
        if w is None:
            w = 8.43
        return w


def estimate_text_lines(text, col_width, font_size=11):
    if not text:
        return 1
    col_width_px = col_width * 7
    chars_per_line = max(1, int(col_width_px / (font_size * 0.6)))
    lines = str(text).split('\n')
    total_lines = 0
    for line in lines:
        line_len = len(line)
        if line_len > chars_per_line:
            total_lines += max(1, math.ceil(line_len / chars_per_line))
        else:
            total_lines += 1
    return total_lines


def get_font_size(cell):
    try:
        if cell.font and cell.font.size:
            return cell.font.size
    except:
        pass
    return 11


def is_meta_data(text):
    if text.startswith('---') or text.startswith('calction'):
        return True
    return False


def is_pure_number(value):
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        if re.match(r'^-?\d+(\.\d+)?$', value.strip()):
            return True
        if re.match(r'^\d+:[A-Z]\d+$', value.strip()):
            return True
    return False


def fix_sheet(ws, sheet_name, max_row_height=200):
    print(f"  Processing sheet: {sheet_name}")

    merged_ranges = list(ws.merged_cells.ranges)
    fixed_cells = []

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            if is_meta_data(text):
                continue
            if is_pure_number(cell.value):
                continue
            if len(text) <= 8:
                continue

            col_width = get_effective_col_width(cell, ws)
            font_size = get_font_size(cell)
            needed_lines = estimate_text_lines(text, col_width, font_size)
            has_newlines = '\n' in text

            if needed_lines > 1 or (has_newlines and len(text) > col_width * 1.5):
                mr = get_merged_range(cell, merged_ranges)

                if mr:
                    current_height = 0
                    for r in range(mr.min_row, mr.max_row + 1):
                        h = ws.row_dimensions[r].height
                        if h is None:
                            h = 15
                        current_height += h
                else:
                    h = ws.row_dimensions[cell.row].height
                    if h is None:
                        h = 15
                    current_height = h

                line_height = font_size * 1.5 + 4
                needed_height = needed_lines * line_height

                if needed_height > current_height * 1.1:
                    fixed_cells.append({
                        'cell': cell.coordinate,
                        'needed_height': needed_height,
                        'is_merged': mr is not None
                    })

                    if cell.alignment:
                        new_alignment = copy(cell.alignment)
                        new_alignment.wrapText = True
                    else:
                        new_alignment = Alignment(wrapText=True)
                    cell.alignment = new_alignment

    print(f"    Found {len(fixed_cells)} overflow cells")

    rows_to_adjust = {}
    for info in fixed_cells:
        coord = info['cell']
        cell = ws[coord]
        mr = get_merged_range(cell, merged_ranges)
        needed_height = info['needed_height']

        if mr:
            min_row = mr.min_row
            max_row = mr.max_row
            row_count = max_row - min_row + 1
            current_total = 0
            for r in range(min_row, max_row + 1):
                h = ws.row_dimensions[r].height
                if h is None:
                    h = 15
                current_total += h
            if needed_height > current_total:
                target_total = max(needed_height, current_total)
                per_row = target_total / row_count
                for r in range(min_row, max_row + 1):
                    h = ws.row_dimensions[r].height
                    if h is None:
                        h = 15
                    new_h = max(h, per_row)
                    if r not in rows_to_adjust or rows_to_adjust[r] < new_h:
                        rows_to_adjust[r] = new_h
        else:
            row = cell.row
            old_h = ws.row_dimensions[row].height
            if old_h is None:
                old_h = 15
            new_h = max(old_h, needed_height)
            if row not in rows_to_adjust or rows_to_adjust[row] < new_h:
                rows_to_adjust[row] = new_h

    for row, new_h in rows_to_adjust.items():
        if new_h > max_row_height:
            new_h = max_row_height
        ws.row_dimensions[row].height = new_h

    return len(fixed_cells)


import shutil

def process_file(input_path, output_dir):
    filename = os.path.basename(input_path)
    print(f"\nProcessing: {filename}")

    if not os.path.exists(input_path):
        print(f"  File not found, skipping")
        return False

    name, ext = os.path.splitext(filename)
    output_name = f"{name}_fixed{ext}"
    output_path = os.path.join(output_dir, output_name)

    try:
        wb = openpyxl.load_workbook(input_path, data_only=False)
    except Exception as e:
        print(f"  Error loading file: {e}, copying original instead")
        try:
            shutil.copy2(input_path, output_path)
            print(f"  Copied original: {output_name}")
            return True
        except Exception as ce:
            print(f"  Copy failed: {ce}")
            return False

    sheet_name = None
    for name_in_wb in wb.sheetnames:
        if "审定表" in name_in_wb:
            sheet_name = name_in_wb
            break

    if not sheet_name:
        print(f"  No '审定表' sheet found, copying original")
        try:
            wb.save(output_path)
            print(f"  Saved original: {output_name}")
            return True
        except Exception as e:
            print(f"  Save failed: {e}")
            return False

    ws = wb[sheet_name]
    fixed_count = fix_sheet(ws, sheet_name)

    try:
        wb.save(output_path)
        if fixed_count > 0:
            print(f"  Saved fixed: {output_name} ({fixed_count} cells fixed)")
        else:
            print(f"  Saved original: {output_name} (no overflow found)")
        return True
    except Exception as e:
        print(f"  Save failed: {e}")
        return False


def main():
    if len(sys.argv) < 2:
        print("Usage: python fix_overflow.py <input_folder> [output_folder]")
        print("Example: python fix_overflow.py C:\\Users\\Dell\\Desktop\\底稿")
        sys.exit(1)

    input_folder = sys.argv[1]

    if len(sys.argv) >= 3:
        output_dir = sys.argv[2]
    else:
        output_dir = os.path.join(input_folder, "fixed")

    os.makedirs(output_dir, exist_ok=True)

    excel_files = []
    for f in os.listdir(input_folder):
        if f.lower().endswith(('.xlsx', '.xls')):
            excel_files.append(os.path.join(input_folder, f))

    print(f"Found {len(excel_files)} Excel files")

    success_count = 0
    for file_path in excel_files:
        if process_file(file_path, output_dir):
            success_count += 1

    print(f"\nDone: {success_count}/{len(excel_files)} files processed")
    print(f"Output directory: {output_dir}")


if __name__ == "__main__":
    main()
