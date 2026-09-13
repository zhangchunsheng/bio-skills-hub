"""Extract text from the critical illness review DOCX to markdown reference."""
from docx import Document

doc = Document(r'C:/Users/m23002/Documents/LNJ/Cliams/重疾理赔审核要点.docx')
outpath = r'C:/Users/m23002/.workbuddy/skills/critical-illness-claims-review/references/critical_illness_review_guide.md'

with open(outpath, 'w', encoding='utf-8') as f:
    f.write('# 重大疾病理赔医学知识解读与审核指引（2021版）\n\n')
    f.write('> 本文件为技能 critical-illness-claims-review 的核心参考依据。\n')
    f.write('> 源自中国**保险股份有限公司理赔相关医学知识解读与审核指引（2021版）（征求意见版）。\n\n')
    f.write('---\n\n')

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            f.write(f'{text}\n\n')

    if doc.tables:
        f.write('\n---\n\n## 附录：表格数据\n\n')
        for ti, table in enumerate(doc.tables):
            f.write(f'\n### 表格 {ti+1}\n\n')
            for ri, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                f.write('| ' + ' | '.join(cells) + ' |\n')
                if ri == 0:
                    f.write('|' + '|'.join(['---' for _ in cells]) + '|\n')

import os
print('Reference file created.')
print(f'Size: {os.path.getsize(outpath)} bytes')
