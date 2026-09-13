import os
import re
from docx import Document


SECTION_TITLES = [
    "主诉",
    "现病史",
    "既往史",
    "过敏史",
    "个人史",
    "家族史",
    "体格检查",
    "辅助检查",
    "诊断",
    "处理",
    "处置意见",
]

def write_markdown(filepath: str, paragraphs: list[str]) -> None:
    """将脱敏后的段落写入Markdown文件。

    Args:
        filepath: 输出文件路径
        paragraphs: 脱敏后的段落列表
    """
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    content = "# 脱敏门诊病历\n\n"
    content += "\n\n".join(_format_markdown_paragraph(para) for para in paragraphs)
    content += "\n"

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)


def _format_markdown_paragraph(paragraph: str) -> str:
    text = paragraph.strip()
    if not text:
        return ""

    for title in SECTION_TITLES:
        match = re.match(rf"^{re.escape(title)}[：:]?\s*(.*)$", text)
        if match:
            content = match.group(1).strip()
            return f"## {title}\n\n{content}" if content else f"## {title}"

    basic_fields = _format_basic_fields(text)
    if basic_fields:
        return basic_fields

    return text


def _format_basic_fields(text: str) -> str:
    starts = [match.start() for match in re.finditer(r"(?=(性别|年龄|科别)[：:])", text)]
    if not starts:
        return ""
    if starts[0] != 0:
        return ""

    parts = []
    for index, start in enumerate(starts):
        end = starts[index + 1] if index + 1 < len(starts) else len(text)
        part = text[start:end].strip()
        if not part:
            continue
        label_match = re.match(r"^(性别|年龄|科别)[：:]?\s*(.*)$", part)
        if not label_match:
            continue
        label, value = label_match.group(1), label_match.group(2).strip()
        parts.append(f"* **{label}：** {value}" if value else f"* **{label}：**")
    return "\n".join(parts)

def write_docx(filepath: str, paragraphs: list[str]) -> None:
    """将脱敏后的段落写入Word文件。

    Args:
        filepath: 输出文件路径
        paragraphs: 脱敏后的段落列表
    """
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    doc = Document()
    doc.add_heading("病历记录", level=1)

    for para in paragraphs:
        if para.strip():
            doc.add_paragraph(para)

    doc.save(filepath)
