from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt

from layout_registry import FONT_MINIMUMS, LAYOUT_REGISTRY
from presentation_contracts import (
    compact_comparison_copy,
    contrast_ratio,
    validate_image_exception,
    validate_theme,
    validate_visual_spec,
)


PLACEHOLDER_PATTERN = re.compile(r"\{[^{}]+\}")
CJK_PATTERN = re.compile(r"[\u3400-\u9fff]")
ORPHAN_PRONE_PUNCTUATION = tuple("，。；：！？、")
ALIGNMENT_PATTERN = re.compile(
    r"ALIGNROW_(?P<group>[A-Za-z0-9-]+)_(?P<part>[A-Za-z0-9-]+)$"
)
ROLE_MINIMUMS = {
    "ROLE_cover_title": FONT_MINIMUMS["cover_title"],
    "ROLE_slide_title": FONT_MINIMUMS["slide_title"],
    "ROLE_subheading": FONT_MINIMUMS["subheading"],
    "ROLE_body": FONT_MINIMUMS["body"],
    "ROLE_footnote": FONT_MINIMUMS["footnote"],
    "ROLE_meta": FONT_MINIMUMS["footnote"],
}


def _intersection_area(first: Any, second: Any) -> int:
    left = max(first.left, second.left)
    top = max(first.top, second.top)
    right = min(first.left + first.width, second.left + second.width)
    bottom = min(first.top + first.height, second.top + second.height)
    return max(0, right - left) * max(0, bottom - top)


def _content_shapes(slide: Any) -> list[Any]:
    return [
        shape
        for shape in slide.shapes
        if shape.name.startswith("ROLE_") or shape.name.startswith("CONTENT_")
    ]


def _rgb_hex(color_format: Any) -> str | None:
    try:
        rgb = color_format.rgb
    except (AttributeError, TypeError, ValueError):
        return None
    return str(rgb) if rgb is not None else None


def _solid_fill_hex(shape: Any) -> str | None:
    try:
        if shape.fill.type is None:
            return None
        return _rgb_hex(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def _slide_background_hex(slide: Any, fallback: str) -> str:
    try:
        return _rgb_hex(slide.background.fill.fore_color) or fallback
    except (AttributeError, TypeError, ValueError):
        return fallback


def _contains(container: Any, content: Any) -> bool:
    center_x = content.left + content.width // 2
    center_y = content.top + content.height // 2
    return (
        container.left <= center_x <= container.left + container.width
        and container.top <= center_y <= container.top + container.height
    )


def _text_background_hex(
    slide: Any, shape: Any, preceding_shapes: list[Any], fallback: str
) -> str:
    background = _slide_background_hex(slide, fallback)
    for candidate in reversed(preceding_shapes):
        if (
            candidate is shape
            or "DECORATIVE_OVERFLOW" in candidate.name
            or not _contains(candidate, shape)
        ):
            continue
        candidate_fill = _solid_fill_hex(candidate)
        if candidate_fill:
            return candidate_fill
    return background


def _run_font_size(run: Any) -> float | None:
    return run.font.size.pt if run.font.size else None


def _shape_font_sizes(shape: Any) -> list[float]:
    if not getattr(shape, "has_text_frame", False):
        return []
    return [
        size
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if (size := _run_font_size(run)) is not None
    ]


def _visual_role_count(slide: Any, role: str) -> int:
    marker = f"VISROLE_{role}"
    return sum(marker in shape.name for shape in slide.shapes)


def _load_theme_for_qa(model: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    theme_name = str(model.get("request", {}).get("theme", "clean"))
    path = Path(__file__).resolve().parents[1] / "references" / "themes.json"
    violations: list[dict[str, Any]] = []
    try:
        themes = json.loads(path.read_text(encoding="utf-8"))
        theme = themes[theme_name]
        validate_theme(theme)
        return theme, violations
    except (KeyError, OSError, TypeError, ValueError, json.JSONDecodeError) as exc:
        violations.append({"theme": theme_name, "reason": str(exc)})
        return {
            "background": "FFFFFF",
            "surface": "FFFFFF",
            "ink": "000000",
        }, violations


def qa_pptx(path: Path, model: dict[str, Any]) -> dict[str, Any]:
    presentation = Presentation(path)
    theme, theme_violations = _load_theme_for_qa(model)
    out_of_bounds: list[dict[str, Any]] = []
    intentional_decorative_overflow: list[dict[str, Any]] = []
    unexpected_overlaps: list[dict[str, Any]] = []
    font_violations: list[dict[str, Any]] = []
    unresolved_placeholders: list[dict[str, Any]] = []
    east_asian_text_violations: list[dict[str, Any]] = []
    alignment_violations: list[dict[str, Any]] = []
    orphan_punctuation_violations: list[dict[str, Any]] = []
    chart_format_violations: list[dict[str, Any]] = []
    contrast_violations: list[dict[str, Any]] = []
    hierarchy_violations: list[dict[str, Any]] = []
    missing_visual_roles: list[dict[str, Any]] = []
    focus_violations: list[dict[str, Any]] = []
    density_violations: list[dict[str, Any]] = []
    repeated_scope_violations: list[dict[str, Any]] = []
    image_exception_violations: list[dict[str, Any]] = []
    native_chart_violations: list[dict[str, Any]] = []
    native_chart_count = 0
    declared_image_exceptions: dict[str, dict[str, Any]] = {}
    observed_image_exceptions: dict[str, int] = {}

    for candidate in [model, *list(model.get("slide_plan", []))]:
        raw_exceptions = candidate.get("image_exceptions", [])
        if raw_exceptions is None:
            continue
        if not isinstance(raw_exceptions, list):
            image_exception_violations.append(
                {"scope": "model", "reason": "image_exceptions_must_be_list"}
            )
            continue
        for exception in raw_exceptions:
            try:
                validate_image_exception(exception)
            except ValueError as exc:
                image_exception_violations.append(
                    {"scope": "model", "reason": str(exc)}
                )
                continue
            image_id = str(exception.get("image_id") or "")
            declared_slide_id = str(
                exception.get("slide_id") or candidate.get("slide_id") or ""
            )
            if not image_id or not declared_slide_id:
                image_exception_violations.append(
                    {
                        "scope": "model",
                        "reason": "image_exception_requires_image_id_and_slide_id",
                    }
                )
                continue
            if image_id in declared_image_exceptions:
                image_exception_violations.append(
                    {"image_id": image_id, "reason": "duplicate_image_exception_id"}
                )
                continue
            declared_image_exceptions[image_id] = {
                "slide_id": declared_slide_id,
                "exception": exception,
            }

    for slide_number, slide in enumerate(presentation.slides, start=1):
        planned_slide = (
            model.get("slide_plan", [])[slide_number - 1]
            if slide_number <= len(model.get("slide_plan", []))
            else {}
        )
        try:
            validate_visual_spec(planned_slide.get("visual_spec", {}))
        except ValueError as exc:
            theme_violations.append(
                {"slide": slide_number, "reason": str(exc)}
            )
        allows_clipped_decoration = planned_slide.get("layout_id") in {
            "cover",
            "data-appendix",
        }
        alignment_groups: dict[str, list[Any]] = {}
        shapes = list(slide.shapes)
        for shape_index, shape in enumerate(shapes):
            if getattr(shape, "shape_type", None) == 13:
                image_match = re.fullmatch(
                    r"(?:DECORATIVE_IMAGE_EXCEPTION|CONTENT_image)_(?P<image_id>[A-Za-z0-9_-]+)(?:_.*)?",
                    shape.name,
                )
                image_id = image_match.group("image_id") if image_match else ""
                declaration = declared_image_exceptions.get(image_id)
                planned_id = str(planned_slide.get("slide_id") or "")
                if not declaration:
                    image_exception_violations.append(
                        {
                            "slide": slide_number,
                            "shape": shape.name,
                            "reason": "image_not_declared_as_decorative_exception",
                        }
                    )
                elif declaration["slide_id"] != planned_id:
                    image_exception_violations.append(
                        {
                            "slide": slide_number,
                            "shape": shape.name,
                            "image_id": image_id,
                            "reason": "image_exception_declared_for_other_slide",
                        }
                    )
                else:
                    observed_image_exceptions[image_id] = (
                        observed_image_exceptions.get(image_id, 0) + 1
                    )
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > presentation.slide_width
                or shape.top + shape.height > presentation.slide_height
            ):
                shape_text = (
                    shape.text.strip()
                    if getattr(shape, "has_text_frame", False)
                    else ""
                )
                is_unlabelled_decoration = (
                    allows_clipped_decoration
                    and not shape_text
                    and not getattr(shape, "has_chart", False)
                    and not getattr(shape, "has_table", False)
                )
                target = (
                    intentional_decorative_overflow
                    if is_unlabelled_decoration
                    else out_of_bounds
                )
                target.append({"slide": slide_number, "shape": shape.name})
            if getattr(shape, "has_chart", False):
                native_chart_count += 1
                chart = shape.chart
                category_size = chart.category_axis.tick_labels.font.size
                value_size = chart.value_axis.tick_labels.font.size
                reasons: list[str] = []
                if category_size is None or category_size.pt < 12:
                    reasons.append("category_axis_font_below_12pt")
                if value_size is None or value_size.pt < 12:
                    reasons.append("value_axis_font_below_12pt")
                if chart.value_axis.tick_labels.number_format in {None, "General"}:
                    reasons.append("missing_number_format")
                if not chart.value_axis.has_title:
                    reasons.append("missing_axis_unit_title")
                if reasons:
                    chart_format_violations.append(
                        {"slide": slide_number, "shape": shape.name, "reasons": reasons}
                    )
            alignment_match = ALIGNMENT_PATTERN.search(shape.name)
            if alignment_match:
                alignment_groups.setdefault(
                    alignment_match.group("group"), []
                ).append(shape)
            if getattr(shape, "has_text_frame", False):
                text = shape.text or ""
                if (
                    "PUNCTGUARD" in shape.name
                    and text.rstrip().endswith(ORPHAN_PRONE_PUNCTUATION)
                ):
                    orphan_punctuation_violations.append(
                        {
                            "slide": slide_number,
                            "shape": shape.name,
                            "text": text,
                        }
                    )
                if PLACEHOLDER_PATTERN.search(text):
                    unresolved_placeholders.append(
                        {"slide": slide_number, "shape": shape.name, "text": text}
                    )
                minimum = next(
                    (
                        value
                        for prefix, value in ROLE_MINIMUMS.items()
                        if shape.name.startswith(prefix)
                    ),
                    None,
                )
                if minimum is not None:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            actual = run.font.size.pt if run.font.size else None
                            if actual is None or actual < minimum:
                                font_violations.append(
                                    {
                                        "slide": slide_number,
                                        "shape": shape.name,
                                        "minimum_pt": minimum,
                                        "actual_pt": actual,
                                    }
                                )
                if CJK_PATTERN.search(text):
                    for paragraph_index, paragraph in enumerate(
                        shape.text_frame.paragraphs
                    ):
                        paragraph_properties = paragraph._p.get_or_add_pPr()
                        missing_rules = [
                            name
                            for name, expected in (
                                ("eaLnBrk", "1"),
                                ("hangingPunct", "1"),
                            )
                            if paragraph_properties.get(name) != expected
                        ]
                        missing_language = any(
                            run._r.get_or_add_rPr().get("lang") != "zh-CN"
                            for run in paragraph.runs
                            if CJK_PATTERN.search(run.text or "")
                        )
                        if missing_rules or missing_language:
                            east_asian_text_violations.append(
                                {
                                    "slide": slide_number,
                                    "shape": shape.name,
                                    "paragraph": paragraph_index,
                                    "missing_rules": missing_rules,
                                    "missing_language": missing_language,
                                }
                            )

                background = _text_background_hex(
                    slide,
                    shape,
                    shapes[:shape_index],
                    theme.get("background", "FFFFFF"),
                )
                for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
                    for run_index, run in enumerate(paragraph.runs):
                        foreground = _rgb_hex(run.font.color)
                        size = _run_font_size(run)
                        if not foreground or not size or not (run.text or "").strip():
                            continue
                        is_metadata = (
                            shape.name.startswith(("ROLE_footnote", "ROLE_meta"))
                            or "VISROLE_eyebrow" in shape.name
                        )
                        is_large = size >= 24 or (size >= 18 and bool(run.font.bold))
                        minimum_contrast = 3.0 if is_large or is_metadata else 4.5
                        ratio = contrast_ratio(foreground, background)
                        if ratio + 1e-9 < minimum_contrast:
                            contrast_violations.append(
                                {
                                    "slide": slide_number,
                                    "shape": shape.name,
                                    "paragraph": paragraph_index,
                                    "run": run_index,
                                    "foreground": foreground,
                                    "background": background,
                                    "ratio": round(ratio, 2),
                                    "minimum": minimum_contrast,
                                }
                            )

        layout_id = planned_slide.get("layout_id")
        engineering_qa = (
            str(model.get("request", {}).get("analysis_intent") or "")
            == "engineering_regression"
            and planned_slide.get("content_exemption") == "qa"
        )
        legacy_qa = (
            planned_slide.get("content_exemption") == "qa"
            and not any(
                "content_refs" in page or "role" in page
                for page in model.get("slide_plan", [])
            )
        )
        enforce_editorial_contract = (
            isinstance(planned_slide.get("visual_spec"), dict)
            and not engineering_qa
            and not legacy_qa
        )
        comparison_contexts = planned_slide.get("visual_spec", {}).get(
            "comparison_contexts", []
        )
        if comparison_contexts:
            context_shapes = [
                shape
                for shape in slide.shapes
                if "VISROLE_comparison-context" in shape.name
            ]
            if len(context_shapes) != 1:
                repeated_scope_violations.append(
                    {
                        "slide": slide_number,
                        "reason": "comparison_context_count",
                        "actual": len(context_shapes),
                    }
                )
            for shape in slide.shapes:
                if (
                    "VISROLE_comparison-context" in shape.name
                    or not getattr(shape, "has_text_frame", False)
                ):
                    continue
                for context in comparison_contexts:
                    current = str(context.get("current_label") or "")
                    baseline = str(context.get("baseline_label") or "")
                    if (
                        current
                        and baseline
                        and current in (shape.text or "")
                        and baseline in (shape.text or "")
                    ):
                        repeated_scope_violations.append(
                            {
                                "slide": slide_number,
                                "shape": shape.name,
                                "reason": "period_repeated_outside_context",
                                "current_label": current,
                                "baseline_label": baseline,
                            }
                        )
        contract = LAYOUT_REGISTRY.get(layout_id, {})
        required_roles = list(contract.get("required_visual_roles", ()))
        if planned_slide.get("content_exemption") == "selection_meta":
            # The no-decision closure intentionally states that no resource
            # action is justified.  It uses the roadmap silhouette for visual
            # closure, but must not fabricate priority rows or success signals.
            required_roles = []
        if layout_id == "kpi-spotlight":
            period_kpis = model.get("period_overview", {}).get("kpis") or model.get("kpis", [])
            if len(period_kpis) <= 2 and "supporting-kpi" in required_roles:
                required_roles.remove("supporting-kpi")
        if enforce_editorial_contract:
            for visual_role in required_roles:
                if _visual_role_count(slide, visual_role) == 0:
                    missing_visual_roles.append(
                        {
                            "slide": slide_number,
                            "layout_id": layout_id,
                            "role": visual_role,
                        }
                    )
            expected_chart_ids = list(planned_slide.get("chart_ids", []))
            # chart_ids can also be evidence references on a non-chart layout
            # (for example risk/opportunity). A native object is mandatory only
            # where the layout contract actually requires the native-chart role.
            if expected_chart_ids and "native-chart" in required_roles:
                native_chart_ids = {
                    str(chart_id)
                    for chart_id in expected_chart_ids
                    if any(
                        getattr(shape, "has_chart", False)
                        and shape.name.startswith(f"CONTENT_chart_{chart_id}")
                        for shape in slide.shapes
                    )
                }
                missing_chart_ids = sorted(
                    set(str(chart_id) for chart_id in expected_chart_ids)
                    - native_chart_ids
                )
                if missing_chart_ids:
                    native_chart_violations.append(
                        {
                            "slide": slide_number,
                            "slide_id": planned_slide.get("slide_id"),
                            "reason": "required_native_chart_missing",
                            "chart_ids": missing_chart_ids,
                        }
                    )

        focus_expectations = {
            "cover": ("judgement", 1, 1),
            "kpi-spotlight": ("primary-kpi", 1, 2),
            "trend-wide": ("focus-callout", 1, 1),
            "chart-insight-action": ("focus-callout", 1, 1),
            "chart-side-kpi": ("focus-callout", 1, 1),
        }
        if enforce_editorial_contract and layout_id in focus_expectations:
            role, minimum_count, maximum_count = focus_expectations[layout_id]
            actual_count = _visual_role_count(slide, role)
            if not minimum_count <= actual_count <= maximum_count:
                focus_violations.append(
                    {
                        "slide": slide_number,
                        "layout_id": layout_id,
                        "role": role,
                        "actual": actual_count,
                        "expected": [minimum_count, maximum_count],
                    }
                )

        density_role = {
            "kpi-spotlight": ("-kpi",),
            "driver-levers": ("lever-node",),
            "action-roadmap": ("priority-row",),
        }.get(layout_id)
        if (
            enforce_editorial_contract
            and density_role
            and planned_slide.get("content_exemption") != "selection_meta"
        ):
            if layout_id == "kpi-spotlight":
                actual_items = sum(
                    "VISROLE_primary-kpi" in shape.name
                    or "VISROLE_supporting-kpi" in shape.name
                    for shape in slide.shapes
                )
            else:
                actual_items = _visual_role_count(slide, density_role[0])
            maximum_items = int(contract.get("max_items", 0))
            if actual_items > maximum_items:
                density_violations.append(
                    {
                        "slide": slide_number,
                        "layout_id": layout_id,
                        "actual": actual_items,
                        "maximum": maximum_items,
                    }
                )

        title_sizes = [
            size
            for shape in slide.shapes
            if shape.name.startswith("ROLE_slide_title")
            for size in _shape_font_sizes(shape)
        ]
        body_sizes = [
            size
            for shape in slide.shapes
            if shape.name.startswith("ROLE_body")
            for size in _shape_font_sizes(shape)
        ]
        if title_sizes and body_sizes and max(title_sizes) <= max(body_sizes):
            hierarchy_violations.append(
                {
                    "slide": slide_number,
                    "reason": "slide_title_not_larger_than_body",
                    "title_pt": max(title_sizes),
                    "body_pt": max(body_sizes),
                }
            )
        if layout_id == "kpi-spotlight":
            primary_sizes = [
                size
                for shape in slide.shapes
                if "VISROLE_primary-kpi" in shape.name
                for size in _shape_font_sizes(shape)
            ]
            supporting_sizes = [
                size
                for shape in slide.shapes
                if "VISROLE_supporting-kpi" in shape.name
                for size in _shape_font_sizes(shape)
            ]
            if primary_sizes and supporting_sizes and max(primary_sizes) <= max(supporting_sizes):
                hierarchy_violations.append(
                    {
                        "slide": slide_number,
                        "reason": "primary_kpi_not_larger_than_supporting_kpi",
                        "primary_pt": max(primary_sizes),
                        "supporting_pt": max(supporting_sizes),
                    }
                )

        for group, members in alignment_groups.items():
            centers = [member.top + member.height // 2 for member in members]
            anchors = [
                member.text_frame.vertical_anchor
                for member in members
                if getattr(member, "has_text_frame", False)
            ]
            reasons: list[str] = []
            if len(members) < 2:
                reasons.append("missing_peer")
            if centers and max(centers) - min(centers) > Pt(0.5):
                reasons.append("centerline_mismatch")
            if anchors and any(anchor != MSO_ANCHOR.MIDDLE for anchor in anchors):
                reasons.append("vertical_anchor_mismatch")
            if reasons:
                alignment_violations.append(
                    {
                        "slide": slide_number,
                        "group": group,
                        "shapes": [member.name for member in members],
                        "reasons": reasons,
                        "centerline_span_pt": round(
                            (max(centers) - min(centers)) / 12700, 2
                        )
                        if centers
                        else None,
                    }
                )

        content = _content_shapes(slide)
        for index, first in enumerate(content):
            for second in content[index + 1 :]:
                if _intersection_area(first, second) > 0:
                    unexpected_overlaps.append(
                        {
                            "slide": slide_number,
                            "first": first.name,
                            "second": second.name,
                        }
                    )

    all_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    missing_locked_content: list[str] = []
    for item in model.get("slide_plan", []):
        if item.get("slide_id") not in all_text:
            missing_locked_content.append(str(item.get("slide_id")))
        visual_spec = item.get("visual_spec", {})
        expected_claim = compact_comparison_copy(
            visual_spec.get("display_claim") or item.get("claim"),
            visual_spec.get("comparison_contexts", []),
        )
        if expected_claim not in all_text:
            missing_locked_content.append(str(expected_claim))

    for image_id, declaration in declared_image_exceptions.items():
        observed_count = observed_image_exceptions.get(image_id, 0)
        if observed_count != 1:
            image_exception_violations.append(
                {
                    "image_id": image_id,
                    "slide_id": declaration["slide_id"],
                    "observed_count": observed_count,
                    "reason": "image_exception_count_mismatch",
                }
            )

    issues = (
        theme_violations
        or out_of_bounds
        or unexpected_overlaps
        or font_violations
        or unresolved_placeholders
        or east_asian_text_violations
        or alignment_violations
        or orphan_punctuation_violations
        or chart_format_violations
        or missing_locked_content
        or contrast_violations
        or hierarchy_violations
        or missing_visual_roles
        or focus_violations
        or density_violations
        or repeated_scope_violations
        or image_exception_violations
        or native_chart_violations
    )
    return {
        "status": "failed" if issues else "passed",
        "slide_count": len(presentation.slides),
        "silhouette_count": len(
            {item.get("silhouette") for item in model.get("slide_plan", [])}
        ),
        "native_chart_count": native_chart_count,
        "native_chart_violations": native_chart_violations,
        "out_of_bounds": out_of_bounds,
        "intentional_decorative_overflow": intentional_decorative_overflow,
        "unexpected_overlaps": unexpected_overlaps,
        "font_violations": font_violations,
        "unresolved_placeholders": unresolved_placeholders,
        "east_asian_text_violations": east_asian_text_violations,
        "alignment_violations": alignment_violations,
        "orphan_punctuation_violations": orphan_punctuation_violations,
        "chart_format_violations": chart_format_violations,
        "missing_locked_content": missing_locked_content,
        "theme_violations": theme_violations,
        "contrast_violations": contrast_violations,
        "hierarchy_violations": hierarchy_violations,
        "missing_visual_roles": missing_visual_roles,
        "focus_violations": focus_violations,
        "density_violations": density_violations,
        "repeated_scope_violations": repeated_scope_violations,
        "image_exception_violations": image_exception_violations,
    }
