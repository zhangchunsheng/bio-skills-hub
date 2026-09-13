"""Check native text/shape placement and chart workbook caches without rendering."""
import io
import math
import posixpath
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pptx import Presentation
from openpyxl import load_workbook

NS={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def inspect_native_scene(path, scene):
    errors=[];checks=0;prs=Presentation(str(path))
    if len(prs.slides)!=len(scene['scene']):return ['native scene page count differs'],0
    for n,(slide,objects) in enumerate(zip(prs.slides,scene['scene']),1):
        expected=[o for o in objects if o['kind'] not in {'chart','chart-meta'}]
        actual=[s for s in slide.shapes if not s.has_chart]
        if len(actual)!=len(expected):errors.append(f'page {n}: native shape count differs');continue
        for shape,o in zip(actual,expected):
            box=([min(o['x1'],o['x2']),min(o['y1'],o['y2']),abs(o['x2']-o['x1']),abs(o['y2']-o['y1'])]
                 if o['kind']=='line' else [o[k] for k in ('x','y','w','h')])
            if any(abs(a-round(b*9525))>1 for a,b in zip((shape.left,shape.top,shape.width,shape.height),box)):
                errors.append(f'page {n}: native object moved or resized')
            if o['kind']=='text':
                if not shape.has_text_frame or ''.join(shape.text.split())!=''.join(o['t'].split()):
                    errors.append(f'page {n}: native scene text differs')
                else:
                    for p in shape.text_frame.paragraphs:
                        if p.font.size is None or abs(p.font.size.pt-o['size']*.75)>.02:
                            errors.append(f'page {n}: native font size differs')
            checks+=1
    # Inspect each embedded workbook through the chart relationship. A correct
    # visible cache with incorrect editable data must not pass.
    with ZipFile(path) as z:
        for name in z.namelist():
            if not re.fullmatch(r'ppt/charts/chart\d+\.xml',name):continue
            chart=ET.fromstring(z.read(name));ext=chart.find('c:externalData',NS)
            if ext is None:errors.append('native chart workbook missing');continue
            relname=posixpath.join(posixpath.dirname(name),'_rels',posixpath.basename(name)+'.rels')
            rels=ET.fromstring(z.read(relname));rid=ext.get('{'+NS['r']+'}id')
            link=next((r for r in rels if r.get('Id')==rid),None)
            if link is None or link.get('TargetMode')=='External':errors.append('workbook not embedded');continue
            target=posixpath.normpath(posixpath.join(posixpath.dirname(name),link.get('Target')))
            wb=load_workbook(io.BytesIO(z.read(target)),read_only=True,data_only=True)
            try:
                for tag,cache in [('numRef','numCache'),('strRef','strCache')]:
                    for ref in chart.findall('.//c:'+tag,NS):
                        formula=ref.findtext('c:f','',NS)
                        if '!' not in formula:errors.append('unresolved chart data reference');continue
                        sheet,area=formula.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
                        if sheet not in wb.sheetnames:errors.append('workbook sheet missing');continue
                        cells=wb[sheet][area.replace('$','')]
                        if not isinstance(cells,tuple):cells=((cells,),)
                        data=[c.value for row in cells for c in row]
                        points=sorted(ref.findall('c:'+cache+'/c:pt',NS),key=lambda p:int(p.get('idx')))
                        for pt in points:
                            index=int(pt.get('idx'));value=pt.findtext('c:v','',NS)
                            good=index<len(data)
                            if good:
                                try:good=(math.isclose(float(data[index]),float(value),rel_tol=1e-13,abs_tol=1e-10)
                                          if tag=='numRef' else str(data[index])==value)
                                except (ValueError,TypeError):good=False
                            if not good:errors.append('editable workbook differs from visible chart cache')
                            checks+=1
            finally:wb.close()
    return errors,checks
