"""Public python-pptx writer for bound chapter scenes; no raw data or host-private runtime."""
from __future__ import annotations
import argparse, hashlib, json, math, re
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE as CT, XL_LABEL_POSITION as LP, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX=9525
def emu(x): return Emu(round(x*PX))
def canonical(d): return hashlib.sha256(json.dumps(d,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
def rgb(c): return RGBColor.from_string(c.lstrip('#'))
def fmt(v): return f'{v:,.2f}'.rstrip('0').rstrip('.') if v!=int(v) else f'{v:,.0f}'
def precise(x): return float(f'{x:.15g}')

def wrap_label(value, columns):
    """Retain the existing author's label wrapping and indivisible numeric tokens."""
    raw=re.findall(r'\d{4}-\d{2}(?:-\d{2})?|[-+]?\d[\d,]*(?:\.\d+)?%?|[A-Za-z]+|[\s\S]',str(value))
    tokens=[]
    for token in raw:
        if token in '，。、；：！？）】》”’' and tokens and tokens[-1]!='\n':tokens[-1]+=token
        else:tokens.append(token)
    rows=[];row='';units=0
    for token in tokens:
        if token=='\n':rows.append(row);row='';units=0;continue
        width=sum(1 if ord(c)>255 else .55 for c in token)
        if width>columns:raise ValueError('Chart label token exceeds layout capacity')
        if row and units+width>columns:rows.append(row);row='';units=0
        row+=token;units+=width
    if row:rows.append(row)
    return '\n'.join(rows)
def axis_range(values):
    lo=min(0,*values);hi=max(0,*values);span=hi-lo or 1;rough=span*1.12/4;p=10**math.floor(math.log10(rough))
    step=next(x*p for x in [1,2,2.5,5,10] if x*p>=rough)
    return math.floor(lo/step)*step,max(step,math.ceil(hi*1.08/step)*step),step
def child(parent,tag,**attrs):
    x=OxmlElement(tag)
    for k,v in attrs.items():x.set(k,str(v))
    parent.append(x);return x
def font_setting(font,size,face,color,bold=False):
    font.name=face;font.size=Pt(size*.75);font.bold=bold;font.color.rgb=rgb(color)
    # East Asian font is explicit, never inherited from an English Office theme.
    r=font._rPr
    r.set('lang','zh-CN')
    for tag in ['a:ea','a:cs']:
        n=r.find(qn(tag))
        if n is None:n=child(r,tag)
        n.set('typeface',face)
def rich_label(labels,idx,text,face,color,size=18):
    n=OxmlElement('c:dLbl');labels.append(n);child(n,'c:idx',val=idx)
    tx=child(n,'c:tx');rich=child(tx,'c:rich');child(rich,'a:bodyPr');child(rich,'a:lstStyle');p=child(rich,'a:p');r=child(p,'a:r');rp=child(r,'a:rPr',sz=round(size*75))
    child(child(rp,'a:solidFill'),'a:srgbClr',val=color.lstrip('#'))
    for tag in ['a:latin','a:ea','a:cs']:child(rp,tag,typeface=face)
    rp.set('lang','zh-CN');child(r,'a:t').text=text;child(n,'c:dLblPos',val='t')
    for flag in ['showLegendKey','showVal','showCatName','showSerName','showPercent','showBubbleSize']:
        child(n,'c:'+flag,val=0)

def draw_chart(slide,o,v,config,series_indices=None):
    kind=v['kind'];palette=config['palette'];face=config['font'];horizontal=kind in ['bar','paired']
    si=list(range(len(v['series']))) if series_indices is None else series_indices
    order=list(range(len(v['labels'])));
    if horizontal:order.reverse()
    xy=kind=='scatter' or (kind=='line' and v['axes']['x']['scale']!='category')
    typ=CT.XY_SCATTER_LINES if kind=='line' and xy else CT.XY_SCATTER if xy else CT.LINE if kind=='line' else CT.DOUGHNUT if kind=='donut' else CT.BAR_CLUSTERED if horizontal else CT.COLUMN_STACKED if kind=='stacked' else CT.COLUMN_CLUSTERED
    data=XyChartData() if xy else CategoryChartData()
    labels=v['labels'][:]
    if kind=='histogram': labels=[f'[{fmt(a)}, {fmt(b)}{chr(93) if i==len(labels)-1 else chr(41)}' for i,(a,b) in enumerate(zip(v['bin_lower'],v['bin_upper']))]
    if not xy:data.categories=[wrap_label(labels[i],15 if horizontal else 10) for i in order]
    for index in si:
        s=v['series'][index];values=[precise(s['values'][i]) for i in order]
        if xy:
            series=data.add_series(s['label']);xs=[x/86400000+25569 if v['axes']['x']['scale']=='time' else x for x in v['x_values']]
            for x,y in zip(xs,values):series.add_data_point(precise(x),y)
        else:data.add_series(s['label'],values)
    chart=slide.shapes.add_chart(typ,emu(o['x']),emu(o['y']),emu(o['w']),emu(o['h']),data).chart
    cs=chart._chartSpace
    fills=[]
    for parent in [cs,cs.chart.plotArea]:
        prop=parent.find(qn('c:spPr'))
        if prop is None:
            prop=OxmlElement('c:spPr')
            parent.insert_element_before(prop,'c:txPr','c:externalData','c:printSettings','c:extLst')
        fills.append(prop)
    for prop in fills:
        child(child(prop,'a:solidFill'),'a:srgbClr',val=palette['background']);child(child(prop,'a:ln'),'a:noFill')
    chart.has_title=False;chart.has_legend=len(si)>1
    if chart.has_legend:
        chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False
        font_setting(chart.legend.font,22,face,palette['ink'])
    plot=chart.plots[0]
    if typ in [CT.BAR_CLUSTERED,CT.COLUMN_CLUSTERED,CT.COLUMN_STACKED]:plot.gap_width=0 if kind=='histogram' else 70;plot.overlap=100 if kind=='stacked' else 0
    if kind=='donut':plot._element.find(qn('c:holeSize')).set('val','62')
    for j,series in enumerate(chart.series):
        color=palette['chart'][si[j]%len(palette['chart'])];series.format.fill.solid();series.format.fill.fore_color.rgb=rgb(color)
        if kind=='line':series.format.line.color.rgb=rgb(color);series.format.line.width=Pt(2.25)
        else:series.format.line.fill.background()
        if xy:
            series.marker.style=XL_MARKER_STYLE.CIRCLE;series.marker.size=6;series.marker.format.fill.solid();series.marker.format.fill.fore_color.rgb=rgb(color);series.marker.format.line.fill.background()
        if kind=='donut':
            for i,point in enumerate(series.points):point.format.fill.solid();point.format.fill.fore_color.rgb=rgb(palette['chart'][i%len(palette['chart'])]);point.format.line.fill.background()
        if kind in ['bar','column','grouped','paired','stacked','histogram']:
            nf='0.0,,"百万"' if max(abs(n) for n in v['series'][si[j]]['values'])>=1e6 else '0' if all(n==int(n) for n in v['series'][si[j]]['values']) else '0.00'
        else:nf='0.00'
    if kind=='scatter':
        plot._element.find(qn('c:scatterStyle')).set('val','marker')
        el=OxmlElement('c:dLbls');chart.series[0]._element.insert_element_before(el,'c:trendline','c:errBars','c:xVal','c:yVal','c:smooth')
        for i in range(len(labels)):
            same=[j for j in range(len(labels)) if v['x_values'][j]==v['x_values'][i] and v['series'][0]['values'][j]==v['series'][0]['values'][i]]
            rich_label(el,i,'/'.join(str(j+1) for j in same) if same[0]==i else '',face,palette['ink'])
        child(el,'c:showLegendKey',val=0);child(el,'c:showVal',val=0);child(el,'c:showCatName',val=0);child(el,'c:showSerName',val=0)
    if not xy:plot.has_data_labels=kind!='donut' and len(labels)<=10
    if not xy and plot.has_data_labels:
        dl=plot.data_labels;dl.show_value=not xy;dl.show_legend_key=False;dl.show_category_name=False;dl.show_series_name=False
        dl.position=LP.CENTER if kind=='stacked' else LP.ABOVE if xy else LP.OUTSIDE_END
        font_setting(dl.font,21,face,'FFFFFF' if kind=='stacked' else palette['ink']);dl.number_format=v.get('_sample_format',nf);dl.number_format_is_linked=False
    if kind!='donut':
        extent=[x for i in si for x in v['series'][i]['values']]
        if kind=='stacked':extent=[x for s in v['series'] for key in ['bases','ends'] for x in s[key]]
        low,high,step=v.get('_sample_axis',axis_range(extent));va=chart.value_axis;va.minimum_scale=low;va.maximum_scale=high;va.major_unit=step;va.tick_labels.number_format='0' if step==int(step) else '0.00'
        va.has_major_gridlines=True;va.major_gridlines.format.line.color.rgb=rgb(palette['line']);va.major_gridlines.format.line.width=Pt(.75)
        for axis in [chart.category_axis,va]:
            font_setting(axis.tick_labels.font,20,face,palette['muted']);axis.format.line.color.rgb=rgb(palette['line']);axis.format.line.width=Pt(.75)
            # Tick labels stay outside chart plot; negative series never shifts categories over bars.
            old=axis._element.find(qn('c:tickLblPos'))
            if old is not None:old.set('val','low')
        if xy:
            xa=chart.category_axis;xs=[x/86400000+25569 if v['axes']['x']['scale']=='time' else x for x in v['x_values']];lo,hi=min(xs),max(xs)
            if lo==hi:lo-=max(1,abs(lo)*.05);hi+=max(1,abs(hi)*.05)
            xa.minimum_scale=lo;xa.maximum_scale=hi;xa.tick_labels.number_format='yyyy-mm' if v['axes']['x']['scale']=='time' else '0.00'
            if v['axes']['x']['scale']=='time':xa.major_unit=max(1,round((hi-lo)/6))
    return {'id':v['id'],'kind':kind,'series_indices':si,'source_order':order}

def render_deck(projection,scene,pages,out):
    if out.exists():raise ValueError('Use a new output path; existing files are preserved')
    expected=projection['sha256'];body={k:v for k,v in projection.items() if k!='sha256'}
    if canonical(body)!=expected or scene['source']!=expected:raise ValueError('Source binding mismatch')
    if projection['selection']['target']!='standard':raise ValueError('Probe supports standard target only')
    if not pages or len(set(pages))!=len(pages) or any(i<1 or i>len(projection['slides']) for i in pages):raise ValueError('Invalid explicit source pages')
    if len(scene['scene'])!=len(pages):raise ValueError('Scene page count mismatch')
    prs=Presentation();prs.slide_width=emu(scene['width']);prs.slide_height=emu(scene['height']);config=projection['selection']['config'];records=[]
    for local_index, srcpage in enumerate(pages):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);record={'source_page':srcpage,'source_id':projection['slides'][srcpage-1]['id'],'charts':[]}
        for o in scene['scene'][local_index]:
            kind=o['kind']
            if kind=='chart-meta':continue
            if kind=='chart':
                view=projection['chart_views'][o['id']]
                if view['kind']=='paired':
                    for j in range(len(view['series'])):record['charts'].append(draw_chart(slide,{**o,'x':o['x']+j*o['w']/2,'w':o['w']/2-20},view,config,[j]))
                else:record['charts'].append(draw_chart(slide,o,view,config))
                continue
            if kind=='text':
                shape=slide.shapes.add_textbox(emu(o['x']),emu(o['y']),emu(o['w']),emu(o['h']));tf=shape.text_frame;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.TOP;tf.text=o['t']
                for para in tf.paragraphs:
                    para.space_before=Pt(0);para.space_after=Pt(0)
                    para._p.get_or_add_pPr().set('eaLnBrk','1');para._p.get_or_add_pPr().set('hangingPunct','1')
                    font_setting(para.font,o['size'],config['font'],o['fill'],o.get('bold',False))
                    for run in para.runs:font_setting(run.font,o['size'],config['font'],o['fill'],o.get('bold',False))
            elif kind=='rect':
                shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,emu(o['x']),emu(o['y']),emu(o['w']),emu(o['h']));shape.fill.solid();shape.fill.fore_color.rgb=rgb(o['fill']);shape.line.fill.background()
            elif kind=='line':
                shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,emu(o['x1']),emu(o['y1']),emu(o['x2']),emu(o['y2']));shape.line.color.rgb=rgb(o['fill']);shape.line.width=Pt(o.get('width',1)*.75)
            else:raise ValueError('Unsupported scene object '+kind)
            # Suppress theme style references which otherwise add shadows to thin rules/bars.
            style=shape._element.find(qn('p:style'))
            if style is not None:shape._element.remove(style)
            child(shape._element.spPr,'a:effectLst')
        page=projection['slides'][srcpage-1]
        slide.notes_slide.notes_text_frame.text=json.dumps({'page':page,'source_model':projection['source_model_sha256'],'projection':expected,'scopes':[projection['scope_records'][e] for e in page['evidence_ids']]},ensure_ascii=False)
        records.append(record)
    out.parent.mkdir(parents=True,exist_ok=True)
    with out.open('xb') as stream:prs.save(stream)
    out.with_suffix('.receipt.json').write_text(json.dumps({'scope':scene.get('purpose','full-deck'),'projection_sha256':expected,'pages':records,'exporter':'python-pptx','chart_data':'frozen values; waterfall is editable shapes'},ensure_ascii=False,indent=2),encoding='utf-8')

def export_bound(model, selection, projection, scene, out, story=None):
    from chapter_presentation import validate_projection
    from chapter_presentation_qa import inspect
    validate_projection(projection, model, selection, story)
    if selection['target'] != 'standard':
        raise ValueError('Standard export cannot substitute for SlideViber')
    purpose = scene.get('purpose')
    indices = scene.get('page_indices')
    if purpose:
        if purpose not in {'design-sample', 'design-review'} or not isinstance(indices, list):
            raise ValueError('Unknown scene purpose or missing page mapping')
        if (not indices or any(type(i) is not int or i < 0 or i >= len(projection['slides']) for i in indices)
                or indices != sorted(set(indices)) or len(indices) != len(scene['scene'])):
            raise ValueError('Invalid sample source mapping')
        sample = [i+1 for i in indices] if purpose == 'design-sample' else None
        review = [projection['slides'][i]['id'] for i in indices] if purpose == 'design-review' else None
    else:
        if indices is not None:raise ValueError('Subset cannot be accepted as full deck')
        indices = list(range(len(projection['slides'])))
        sample = review = None
    if scene.get('source') != projection['sha256'] or (scene.get('width'),scene.get('height')) != (1280,720):
        raise ValueError('Scene source or dimensions differ')
    if len(scene['scene']) != len(indices):raise ValueError('Scene page count mismatch')
    # Renderer works on a subset while notes remain bound to full source pages.
    render_deck(projection, scene, [i+1 for i in indices], out)
    qa = inspect(out, projection, model, selection, scene, sample, story, review)
    from chapter_native_scene_qa import inspect_native_scene
    errors, checks = inspect_native_scene(out, scene)
    qa['errors'].extend(errors);qa['native_scene_checks']=checks;qa['passed']=not qa['errors']
    with out.with_suffix('.qa.json').open('x',encoding='utf-8') as stream:
        json.dump(qa, stream, ensure_ascii=False, indent=2)
    if not qa['passed']:raise ValueError('PPT QA failed: '+ '; '.join(qa['errors'][:5]))
    return qa


if __name__ == '__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    for field in ['model','selection','projection','scene','output']:
        ap.add_argument('--'+field,type=Path,required=True)
    ap.add_argument('--story',type=Path)
    a=ap.parse_args();read=lambda p:json.loads(p.read_text(encoding='utf-8'))
    result=export_bound(read(a.model),read(a.selection),read(a.projection),read(a.scene),a.output,read(a.story) if a.story else None)
    print(json.dumps(result,ensure_ascii=False))
