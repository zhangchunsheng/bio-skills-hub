from __future__ import annotations

import argparse
import json
import os
import re
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation

from presentation_contracts import compact_comparison_copy, compact_report_subtitle
from content_resolver import project_legacy_content_refs, resolve_model_for_page


PLACEHOLDER_PATTERN = re.compile(r"\{[^{}]+\}")
CJK_PATTERN = re.compile(r"[\u3400-\u9fff]")
HEADING_TERMINAL_PUNCTUATION = "。！？；，,.!?;：:、"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
A = f"{{{A_NS}}}"
NEAR_OPAQUE_ALPHA = 85_000


def _fmt(value: Any) -> str:
    if value is None:
        return "—"
    return f"{float(value):,.2f}"


def _format_metric_value(value: Any, unit: str) -> str:
    """Mirror the standard-deck display contract without coupling to its renderer."""
    if value is None:
        return "—"
    number = float(value)
    if unit == "%":
        return f"{number:,.2f}%"
    if unit == "x":
        return f"{number:,.2f}倍"
    if abs(number) >= 10_000:
        return f"{number / 10_000:,.2f}万"
    if unit == "单" or abs(number) >= 1_000:
        return f"{number:,.0f}"
    return f"{number:,.2f}"


def _headline(item: dict[str, Any], fallback_key: str) -> str:
    return str(item.get("headline") or item[fallback_key])


def _compact_heading(value: Any) -> str:
    return str(value or "").rstrip().rstrip(HEADING_TERMINAL_PUNCTUATION)


def _metric_lock(item: dict[str, Any], *, period: str, scope: str) -> dict[str, Any]:
    return {
        "lock_type": "metric",
        "lock_id": str(item.get("kpi_id") or item.get("name")),
        "name": str(item.get("display_name") or item.get("name")),
        "metric_name": str(item.get("name")),
        "value": item.get("value"),
        "unit": str(item.get("unit", "")),
        "period": period,
        "scope": scope,
        "evidence_ids": list(item.get("evidence_ids", [])),
    }


def _statement_lock(
    lock_id: str,
    statement: str,
    *,
    evidence_ids: list[str] | None = None,
    lock_type: str = "statement",
    required_tokens: list[str] | None = None,
) -> dict[str, Any]:
    return {
        "lock_type": lock_type,
        "lock_id": lock_id,
        "statement": statement,
        "required_tokens": list(required_tokens or [statement]),
        "evidence_ids": list(evidence_ids or []),
    }


def _content_binding(
    binding_type: str, binding_id: str, required_tokens: list[str]
) -> dict[str, Any]:
    return {
        "binding_type": binding_type,
        "binding_id": binding_id,
        "required_tokens": [str(token) for token in required_tokens if str(token)],
    }


def _metric_contract_lock(contract: dict[str, Any]) -> dict[str, Any]:
    """Lock the exact appendix contract selected by its stable contract ID."""

    name = str(contract.get("display_name") or contract.get("name") or "")
    formula = str(contract.get("formula_display") or contract.get("formula") or "")
    return {
        "lock_type": "metric_contract",
        "lock_id": str(contract.get("metric_id") or contract.get("name") or ""),
        "name": name,
        "formula": formula,
        "statement": f"{name}：{formula}",
    }


def _risk_opportunity_items(model: dict[str, Any]) -> list[dict[str, Any]]:
    return list(model.get("opportunity_insights", [])) + list(model.get("risk_insights", []))


def _chart_lock(chart: dict[str, Any]) -> dict[str, Any]:
    return {
        "chart_id": str(chart["chart_id"]),
        "metric_name": str(chart.get("metric_name", "")),
        "chart_type": str(chart.get("chart_type", "")),
        "display_scale": float(chart.get("display_scale", 1) or 1),
        "categories": [str(value) for value in chart.get("categories", [])],
        "series": [
            {
                "name": str(series.get("name", "")),
                "values": [float(value) for value in series.get("values", [])],
            }
            for series in chart.get("series", [])
        ],
        "evidence_ids": list(chart.get("evidence_ids", [])),
    }


def build_slide_lock_spec(model: dict[str, Any]) -> list[dict[str, Any]]:
    source_model = project_legacy_content_refs(model)
    specifications: list[dict[str, Any]] = []
    for slide in source_model["slide_plan"]:
        model = resolve_model_for_page(source_model, slide)
        chart_by_id = {item["chart_id"]: item for item in model.get("charts", [])}
        tokens = [slide["slide_id"]]
        visual_spec = slide.get("visual_spec", {})
        comparison_contexts = visual_spec.get("comparison_contexts", [])
        token_alternatives = [
            list(
                dict.fromkeys(
                    value
                    for value in (
                        str(slide["claim"]),
                        str(visual_spec.get("display_claim") or ""),
                    )
                    if value
                )
            )
        ]
        for context in comparison_contexts:
            tokens.extend(
                [str(context["current_label"]), str(context["baseline_label"])]
            )
        semantic_locks: list[dict[str, Any]] = []
        content_bindings: list[dict[str, Any]] = []
        layout = slide["layout_id"]
        if layout == "cover":
            subtitle = str(model.get("report_subtitle") or "")
            if subtitle:
                token_alternatives.append(list(dict.fromkeys(
                    [subtitle, compact_report_subtitle(subtitle)]
                )))
            tokens.append(model.get("judgement_headline", ""))
            semantic_locks.append(
                _statement_lock(
                    "report-purpose",
                    str(model.get("judgement_headline") or model["request"]["objective"]),
                    lock_type="judgement",
                )
            )
        elif layout == "executive-summary":
            conclusions = list(model.get("executive_summary", {}).get("key_conclusions", []))
            priorities = list(model.get("executive_summary", {}).get("priority_actions", []))
            conclusion_evidence = list(
                model.get("executive_summary", {}).get("key_conclusion_evidence_ids", [])
            )
            token_alternatives.extend(
                list(
                    dict.fromkeys(
                        (
                            statement,
                            compact_comparison_copy(statement, comparison_contexts),
                        )
                    )
                )
                for statement in conclusions
            )
            tokens.extend(priorities)
            semantic_locks.extend(
                _statement_lock(
                    f"conclusion-{index + 1:02d}",
                    statement,
                    evidence_ids=(
                        conclusion_evidence[index]
                        if index < len(conclusion_evidence)
                        and isinstance(conclusion_evidence[index], list)
                        else []
                    ),
                    lock_type="conclusion",
                    required_tokens=[
                        compact_comparison_copy(statement, comparison_contexts)
                    ],
                )
                for index, statement in enumerate(conclusions)
            )
            semantic_locks.extend(
                _statement_lock(
                    f"priority-{index + 1:02d}", statement, lock_type="action"
                )
                for index, statement in enumerate(priorities)
            )
        elif layout == "kpi-spotlight":
            overview = model.get("period_overview", {})
            period = str(overview.get("label", ""))
            kpis = list(overview.get("kpis") or model.get("kpis", []))
            for item in kpis:
                semantic_locks.append(
                    _metric_lock(item, period=period, scope="analysis_window")
                )
                tokens.extend(
                    [
                        str(item.get("display_name") or item["name"]),
                        _format_metric_value(item.get("value"), str(item.get("unit", ""))),
                    ]
                )
                content_bindings.append(
                    _content_binding(
                        "metric",
                        str(item.get("kpi_id") or item.get("name")),
                        [
                            str(item.get("display_name") or item["name"]),
                            _format_metric_value(
                                item.get("value"), str(item.get("unit", ""))
                            ),
                        ],
                    )
                )
        elif layout == "driver-levers":
            for item in model.get("levers", []):
                tokens.append(str(item["label"]))
                token_alternatives.append(
                    list(
                        dict.fromkeys(
                            (
                                str(item["headline"]),
                                compact_comparison_copy(
                                    item["headline"], comparison_contexts
                                ),
                            )
                        )
                    )
                )
                semantic_locks.append(
                    _statement_lock(
                        str(item.get("lever_id") or item["label"]),
                        str(item["headline"]),
                        evidence_ids=list(item.get("evidence_ids", [])),
                        lock_type="lever",
                        required_tokens=[compact_comparison_copy(
                            item["headline"], comparison_contexts
                        )],
                    )
                )
        elif layout == "risk-opportunity":
            visible_items = _risk_opportunity_items(model)
            token_alternatives.extend(
                list(
                    dict.fromkeys(
                        (
                            str(item["statement"]),
                            compact_comparison_copy(
                                item["statement"], comparison_contexts
                            ),
                        )
                    )
                )
                for item in visible_items
            )
            semantic_locks.extend(
                _statement_lock(
                    str(item.get("insight_id") or f"insight-{index + 1:02d}"),
                    str(item["statement"]),
                    evidence_ids=list(item.get("evidence_ids", [])),
                    lock_type="insight",
                )
                for index, item in enumerate(visible_items)
            )
        elif layout == "action-roadmap":
            visible_actions = list(model.get("actions", []))
            for index, action in enumerate(visible_actions):
                heading = _compact_heading(_headline(action, "text"))
                signal = str(action.get("verification_signal", ""))
                tokens.extend([heading, signal])
                semantic_locks.append(
                    _statement_lock(
                        str(action.get("action_id") or f"action-{index + 1:02d}"),
                        str(action.get("text", "")),
                        evidence_ids=list(action.get("evidence_ids", [])),
                        lock_type="action",
                        required_tokens=[token for token in (heading, signal) if token],
                    )
                )
                content_bindings.append(
                    _content_binding(
                        "action",
                        str(action.get("action_id") or f"action-{index + 1:02d}"),
                        [heading, signal],
                    )
                )
        elif layout == "data-appendix":
            tokens.extend(
                [
                    model["data_profile"]["date_min"],
                    model["data_profile"]["date_max"],
                    f"{model['data_profile']['row_count']:,}",
                ]
            )
            for contract in model.get("metric_contracts", []):
                contract_lock = _metric_contract_lock(contract)
                tokens.extend([contract_lock["name"], contract_lock["formula"]])
                semantic_locks.append(contract_lock)
        chart_locks = [
            _chart_lock(chart_by_id[chart_id])
            for chart_id in slide.get("chart_ids", [])
            if chart_id in chart_by_id
        ]
        specifications.append(
            {
                "slide_id": slide["slide_id"],
                "locked_tokens": list(dict.fromkeys(str(token) for token in tokens if token)),
                "locked_token_alternatives": [
                    alternatives
                    for alternatives in token_alternatives
                    if alternatives
                ],
                "semantic_locks": semantic_locks,
                "chart_locks": chart_locks,
                "content_bindings": content_bindings,
                "evidence_ids": list(slide.get("evidence_ids", [])),
            }
        )
    return specifications


def _normalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(value))
    return "".join(
        character
        for character in normalized
        if not character.isspace() and character not in {",", "，"}
    )


def _semantic_lock_token(lock: dict[str, Any]) -> str:
    """Return the human-visible value that carries a locked semantic meaning."""

    return str(lock.get("formula") or lock.get("statement") or "")


def _chart_payload_from_native_shape(shape: Any) -> dict[str, Any]:
    chart = shape.chart
    chart_type = (
        "bar"
        if chart._chartSpace.xpath(".//c:barChart")
        else "line"
        if chart._chartSpace.xpath(".//c:lineChart")
        else ""
    )
    series_nodes = chart._chartSpace.xpath(".//c:ser")
    series: list[dict[str, Any]] = []
    categories: list[str] = []
    for index, node in enumerate(series_nodes):
        name_nodes = node.xpath("./c:tx//c:v")
        value_nodes = node.xpath("./c:val//c:pt/c:v")
        category_nodes = node.xpath("./c:cat//c:pt/c:v")
        if index == 0:
            categories = [str(value.text or "") for value in category_nodes]
        series.append(
            {
                "name": str(name_nodes[0].text or "") if name_nodes else "",
                "values": [float(value.text) for value in value_nodes],
            }
        )
    return {"chart_type": chart_type, "categories": categories, "series": series}


def _same_chart_payload(actual: dict[str, Any], expected: dict[str, Any]) -> bool:
    if actual.get("chart_type") != expected.get("chart_type"):
        return False
    if actual.get("categories") != expected.get("categories"):
        return False
    actual_series = actual.get("series", [])
    expected_series = expected.get("series", [])
    if len(actual_series) != len(expected_series):
        return False
    for actual_item, expected_item in zip(actual_series, expected_series):
        if actual_item.get("name") != expected_item.get("name"):
            return False
        if len(actual_item.get("values", [])) != len(expected_item.get("values", [])):
            return False
        display_scale = float(expected.get("display_scale", 1) or 1)
        if any(
            abs(float(left) - float(right) / display_scale) > 1e-9
            for left, right in zip(actual_item["values"], expected_item["values"])
        ):
            return False
    return True


def _shape_is_hidden(shape: Any) -> bool:
    if str(shape._element.get("hidden") or "").lower() in {"1", "true"}:
        return True
    return any(
        str(node.get("hidden") or "").lower() in {"1", "true"}
        for node in shape._element.xpath(".//p:cNvPr")
    )


def _drawingml_solid_fill_alpha(shape: Any) -> int:
    """Return DrawingML solid-fill opacity (100000 is fully opaque).

    ``python-pptx`` does not expose a reliable ``FillFormat.transparency``;
    alpha lives under the shape's direct ``spPr`` solid-fill color node in the
    source OOXML instead. Text-body and line alpha transforms must not affect
    the rectangle's own fill opacity.
    Unknown/malformed alpha is treated as opaque so it cannot bypass a content
    visibility lock.
    """

    alpha_nodes = shape._element.xpath(
        "./*[local-name()='spPr']/a:solidFill/*/a:alpha"
    )
    if not alpha_nodes:
        return 100_000
    try:
        alpha = int(alpha_nodes[-1].get("val"))
    except (TypeError, ValueError):
        return 100_000
    return alpha if 0 <= alpha <= 100_000 else 100_000


def _opaque_foreground_overlap(target: Any, cover: Any) -> tuple[int, int, int, int] | None:
    """Return clipped near-opaque overlap, excluding transparent and line-only shapes."""

    try:
        if (
            int(cover.fill.type) != 1
            or _drawingml_solid_fill_alpha(cover) < NEAR_OPAQUE_ALPHA
        ):
            return None
    except (AttributeError, TypeError, ValueError):
        return None
    left = max(target.left, cover.left)
    top = max(target.top, cover.top)
    right = min(target.left + target.width, cover.left + cover.width)
    bottom = min(target.top + target.height, cover.top + cover.height)
    if right <= left or bottom <= top:
        return None
    return left, top, right, bottom


def _rectangular_union_area(rectangles: list[tuple[int, int, int, int]]) -> int:
    """Compute an axis-aligned rectangle union without double-counting overlaps."""

    x_edges = sorted({edge for rectangle in rectangles for edge in (rectangle[0], rectangle[2])})
    area = 0
    for left, right in zip(x_edges, x_edges[1:]):
        if right <= left:
            continue
        intervals = sorted(
            (top, bottom)
            for rect_left, top, rect_right, bottom in rectangles
            if rect_left < right and rect_right > left
        )
        covered_height = 0
        previous_top = previous_bottom = None
        for top, bottom in intervals:
            if previous_top is None:
                previous_top, previous_bottom = top, bottom
            elif top > previous_bottom:
                covered_height += previous_bottom - previous_top
                previous_top, previous_bottom = top, bottom
            else:
                previous_bottom = max(previous_bottom, bottom)
        if previous_top is not None:
            covered_height += previous_bottom - previous_top
        area += (right - left) * covered_height
    return area


def _opaque_foreground_union_coverage(target: Any, covers: list[Any]) -> float:
    rectangles = [
        overlap
        for cover in covers
        if not _shape_is_hidden(cover)
        for overlap in [_opaque_foreground_overlap(target, cover)]
        if overlap is not None
    ]
    if not rectangles:
        return 0.0
    area = _rectangular_union_area(rectangles)
    target_area = max(1, target.width * target.height)
    return area / target_area


def _is_visible_binding_shape(shape: Any, presentation: Presentation, slide: Any) -> bool:
    """Reject hidden, off-canvas, or metadata-sized shapes as visible content."""

    if _shape_is_hidden(shape):
        return False
    minimum_width = 457200  # 0.5in
    minimum_height = 182880  # 0.2in
    on_canvas = (
        shape.left >= 0
        and shape.top >= 0
        and shape.left + shape.width <= presentation.slide_width
        and shape.top + shape.height <= presentation.slide_height
        and shape.width >= minimum_width
        and shape.height >= minimum_height
    )
    if not on_canvas:
        return False
    shapes = list(slide.shapes)
    try:
        z_index = shapes.index(shape)
    except ValueError:
        return False
    return _opaque_foreground_union_coverage(shape, shapes[z_index + 1 :]) < 0.85


def normalize_east_asian_text(path: Path) -> int:
    """Add PowerPoint East Asian line-break metadata without rebuilding the deck."""
    source_path = Path(path)
    with tempfile.NamedTemporaryFile(
        prefix=f".{source_path.stem}-", suffix=".pptx", dir=source_path.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)

    changed = 0
    try:
        with zipfile.ZipFile(source_path, "r") as source, zipfile.ZipFile(
            temporary_path, "w"
        ) as target:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", entry.filename):
                    root = etree.fromstring(payload)
                    for paragraph in root.findall(f".//{A}p"):
                        paragraph_text = "".join(
                            node.text or "" for node in paragraph.findall(f".//{A}t")
                        )
                        if not CJK_PATTERN.search(paragraph_text):
                            continue
                        paragraph_properties = paragraph.find(f"{A}pPr")
                        if paragraph_properties is None:
                            paragraph_properties = etree.Element(f"{A}pPr")
                            paragraph.insert(0, paragraph_properties)
                        paragraph_changed = False
                        for name, value in (
                            ("eaLnBrk", "1"),
                            ("latinLnBrk", "0"),
                            ("hangingPunct", "1"),
                        ):
                            if paragraph_properties.get(name) != value:
                                paragraph_properties.set(name, value)
                                paragraph_changed = True
                        for run in paragraph.findall(f"{A}r"):
                            run_text = "".join(
                                node.text or "" for node in run.findall(f"{A}t")
                            )
                            if not CJK_PATTERN.search(run_text):
                                continue
                            run_properties = run.find(f"{A}rPr")
                            if run_properties is None:
                                run_properties = etree.Element(f"{A}rPr")
                                run.insert(0, run_properties)
                            if run_properties.get("lang") != "zh-CN":
                                run_properties.set("lang", "zh-CN")
                                paragraph_changed = True
                        if paragraph_changed:
                            changed += 1
                    payload = etree.tostring(
                        root,
                        encoding="UTF-8",
                        xml_declaration=True,
                        standalone=True,
                    )
                target.writestr(entry, payload)
        os.replace(temporary_path, source_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return changed


def qa_slideviber_pptx(path: Path, model: dict[str, Any]) -> dict[str, Any]:
    presentation = Presentation(path)
    lock_spec = build_slide_lock_spec(model)
    missing_locked_content: list[dict[str, Any]] = []
    unresolved_placeholders: list[dict[str, Any]] = []
    chart_lock_warnings: list[dict[str, Any]] = []
    semantic_lock_violations: list[dict[str, Any]] = []
    chart_lock_violations: list[dict[str, Any]] = []
    content_binding_violations: list[dict[str, Any]] = []
    slide_count_matches = len(presentation.slides) == len(lock_spec)
    slides_by_id: dict[str, Any] = {}
    slide_id_metadata_issues: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        name = str(slide.element.cSld.get("name") or "")
        prefix = "sheet-to-report:"
        slide_id = name[len(prefix):] if name.startswith(prefix) else ""
        if not slide_id:
            slide_id_metadata_issues.append({"slide_number": index, "reason": "missing_slide_id_metadata"})
        elif slide_id in slides_by_id:
            slide_id_metadata_issues.append({"slide_number": index, "slide_id": slide_id, "reason": "duplicate_slide_id_metadata"})
        else:
            slides_by_id[slide_id] = slide

    for specification in lock_spec:
        slide = slides_by_id.get(specification["slide_id"])
        if slide is None:
            for token in specification["locked_tokens"]:
                missing_locked_content.append(
                    {
                        "slide_id": specification["slide_id"],
                        "token": token,
                        "reason": "missing_slide",
                    }
                )
            for alternatives in specification.get("locked_token_alternatives", []):
                missing_locked_content.append(
                    {
                        "slide_id": specification["slide_id"],
                        "token": " | ".join(alternatives),
                        "reason": "missing_slide",
                    }
                )
            continue
        text = "\n".join(
            shape.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
        normalized_text = _normalize(text)
        shape_names = {
            str(getattr(shape, "name", ""))
            for shape in slide.shapes
        }
        audience_text = "\n".join(
            shape.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
            and not str(getattr(shape, "name", "")).startswith("CONTENT_chart_metadata_")
        )
        if PLACEHOLDER_PATTERN.search(audience_text):
            unresolved_placeholders.append(
                {"slide_id": specification["slide_id"], "text": audience_text}
            )
        for token in specification["locked_tokens"]:
            if _normalize(token) not in normalized_text:
                missing_locked_content.append(
                    {
                        "slide_id": specification["slide_id"],
                        "token": token,
                        "reason": "not_found",
                    }
                )
        for alternatives in specification.get("locked_token_alternatives", []):
            if not any(_normalize(token) in normalized_text for token in alternatives):
                missing_locked_content.append(
                    {
                        "slide_id": specification["slide_id"],
                        "token": " | ".join(alternatives),
                        "reason": "no_semantic_alternative_found",
                    }
                )
        for lock in specification.get("semantic_locks", []):
            required_tokens = list(lock.get("required_tokens") or [_semantic_lock_token(lock)])
            for token in required_tokens:
                token = str(token)
                if token and _normalize(token) not in normalized_text:
                    semantic_lock_violations.append(
                        {
                            "slide_id": specification["slide_id"],
                            "lock_id": str(lock.get("lock_id") or ""),
                            "lock_type": str(lock.get("lock_type") or ""),
                            "token": token,
                            "reason": "semantic_lock_missing",
                        }
                    )
        for binding in specification.get("content_bindings", []):
            prefix = (
                f"CONTENT_binding_{binding['binding_type']}_{binding['binding_id']}"
            )
            binding_shapes = [
                shape
                for shape in slide.shapes
                if str(getattr(shape, "name", "")).startswith(prefix)
                and getattr(shape, "has_text_frame", False)
                and _is_visible_binding_shape(shape, presentation, slide)
            ]
            binding_text = "\n".join(
                shape.text or ""
                for shape in binding_shapes
            )
            missing_tokens = [
                token
                for token in binding.get("required_tokens", [])
                if _normalize(token) not in _normalize(binding_text)
            ]
            if missing_tokens or not binding_shapes:
                content_binding_violations.append(
                    {
                        "slide_id": specification["slide_id"],
                        "binding_type": binding["binding_type"],
                        "binding_id": binding["binding_id"],
                        "missing_tokens": missing_tokens,
                        "reason": "missing_visible_binding" if not binding_shapes else "binding_token_mismatch",
                    }
                )
        for chart_lock in specification.get("chart_locks", []):
            chart_id = chart_lock["chart_id"]
            expected_shape_name = f"CONTENT_chart_{chart_id}"
            native = next(
                (
                    shape
                    for shape in slide.shapes
                    if getattr(shape, "has_chart", False)
                    and str(getattr(shape, "name", "")).startswith(expected_shape_name)
                ),
                None,
            )
            evidence = next(
                (
                    shape
                    for shape in slide.shapes
                    if str(getattr(shape, "name", ""))
                    == f"{expected_shape_name}_VISROLE_chart-evidence"
                    and not (getattr(shape, "has_text_frame", False) and (shape.text or "").strip())
                ),
                None,
            )
            metadata = next(
                (
                    shape
                    for shape in slide.shapes
                    if str(getattr(shape, "name", ""))
                    == f"CONTENT_chart_metadata_{chart_id}"
                    and getattr(shape, "has_text_frame", False)
                ),
                None,
            )
            if native is not None:
                if not _same_chart_payload(_chart_payload_from_native_shape(native), chart_lock):
                    chart_lock_violations.append(
                        {
                            "slide_id": specification["slide_id"],
                            "chart_id": chart_id,
                            "reason": "native_chart_payload_mismatch",
                        }
                    )
            elif evidence is not None and metadata is not None:
                try:
                    actual_metadata = json.loads(metadata.text)
                except json.JSONDecodeError:
                    actual_metadata = {}
                if actual_metadata != chart_lock:
                    chart_lock_violations.append(
                        {
                            "slide_id": specification["slide_id"],
                            "chart_id": chart_id,
                            "reason": "chart_metadata_mismatch",
                        }
                    )
            else:
                chart_lock_violations.append(
                    {
                        "slide_id": specification["slide_id"],
                        "chart_id": chart_id,
                        "reason": "missing_structured_chart_evidence",
                    }
                )
            if expected_shape_name not in shape_names and _normalize(chart_id) not in normalized_text:
                chart_lock_warnings.append(
                    {
                        "slide_id": specification["slide_id"],
                        "chart_id": chart_id,
                        "reason": "chart_identity_not_machine_readable",
                    }
                )

    passed = (
        slide_count_matches
        and not missing_locked_content
        and not unresolved_placeholders
        and not slide_id_metadata_issues
        and not semantic_lock_violations
        and not chart_lock_violations
        and not content_binding_violations
    )
    return {
        "status": "passed" if passed else "failed",
        "expected_slide_count": len(lock_spec),
        "actual_slide_count": len(presentation.slides),
        "slide_count_matches": slide_count_matches,
        "missing_locked_content": missing_locked_content,
        "unresolved_placeholders": unresolved_placeholders,
        "chart_lock_warnings": chart_lock_warnings,
        "slide_id_metadata_issues": slide_id_metadata_issues,
        "semantic_lock_violations": semantic_lock_violations,
        "chart_lock_violations": chart_lock_violations,
        "content_binding_violations": content_binding_violations,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Verify a SlideViber-polished PPTX against report_model v2 locks."
    )
    parser.add_argument("--model", required=True, type=Path)
    parser.add_argument("--pptx", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    model = json.loads(args.model.read_text(encoding="utf-8"))
    normalized_count = normalize_east_asian_text(args.pptx)
    result = qa_slideviber_pptx(args.pptx, model)
    result["east_asian_text_normalized"] = normalized_count
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
