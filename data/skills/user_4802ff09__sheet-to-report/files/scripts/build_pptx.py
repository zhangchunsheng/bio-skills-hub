from __future__ import annotations

import argparse
import json
import re
import zlib
from pathlib import Path
from typing import Any, Callable

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from layout_registry import FONT_MINIMUMS, validate_layout_capacity
from model_validation import validate_report_model
from content_resolver import project_legacy_content_refs, resolve_model_for_page
from presentation_contracts import (
    compact_report_subtitle as _compact_report_subtitle,
    compact_comparison_copy,
    comparison_context_text,
    validate_image_exception,
    validate_theme,
    validate_visual_spec,
)


SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5
FONT_NAME = "Microsoft YaHei"
HEADING_TERMINAL_PUNCTUATION = "，。；：！？、"


def _rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def load_theme(name: str) -> dict[str, Any]:
    path = Path(__file__).resolve().parents[1] / "references" / "themes.json"
    themes = json.loads(path.read_text(encoding="utf-8"))
    if name not in themes:
        raise ValueError(f"unknown theme: {name}")
    theme = themes[name]
    validate_theme(theme)
    return theme


def _format_metric_value(value: Any, unit: str) -> str:
    if value is None:
        return "—"
    number = float(value)
    if unit == "%":
        return f"{number:,.2f}%"
    if unit == "x":
        return f"{number:,.2f}倍"
    if abs(number) >= 10_000:
        return f"{number / 10_000:,.2f}万"
    if unit == "单" or abs(number) >= 1_000:
        return f"{number:,.0f}"
    return f"{number:,.2f}"


def _comparison_text(item: dict[str, Any], basis: str) -> str:
    change = item.get("change_pct")
    if change is None:
        return f"暂无{basis}可比数据"
    if str(item.get("unit", "")) == "%" and item.get("change_delta") is not None:
        return f"较{basis} {float(item['change_delta']):+.2f} 个百分点"
    return f"较{basis} {float(change):+.2f}%"


def _localized_period(value: Any) -> str:
    return {"monthly": "月度", "weekly": "周度", "custom": "自定义周期"}.get(
        str(value), str(value)
    )


def _localized_incomplete_policy(value: Any) -> str:
    return {
        "exclude": "排除并提示",
        "keep": "保留并提示",
        "include": "保留并提示",
    }.get(str(value), str(value))


def _compact_ppt_heading(value: Any) -> str:
    """Keep title-style card copy free of orphan-prone terminal punctuation."""
    return str(value or "").rstrip().rstrip(HEADING_TERMINAL_PUNCTUATION)


def _set_background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _style_run(run, size: float, color: str, bold: bool) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    properties = run._r.get_or_add_rPr()
    properties.set("lang", "zh-CN")
    properties.set("altLang", "en-US")
    for element in properties.findall(qn("a:ea")):
        properties.remove(element)
    properties.append(properties.makeelement(qn("a:ea"), {"typeface": FONT_NAME}))


def _add_text(
    slide,
    text: str | list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    role: str,
    size: float,
    color: str,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    vertical=MSO_ANCHOR.TOP,
    margin: float = 0.04,
    visual_role: str | None = None,
) -> Any:
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    box.name = f"ROLE_{role}_{len(slide.shapes):03d}"
    if visual_role:
        box.name = f"{box.name}_VISROLE_{visual_role}"
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = vertical
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    lines = [text] if isinstance(text, str) else list(text)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(4 if role == "body" else 0)
        paragraph_properties = paragraph._p.get_or_add_pPr()
        paragraph_properties.set("eaLnBrk", "1")
        paragraph_properties.set("latinLnBrk", "0")
        paragraph_properties.set("hangingPunct", "1")
        run = paragraph.add_run()
        run.text = str(line)
        _style_run(run, size, color, bold)
    return box


def _add_aligned_row_text(
    slide,
    text: str | list[str],
    left: float,
    width: float,
    *,
    row_top: float,
    row_height: float,
    group: str,
    part: str,
    role: str,
    size: float,
    color: str,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
) -> Any:
    """Add one member of a semantic row under a shared centerline contract."""
    shape = _add_text(
        slide,
        text,
        left,
        row_top,
        width,
        row_height,
        role=role,
        size=size,
        color=color,
        bold=bold,
        align=align,
        vertical=MSO_ANCHOR.MIDDLE,
        margin=0,
    )
    for paragraph in shape.text_frame.paragraphs:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
    shape.name = f"{shape.name}_ALIGNROW_{group}_{part}"
    return shape


def _add_rule(
    slide,
    left: float,
    top: float,
    width: float,
    color: str,
    *,
    visual_role: str | None = None,
) -> Any:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(0.025),
    )
    shape.name = f"DECOR_rule_{len(slide.shapes):03d}"
    if visual_role:
        shape.name = f"{shape.name}_VISROLE_{visual_role}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    return shape


def _add_panel(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    fill_color: str,
    *,
    visual_role: str | None = None,
) -> Any:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.name = f"DECOR_panel_{len(slide.shapes):03d}"
    if visual_role:
        shape.name = f"{shape.name}_VISROLE_{visual_role}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_color)
    shape.line.fill.background()
    return shape


def _add_circle(
    slide,
    left: float,
    top: float,
    diameter: float,
    fill_color: str,
    *,
    visual_role: str | None = None,
) -> Any:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(left),
        Inches(top),
        Inches(diameter),
        Inches(diameter),
    )
    shape.name = f"DECOR_circle_{len(slide.shapes):03d}"
    if visual_role:
        shape.name = f"{shape.name}_VISROLE_{visual_role}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_color)
    shape.line.fill.background()
    return shape


def _add_slide_title(
    slide,
    item: dict[str, Any],
    theme: dict[str, Any],
    *,
    force_compact: bool = False,
) -> None:
    visual_spec = item.get("visual_spec", {})
    compact_density = visual_spec.get("density") == "compact" or force_compact
    _add_text(
        slide,
        visual_spec.get("eyebrow", "经营分析"),
        0.76,
        0.22,
        3.2,
        0.24,
        role="meta",
        size=11,
        color=theme["accent"],
        bold=True,
        visual_role="eyebrow",
    )
    context_text = comparison_context_text(visual_spec.get("comparison_contexts"))
    if context_text:
        _add_text(
            slide,
            context_text,
            5.42,
            0.22,
            6.78,
            0.24,
            role="meta",
            size=10,
            color=theme["muted"],
            align=PP_ALIGN.RIGHT,
            visual_role="comparison-context",
        )
    title_shape = _add_text(
        slide,
        visual_spec.get("display_claim") or item["claim"],
        0.60 if compact_density else 0.72,
        0.48,
        12.12 if compact_density else 11.9,
        1.32 if compact_density else 0.68,
        role="slide_title",
        size=FONT_MINIMUMS["slide_title"],
        color=theme["ink"],
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
        visual_role="title",
    )
    if compact_density:
        title_shape.name = f"{title_shape.name}_DENSITY_compact"


def _apply_compact_title_fallback(slide: Any, item: dict[str, Any]) -> None:
    """Reserve real vertical space for a two-line compact title before QA."""

    if item.get("visual_spec", {}).get("density") != "compact":
        return
    title = next(
        (shape for shape in slide.shapes if "DENSITY_compact" in shape.name),
        None,
    )
    if title is None:
        return
    offset = Inches(0.64)
    content_start = Inches(1.16)
    for shape in slide.shapes:
        if shape is title or shape.top < content_start:
            continue
        if shape.top + shape.height + offset <= Inches(SLIDE_HEIGHT):
            shape.top += offset


def _display_copy(item: dict[str, Any], text: Any) -> str:
    contexts = item.get("visual_spec", {}).get("comparison_contexts", [])
    return compact_comparison_copy(text, contexts)


def _chart_metadata_line(item: dict[str, Any], chart: dict[str, Any]) -> str:
    scope = str(chart.get("scope_label") or "")
    for context in item.get("visual_spec", {}).get("comparison_contexts", []):
        current = str(context.get("current_label") or "")
        baseline = str(context.get("baseline_label") or "")
        if current and baseline and current in scope and baseline in scope:
            scope = ""
            break
    unit = str(chart.get("display_unit") or chart.get("unit") or "")
    return "｜".join(
        value for value in (scope, f"单位：{unit}" if unit else "") if value
    )


def _footer_text(item: dict[str, Any]) -> str:
    return str(item["slide_id"])


def _add_footer(slide, item: dict[str, Any], theme: dict[str, Any]) -> None:
    _add_text(
        slide,
        _footer_text(item),
        0.72,
        7.08,
        11.9,
        0.22,
        role="footnote",
        size=FONT_MINIMUMS["footnote"],
        color=theme["muted"],
    )


def _chart_by_id(model: dict[str, Any], chart_id: str) -> dict[str, Any]:
    for chart in model.get("charts", []):
        if chart.get("chart_id") == chart_id:
            return chart
    raise ValueError(f"slide_plan 引用了不存在的 chart：{chart_id}")


def _normalize_chart_axis_ids(chart: Any, chart_id: str) -> None:
    """Keep native chart axis references valid for strict UInt32 OOXML readers."""
    axis_nodes = chart._chartSpace.xpath(".//c:axId")
    cross_nodes = chart._chartSpace.xpath(".//c:crossAx")
    replacements: dict[str, str] = {}
    for node in [*axis_nodes, *cross_nodes]:
        original = str(node.get("val"))
        if original not in replacements:
            normalized = zlib.crc32(f"{chart_id}:{original}".encode("utf-8"))
            replacements[original] = str(normalized or 1)
        node.set("val", replacements[original])


def _add_native_chart(
    slide,
    chart_spec: dict[str, Any],
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict[str, Any],
    chart_focus: dict[str, Any] | None = None,
) -> Any:
    data = ChartData()
    data.categories = chart_spec.get("categories", [])
    display_scale = float(chart_spec.get("display_scale", 1) or 1)
    for series in chart_spec.get("series", []):
        data.add_series(
            str(series["name"]),
            [float(value) / display_scale for value in series.get("values", [])],
        )
    chart_type = (
        XL_CHART_TYPE.BAR_CLUSTERED
        if chart_spec.get("chart_type") == "bar"
        else XL_CHART_TYPE.LINE_MARKERS
    )
    shape = slide.shapes.add_chart(
        chart_type,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
        data,
    )
    shape.name = f"CONTENT_chart_{chart_spec['chart_id']}_VISROLE_native-chart"
    chart = shape.chart
    _normalize_chart_axis_ids(chart, str(chart_spec["chart_id"]))
    chart.has_legend = len(chart.series) > 1
    chart.has_title = False
    chart.chart_style = 2
    chart.font.name = FONT_NAME
    chart.font.size = Pt(12)
    chart.category_axis.tick_labels.font.name = FONT_NAME
    chart.category_axis.tick_labels.font.size = Pt(12)
    chart.value_axis.tick_labels.font.name = FONT_NAME
    chart.value_axis.tick_labels.font.size = Pt(12)
    chart.value_axis.tick_labels.number_format = str(
        chart_spec.get("number_format", "#,##0")
    )
    chart.value_axis.tick_labels.number_format_is_linked = False
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = str(
        chart_spec.get("display_unit") or chart_spec.get("unit") or "数值"
    )
    for paragraph in chart.value_axis.axis_title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(12)
    chart.value_axis.has_major_gridlines = True
    try:
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(theme["line"])
        chart.value_axis.major_gridlines.format.line.width = Pt(0.75)
        chart.category_axis.format.line.color.rgb = _rgb(theme["line"])
        chart.value_axis.format.line.color.rgb = _rgb(theme["line"])
    except (AttributeError, ValueError):
        pass
    if chart.has_legend:
        chart.legend.font.name = FONT_NAME
        chart.legend.font.size = Pt(12)
    focus_series = str((chart_focus or {}).get("series_name") or "")
    focus_category = str((chart_focus or {}).get("category_name") or "")
    categories = [str(value) for value in chart_spec.get("categories", [])]
    for index, series in enumerate(chart.series):
        source_name = str(chart_spec.get("series", [])[index].get("name", ""))
        is_focus = not focus_series or source_name == focus_series
        color = _rgb(
            theme["chart"][index % len(theme["chart"])]
            if is_focus
            else theme["chart_muted"]
        )
        if chart_type == XL_CHART_TYPE.BAR_CLUSTERED:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
            series.format.line.fill.background()
            if focus_category and focus_category in categories:
                focus_index = categories.index(focus_category)
                for point_index, point in enumerate(series.points):
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = _rgb(
                        theme["accent"]
                        if point_index == focus_index
                        else theme["chart_muted"]
                    )
                    point.format.line.fill.background()
        else:
            series.format.line.color.rgb = color
            series.format.line.width = Pt(2.5)
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.color.rgb = color
    return shape


def _chart_focus_value(
    chart_spec: dict[str, Any], chart_focus: dict[str, Any] | None
) -> str:
    series_specs = list(chart_spec.get("series", []))
    if not series_specs:
        return "—"
    focus_series = str((chart_focus or {}).get("series_name") or "")
    series = next(
        (item for item in series_specs if str(item.get("name")) == focus_series),
        series_specs[0],
    )
    values = list(series.get("values", []))
    if not values:
        return "—"
    focus_category = str((chart_focus or {}).get("category_name") or "")
    categories = [str(value) for value in chart_spec.get("categories", [])]
    index = categories.index(focus_category) if focus_category in categories else len(values) - 1
    value = float(values[index]) / float(chart_spec.get("display_scale", 1) or 1)
    unit = str(chart_spec.get("display_unit") or chart_spec.get("unit") or "")
    if unit in {"%", "百分点"}:
        return f"{value:,.2f}{unit}"
    return f"{value:,.1f}{unit}" if abs(value) < 1000 else f"{value:,.0f}{unit}"


def _kpi_role(kpi: dict[str, Any], model: dict[str, Any]) -> str:
    if kpi.get("role"):
        return str(kpi["role"])
    contract = next(
        (
            value
            for value in model.get("metric_contracts", [])
            if value.get("name") == kpi.get("name")
        ),
        {},
    )
    return str(contract.get("role") or "unknown")


def _rank_kpis(kpis: list[dict[str, Any]], model: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not kpis:
        return [], []
    outcome = next(
        (kpi for kpi in kpis if _kpi_role(kpi, model) in {"result", "outcome"}),
        kpis[0],
    )
    primary = [outcome]
    if len(kpis) > 1:
        top_evidence = set(
            model.get("insights", [{}])[0].get("evidence_ids", [])
            if model.get("insights")
            else []
        )
        preferred_roles = {"quality", "efficiency"}
        second = next(
            (
                kpi
                for kpi in kpis
                if kpi is not outcome
                and _kpi_role(kpi, model) in preferred_roles
                and (
                    not top_evidence
                    or top_evidence & set(kpi.get("evidence_ids", []))
                )
            ),
            next((kpi for kpi in kpis if kpi is not outcome), None),
        )
        if second is not None:
            primary.append(second)
    supporting = [kpi for kpi in kpis if kpi not in primary]
    return primary, supporting


def _render_cover(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["cover_background"])
    halo = _add_circle(slide, 10.4, -0.7, 3.6, theme["accent"])
    halo.name = f"{halo.name}_DECORATIVE_OVERFLOW"
    echo = _add_circle(slide, 11.55, 5.65, 2.2, theme["secondary"])
    echo.name = f"{echo.name}_DECORATIVE_OVERFLOW"
    _add_text(
        slide,
        item.get("visual_spec", {}).get("eyebrow", "经营复盘"),
        0.86,
        0.58,
        3.2,
        0.28,
        role="meta",
        size=12,
        color=theme["cover_muted"],
        bold=True,
        visual_role="eyebrow",
    )
    _add_text(
        slide,
        item["claim"],
        0.82,
        1.02,
        9.5,
        1.20,
        role="cover_title",
        size=54,
        color=theme["cover_ink"],
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
        visual_role="title",
    )
    scope = _add_text(
        slide,
        _compact_report_subtitle(model.get("report_subtitle", "")),
        0.82,
        2.34,
        8.8,
        0.5,
        role="subheading",
        size=FONT_MINIMUMS["subheading"],
        color=theme["cover_muted"],
        visual_role="scope",
    )
    scope.name = f"{scope.name}_PUNCTGUARD"
    _add_rule(slide, 0.86, 3.24, 1.35, theme["secondary"])
    _add_text(
        slide,
        "核心判断",
        0.82,
        3.55,
        2.0,
        0.42,
        role="meta",
        size=14,
        color=theme["cover_muted"],
        bold=True,
    )
    _add_text(
        slide,
        model.get("judgement_headline", model["request"]["objective"]),
        0.82,
        4.08,
        10.55,
        1.12,
        role="subheading",
        size=26,
        color=theme["cover_ink"],
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
        visual_role="judgement",
    )
    _add_text(
        slide,
        f"汇报对象：{model['request']['audience']}",
        0.82,
        6.28,
        9.2,
        0.34,
        role="meta",
        size=14,
        color=theme["cover_muted"],
    )
    _add_text(
        slide,
        item["slide_id"],
        11.65,
        7.08,
        0.75,
        0.2,
        role="footnote",
        size=FONT_MINIMUMS["footnote"],
        color=theme["cover_muted"],
        align=PP_ALIGN.RIGHT,
    )


def _render_summary(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])
    _add_slide_title(slide, item, theme)
    conclusions = list(model.get("executive_summary", {}).get("key_conclusions", []))
    row_colors = (theme["secondary"], theme["negative"], theme["accent"])
    for index, conclusion in enumerate(conclusions):
        top = 1.42 + index * 1.22
        group = f"{item['slide_id']}-summary-{index + 1:02d}"
        _add_aligned_row_text(
            slide,
            f"0{index + 1}",
            0.82,
            0.62,
            row_top=top,
            row_height=0.82,
            group=group,
            part="number",
            role="subheading",
            size=30,
            color=row_colors[index],
            bold=True,
        )
        _add_rule(slide, 1.54, top + 0.08, 0.035, row_colors[index])
        body = _add_aligned_row_text(
            slide,
            _display_copy(item, conclusion),
            1.82,
            10.18,
            row_top=top,
            row_height=0.82,
            group=group,
            part="body",
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["ink"],
            bold=True,
        )
        body.name = body.name.replace(
            "_ALIGNROW_", "_VISROLE_numbered-judgement_ALIGNROW_"
        )
        if index < len(conclusions) - 1:
            _add_rule(slide, 1.82, top + 1.00, 10.18, theme["line"])
    priority_actions = list(model.get("executive_summary", {}).get("priority_actions", []))
    if priority_actions:
        _add_panel(
            slide,
            0.82,
            5.18,
            11.55,
            1.12,
            theme["cover_background"],
            visual_role="action-band",
        )
        _add_aligned_row_text(
            slide,
            "下一步重点",
            1.05,
            1.62,
            row_top=5.18,
            row_height=1.12,
            group=f"{item['slide_id']}-actions",
            part="label",
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["cover_muted"],
            bold=True,
        )
        _add_aligned_row_text(
            slide,
            [f"• {action}" for action in priority_actions],
            2.72,
            9.2,
            row_top=5.18,
            row_height=1.12,
            group=f"{item['slide_id']}-actions",
            part="body",
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["cover_ink"],
        )
    _add_footer(slide, item, theme)


def _render_kpi(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])
    _add_slide_title(slide, item, theme)
    period_overview = model.get("period_overview", {})
    kpis = list(period_overview.get("kpis") or model["kpis"])
    primary, supporting = _rank_kpis(kpis, model)
    _add_text(
        slide,
        f"分析期：{period_overview.get('label', '—')}（累计／整体口径）",
        8.30,
        1.22,
        3.90,
        0.3,
        role="meta",
        size=12,
        color=theme["muted"],
        align=PP_ALIGN.RIGHT,
    )
    primary_gap = 0.28
    primary_width = (
        (11.55 - primary_gap * (len(primary) - 1)) / len(primary)
        if primary
        else 11.55
    )
    for index, kpi in enumerate(primary):
        left = 0.82 + index * (primary_width + primary_gap)
        fill = theme["accent_soft"] if index == 0 else theme["secondary_soft"]
        card = _add_panel(
            slide,
            left,
            1.48,
            primary_width,
            2.05 if supporting else 3.65,
            fill,
            visual_role=f"primary-kpi_{kpi.get('display_name') or kpi['name']}",
        )
        card.name = card.name.replace(" ", "_")
        label_shape = _add_text(
            slide,
            kpi.get("display_name") or kpi["name"],
            left + 0.26,
            1.74,
            primary_width - 0.52,
            0.34,
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["muted"],
        )
        binding_id = str(kpi.get("kpi_id") or kpi["name"])
        label_shape.name = f"CONTENT_binding_metric_{binding_id}_label"
        value = _format_metric_value(kpi.get("value"), str(kpi.get("unit", "")))
        value_shape = _add_text(
            slide,
            value,
            left + 0.24,
            2.15,
            primary_width - 0.48,
            0.72,
            role="subheading",
            size=38,
            color=theme["ink"],
            bold=True,
        )
        value_shape.name = f"CONTENT_binding_metric_{binding_id}_value_KPI_VALUE_PRIMARY"
        role_label = _kpi_role(kpi, model)
        _add_text(
            slide,
            "核心结果" if role_label in {"result", "outcome"} else "质量／效率焦点",
            left + 0.26,
            3.02,
            primary_width - 0.52,
            0.28,
            role="meta",
            size=11,
            color=theme["accent"] if index == 0 else theme["secondary"],
            bold=True,
        )
    support_count = len(supporting)
    if support_count:
        gap = 0.20
        width = (11.55 - gap * (support_count - 1)) / support_count
        for index, kpi in enumerate(supporting):
            left = 0.82 + index * (width + gap)
            card = _add_panel(
                slide,
                left,
                3.88,
                width,
                1.72,
                theme["surface"],
                visual_role=f"supporting-kpi_{kpi.get('display_name') or kpi['name']}",
            )
            card.name = card.name.replace(" ", "_")
            _add_rule(slide, left, 3.88, width, theme["accent"])
            label_shape = _add_text(
                slide,
                kpi.get("display_name") or kpi["name"],
                left + 0.18,
                4.18,
                width - 0.36,
                0.30,
                role="body",
                size=FONT_MINIMUMS["body"],
                color=theme["muted"],
            )
            binding_id = str(kpi.get("kpi_id") or kpi["name"])
            label_shape.name = f"CONTENT_binding_metric_{binding_id}_label"
            value = _format_metric_value(kpi.get("value"), str(kpi.get("unit", "")))
            value_shape = _add_text(
                slide,
                value,
                left + 0.18,
                4.62,
                width - 0.36,
                0.58,
                role="subheading",
                size=26,
                color=theme["ink"],
                bold=True,
            )
            value_shape.name = f"CONTENT_binding_metric_{binding_id}_value_KPI_VALUE_SUPPORTING"
    _add_footer(slide, item, theme)


def _render_trend(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])
    _add_slide_title(slide, item, theme)
    chart = _chart_by_id(model, item["chart_ids"][0])
    chart_focus = item.get("visual_spec", {}).get("chart_focus")
    _add_text(
        slide,
        chart.get("title", "关键趋势"),
        0.82,
        1.22,
        6.5,
        0.30,
        role="subheading",
        size=FONT_MINIMUMS["subheading"],
        color=theme["ink"],
        bold=True,
    )
    _add_text(
        slide,
        _chart_metadata_line(item, chart),
        7.50,
        1.22,
        4.70,
        0.28,
        role="meta",
        size=11,
        color=theme["muted"],
        align=PP_ALIGN.RIGHT,
    )
    _add_native_chart(slide, chart, 0.72, 1.60, 11.9, 3.74, theme, chart_focus)
    insight = next(
        (
            value for value in model.get("insights", [])
            if set(value.get("evidence_ids", [])) & set(item.get("evidence_ids", []))
        ),
        None,
    )
    if insight is None:
        raise ValueError(f"趋势页 {item.get('slide_id')} 缺少本页 insight 引用")
    _add_panel(
        slide,
        0.82,
        5.56,
        11.55,
        0.86,
        theme["accent_soft"],
        visual_role="focus-callout",
    )
    _add_text(
        slide,
        _chart_focus_value(chart, chart_focus),
        1.05,
        5.70,
        2.08,
        0.54,
        role="subheading",
        size=26,
        color=theme["accent"],
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
    )
    _add_aligned_row_text(
        slide,
        "核心判断",
        3.25,
        1.35,
        row_top=5.56,
        row_height=0.86,
        group=f"{item['slide_id']}-core",
        part="label",
        role="subheading",
        size=FONT_MINIMUMS["subheading"],
        color=theme["accent"],
        bold=True,
    )
    _add_aligned_row_text(
        slide,
        _display_copy(item, insight["statement"]),
        4.72,
        7.30,
        row_top=5.56,
        row_height=0.86,
        group=f"{item['slide_id']}-core",
        part="body",
        role="body",
        size=FONT_MINIMUMS["body"],
        color=theme["ink"],
        bold=True,
    )
    _add_footer(slide, item, theme)


def _chart_side_content(item: dict, model: dict, chart: dict) -> dict[str, str]:
    """Resolve a chart callout only from this page's selected evidence."""

    selected_ids = list(dict.fromkeys(
        list(item.get("claim_evidence_ids", []))
        + list(item.get("evidence_ids", []))
        + list(chart.get("evidence_ids", []))
    ))
    evidence_items = [
        model.get("evidence_index", {}).get(evidence_id, {})
        for evidence_id in selected_ids
    ]
    evidence_items = [item for item in evidence_items if item]
    matched_insight = next(
        (
            insight for insight in model.get("insights", [])
            if set(insight.get("evidence_ids", [])) & set(selected_ids)
        ),
        None,
    )
    note = str(
        (matched_insight or {}).get("implication")
        or chart.get("takeaway")
        or item.get("claim")
        or "本页图表用于呈现已选证据。"
    )
    diagnostic = next(
        (evidence for evidence in evidence_items if evidence.get("kind") == "dimension_yoy"),
        None,
    )
    if diagnostic:
        unit = str(diagnostic.get("unit", ""))
        latest = float(diagnostic.get("latest") or diagnostic.get("metric") or 0)
        return {
            "heading": "重点诊断项",
            "value": str(diagnostic.get("value") or "—"),
            "metric_line": f"{diagnostic.get('metric', '指标')} {latest:,.2f}{unit}",
            "note": note,
        }

    contribution = next(
        (
            evidence for evidence in evidence_items
            if str(evidence.get("kind", "")).startswith("dimension")
            and evidence.get("rows")
        ),
        None,
    )
    focus = item.get("visual_spec", {}).get("chart_focus") or {}
    focus_category = str(focus.get("category_name") or "")
    if contribution:
        rows = [row for row in contribution.get("rows", []) if not row.get("is_other")]
        selected_row = next(
            (row for row in rows if str(row.get("value")) == focus_category),
            rows[0] if rows else {},
        )
        metric_value = float(selected_row.get("metric") or 0)
        unit = str(contribution.get("unit", ""))
        dimension = str(contribution.get("dimension") or "结构")
        metric_name = str(contribution.get("metric") or "贡献值")
        return {
            "heading": f"{dimension}领先项",
            "value": str(selected_row.get("value") or "—"),
            "metric_line": f"{metric_name} {metric_value:,.2f}{unit}",
            "note": note,
        }

    categories = [str(value) for value in chart.get("categories", [])]
    series = list(chart.get("series", []))
    selected_series = next(
        (value for value in series if str(value.get("name")) == str(focus.get("series_name") or "")),
        series[0] if series else {},
    )
    values = list(selected_series.get("values", []))
    index = categories.index(focus_category) if focus_category in categories else max(0, len(values) - 1)
    value = float(values[index]) if values and index < len(values) else 0.0
    category = categories[index] if categories and index < len(categories) else "—"
    unit = str(chart.get("display_unit") or chart.get("unit") or "")
    return {
        "heading": "图表焦点",
        "value": category,
        "metric_line": f"{selected_series.get('name', '数值')} {value:,.2f}{unit}",
        "note": note,
    }


def _render_chart_side(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])
    _add_slide_title(slide, item, theme)
    chart = _chart_by_id(model, item["chart_ids"][0])
    chart_focus = item.get("visual_spec", {}).get("chart_focus")
    _add_native_chart(slide, chart, 0.72, 1.5, 8.4, 4.72, theme, chart_focus)
    _add_panel(
        slide,
        9.38,
        1.5,
        3.22,
        4.72,
        theme["accent_soft"],
        visual_role="focus-callout",
    )
    sidebar = _chart_side_content(item, model, chart)
    sidebar_heading = sidebar["heading"]
    metric_line = sidebar["metric_line"]
    sidebar_note = sidebar["note"]
    _add_text(
        slide,
        sidebar_heading,
        9.66,
        1.78,
        2.66,
        0.4,
        role="subheading",
        size=FONT_MINIMUMS["subheading"],
        color=theme["accent"],
        bold=True,
    )
    _add_text(
        slide,
        sidebar["value"],
        9.66,
        2.42,
        2.66,
        0.62,
        role="subheading",
        size=30,
        color=theme["ink"],
        bold=True,
    )
    _add_text(
        slide,
        metric_line,
        9.66,
        3.08,
        2.66,
        0.36,
        role="body",
        size=FONT_MINIMUMS["body"],
        color=theme["muted"],
    )
    _add_text(
        slide,
        sidebar_note,
        9.66,
        3.78,
        2.54,
        1.45,
        role="body",
        size=FONT_MINIMUMS["body"],
        color=theme["ink"],
    )
    _add_footer(slide, item, theme)


def _render_levers(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])
    _add_slide_title(slide, item, theme)
    levers = model.get("levers", [])
    visible_levers = levers
    if visible_levers:
        _add_rule(
            slide,
            1.18,
            3.16,
            10.92,
            theme["line"],
            visual_role="lever-axis",
        )
    palette = (theme["accent"], theme["secondary"], theme["positive"], theme["negative"])
    count = len(visible_levers)
    for index, lever in enumerate(visible_levers):
        center = 1.45 + (10.35 * index / max(count - 1, 1))
        left = max(0.72, min(center - 1.28, 10.04))
        _add_text(
            slide,
            lever["label"],
            left,
            1.86,
            2.56,
            0.42,
            role="subheading",
            size=FONT_MINIMUMS["subheading"],
            color=theme["ink"],
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        node = _add_circle(
            slide,
            center - 0.28,
            2.88,
            0.56,
            palette[index],
            visual_role="lever-node",
        )
        node.name = f"{node.name}_{index + 1:02d}"
        _add_text(
            slide,
            f"{index + 1:02d}",
            center - 0.28,
            2.88,
            0.56,
            0.56,
            role="body",
            size=FONT_MINIMUMS["body"],
            color="FFFFFF",
            bold=True,
            align=PP_ALIGN.CENTER,
            vertical=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
        _add_text(
            slide,
            _display_copy(item, lever["headline"]),
            left,
            3.68,
            2.56,
            0.82,
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["ink"],
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        _add_text(
            slide,
            _display_copy(
                item,
                lever.get("takeaway") or lever.get("headline") or "",
            ),
            left,
            4.68,
            2.56,
            1.02,
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["muted"],
            align=PP_ALIGN.CENTER,
        )
    _add_footer(slide, item, theme)


def _render_risk(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])
    _add_slide_title(slide, item, theme)
    _add_panel(
        slide,
        0.78,
        1.46,
        5.92,
        4.86,
        theme["positive_soft"],
        visual_role="opportunity-column",
    )
    _add_panel(
        slide,
        6.92,
        1.46,
        5.64,
        4.86,
        theme["negative_soft"],
        visual_role="risk-column",
    )
    _add_panel(slide, 0.78, 1.46, 5.92, 0.68, theme["positive"])
    _add_panel(slide, 6.92, 1.46, 5.64, 0.68, theme["negative"])
    _add_text(
        slide,
        "机会｜先验证质量，再决定扩张",
        1.08,
        1.58,
        5.26,
        0.42,
        role="meta",
        size=14,
        color="FFFFFF",
        bold=True,
    )
    opportunity_lines = [insight["statement"] for insight in model.get("opportunity_insights", [])]
    risk_lines = [insight["statement"] for insight in model.get("risk_insights", [])]
    # A bounded empty-state preserves the two-column decision frame without
    # leaving an unexplained blank visual slot when one signal class is absent.
    if not opportunity_lines:
        opportunity_lines = [
            "暂无新增机会信号；维持当前验证节奏。"
            if model.get("opportunity_globally_absent", False)
            else "本页未引用机会信号；详见机会判断页。"
        ]
    if not risk_lines:
        risk_lines = [
            "暂无新增风险信号；持续监控关键指标。"
            if model.get("risk_globally_absent", False)
            else "本页未引用风险信号；详见风险判断页。"
        ]
    for index in range(len(opportunity_lines)):
        top = 2.56 + index * 1.45
        _add_circle(slide, 1.04, top, 0.44, theme["positive"])
        _add_text(
            slide,
            f"{index + 1:02d}",
            1.04,
            top,
            0.44,
            0.44,
            role="meta",
            size=10,
            color="FFFFFF",
            bold=True,
            align=PP_ALIGN.CENTER,
            vertical=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    opportunity_box = _add_text(
        slide,
        [f"• {_display_copy(item, line)}" for line in opportunity_lines],
        1.58,
        2.42,
        4.70,
        3.10,
        role="body",
        size=FONT_MINIMUMS["body"],
        color=theme["ink"],
    )
    for paragraph in opportunity_box.text_frame.paragraphs:
        paragraph.space_after = Pt(8)
    _add_text(
        slide,
        "风险｜先止损，再讨论恢复",
        7.24,
        1.58,
        5.00,
        0.42,
        role="meta",
        size=14,
        color="FFFFFF",
        bold=True,
    )
    for index in range(len(risk_lines)):
        top = 2.56 + index * 1.45
        _add_circle(slide, 7.18, top, 0.44, theme["negative"])
        _add_text(
            slide,
            f"{index + 1:02d}",
            7.18,
            top,
            0.44,
            0.44,
            role="meta",
            size=10,
            color="FFFFFF",
            bold=True,
            align=PP_ALIGN.CENTER,
            vertical=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    risk_box = _add_text(
        slide,
        [f"• {_display_copy(item, line)}" for line in risk_lines],
        7.72,
        2.42,
        4.34,
        3.10,
        role="body",
        size=FONT_MINIMUMS["body"],
        color=theme["ink"],
    )
    for paragraph in risk_box.text_frame.paragraphs:
        paragraph.space_after = Pt(8)
    _add_footer(slide, item, theme)


def _render_actions(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])
    title_text = str(item.get("visual_spec", {}).get("display_claim") or item["claim"])
    compact_title = len(title_text) > 32
    _add_slide_title(slide, item, theme, force_compact=compact_title)
    if item.get("content_exemption") == "selection_meta":
        _add_panel(slide, 0.82, 1.70, 11.55, 2.50, theme["accent_soft"], visual_role="selection-meta")
        _add_text(slide, item["claim"], 1.16, 2.05, 10.80, 0.60, role="subheading", size=FONT_MINIMUMS["subheading"], color=theme["accent"], bold=True)
        _add_text(slide, item["purpose"], 1.16, 2.82, 10.80, 0.90, role="body", size=FONT_MINIMUMS["body"], color=theme["ink"])
        _add_footer(slide, item, theme)
        return
    actions = model["actions"]
    row_colors = (theme["accent"], theme["secondary"], theme["negative"])
    for index, action in enumerate(actions):
        top = (1.92 + index * 1.55) if compact_title else (1.42 + index * 1.66)
        _add_panel(
            slide,
            0.82,
            top,
            11.55,
            1.40,
            theme["surface"],
            visual_role="priority-row",
        )
        _add_panel(slide, 0.82, top, 1.18, 1.40, row_colors[index])
        _add_text(
            slide,
            f"优先级\nP{index + 1}",
            0.82,
            top,
            1.18,
            1.40,
            role="body",
            size=FONT_MINIMUMS["body"],
            color="FFFFFF",
            bold=True,
            align=PP_ALIGN.CENTER,
            vertical=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
        action_heading = _add_text(
            slide,
            _compact_ppt_heading(action.get("headline") or action["text"]),
            2.28,
            top + 0.22,
            3.70,
            1.00,
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["ink"],
            bold=True,
            vertical=MSO_ANCHOR.MIDDLE,
        )
        action_id = str(action.get("action_id") or f"action-{index + 1:02d}")
        action_heading.name = f"CONTENT_binding_action_{action_id}_title_PUNCTGUARD"
        _add_panel(
            slide,
            6.22,
            top + 0.06,
            5.95,
            1.28,
            theme["accent_soft"] if index == 0 else theme["secondary_soft"],
            visual_role="success-signal",
        )
        _add_text(
            slide,
            "验证信号",
            6.42,
            top + 0.08,
            5.55,
            0.22,
            role="meta",
            size=10,
            color=row_colors[index],
            bold=True,
        )
        signal_shape = _add_text(
            slide,
            action.get("verification_signal", ""),
            6.42,
            top + 0.30,
            5.55,
            1.04,
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["ink"],
            bold=True,
        )
        signal_shape.name = f"CONTENT_binding_action_{action_id}_validation_signal"
    _add_footer(slide, item, theme)


def _render_appendix(prs, item, model, theme) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["cover_background"])
    profile = model["data_profile"]
    contracts = model["metric_contracts"]
    selected_finding_ids = set(str(value) for value in item.get("finding_ids", []))
    appendix_insights = [
        record for record in model.get("insights", [])
        if str(record.get("insight_id") or "") in selected_finding_ids
    ]
    _add_text(
        slide,
        item.get("visual_spec", {}).get("eyebrow", "数据与口径"),
        0.82,
        0.38,
        3.2,
        0.26,
        role="meta",
        size=12,
        color=theme["cover_muted"],
        bold=True,
        visual_role="eyebrow",
    )
    context_text = comparison_context_text(
        item.get("visual_spec", {}).get("comparison_contexts")
    )
    if context_text:
        _add_text(
            slide,
            context_text,
            5.42,
            0.38,
            6.78,
            0.24,
            role="meta",
            size=10,
            color=theme["cover_muted"],
            align=PP_ALIGN.RIGHT,
            visual_role="comparison-context",
        )
    _add_text(
        slide,
        _display_copy(
            item,
            item.get("visual_spec", {}).get("display_claim") or item["claim"],
        ),
        0.82,
        0.78,
        11.2,
        0.68,
        role="slide_title",
        size=40,
        color=theme["cover_ink"],
        bold=True,
        visual_role="title",
    )
    _add_text(
        slide,
        f"{profile['row_count']:,}",
        0.82,
        1.78,
        4.2,
        0.92,
        role="cover_title",
        size=50,
        color=theme["cover_ink"],
        bold=True,
        visual_role="data-scope",
    )
    _add_text(
        slide,
        f"行数据｜{profile['column_count']} 列｜重复行 {profile['duplicate_rows']}",
        0.82,
        2.70,
        4.30,
        0.34,
        role="body",
        size=FONT_MINIMUMS["body"],
        color=theme["cover_muted"],
    )
    _add_rule(slide, 0.82, 3.38, 4.15, theme["secondary"])
    timeline_lines = [
        f"时间范围　{profile['date_min']} → {profile['date_max']}",
        f"分析周期　{_localized_period(model['periods']['type'])}",
        f"未完成周期　{_localized_incomplete_policy(model['periods']['incomplete_period_policy'])}",
        f"数据来源　{model['source']['file_name']}",
        f"敏感字段　{', '.join(profile['sensitive_columns']) or '无'}",
    ]
    _add_text(
        slide,
        timeline_lines,
        0.82,
        3.68,
        4.30,
        2.20,
        role="body",
        size=FONT_MINIMUMS["body"],
        color=theme["cover_ink"],
    )
    _add_panel(
        slide,
        5.45,
        1.62,
        6.92,
        4.85,
        theme["surface"],
        visual_role="formula-grid",
    )
    _add_text(
        slide,
        "指标口径",
        5.76,
        1.92,
        2.8,
        0.38,
        role="subheading",
        size=FONT_MINIMUMS["subheading"],
        color=theme["accent"],
        bold=True,
    )
    contract_lines = [
        f"{contract.get('display_name') or contract['name']}：{contract.get('formula_display') or contract['formula']}（{contract['unit']}）"
        for contract in contracts
    ]
    columns = 2 if len(contract_lines) >= 4 else 1
    rows = (len(contract_lines) + columns - 1) // columns
    card_width = 3.0 if columns == 2 else 6.12
    for index, line in enumerate(contract_lines):
        column = index % columns
        row = index // columns
        left = 5.76 + column * 3.18
        top = 2.62 + row * (3.24 / max(rows, 1))
        _add_panel(slide, left, top, card_width, 0.86, theme["accent_soft"])
        _add_text(
            slide,
            line,
            left + 0.16,
            top + 0.10,
            card_width - 0.32,
            0.66,
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["ink"],
            vertical=MSO_ANCHOR.MIDDLE,
        )
    if appendix_insights:
        _add_text(
            slide,
            [
                _display_copy(
                    item,
                    record.get("headline") or record.get("statement") or "",
                )
                for record in appendix_insights
            ],
            0.82,
            5.92,
            11.0,
            0.78,
            role="body",
            size=FONT_MINIMUMS["body"],
            color=theme["ink"],
        )
    _add_text(
        slide,
        item["slide_id"],
        11.65,
        7.08,
        0.75,
        0.2,
        role="footnote",
        size=FONT_MINIMUMS["footnote"],
        color=theme["cover_muted"],
        align=PP_ALIGN.RIGHT,
    )


BUILDERS: dict[str, Callable] = {
    "cover": _render_cover,
    "executive-summary": _render_summary,
    "kpi-spotlight": _render_kpi,
    "trend-wide": _render_trend,
    "chart-insight-action": _render_trend,
    "chart-side-kpi": _render_chart_side,
    "driver-levers": _render_levers,
    "risk-opportunity": _render_risk,
    "action-roadmap": _render_actions,
    "data-appendix": _render_appendix,
}


def _validate_image_exceptions(model: dict[str, Any]) -> None:
    """Image use is opt-in and never a carrier for business evidence."""

    for candidate in [model, *list(model.get("slide_plan", []))]:
        raw_exceptions = candidate.get("image_exceptions", [])
        if raw_exceptions is None:
            continue
        if not isinstance(raw_exceptions, list):
            raise ValueError("image_exceptions must be a list")
        for exception in raw_exceptions:
            validate_image_exception(exception)


def build_pptx(model: dict[str, Any], output_path: Path, *, theme_selection: dict | None = None) -> None:
    validate_report_model(model)
    validate_layout_capacity(model)
    model = project_legacy_content_refs(model)
    _validate_image_exceptions(model)
    theme = load_theme(model["request"].get("theme", "clean"))
    if theme_selection is not None:
        from ppt_theme_selection import validate_selection
        selected = validate_selection(theme_selection)
        if theme_selection['target'] != 'standard':
            raise ValueError('theme selection target mismatch')
        if selected['id'] == 'editorial':
            raise ValueError('editorial layout is not supported by legacy PPT renderer; use the explicit themed presentation consumer')
        theme = selected['palette']
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH)
    presentation.slide_height = Inches(SLIDE_HEIGHT)
    for item in model["slide_plan"]:
        validate_visual_spec(item.get("visual_spec", {}))
        builder = BUILDERS.get(item["layout_id"])
        if builder is None:
            raise ValueError(f"未知 layout：{item['layout_id']}")
        builder(presentation, item, resolve_model_for_page(model, item), theme)
        _apply_compact_title_fallback(presentation.slides[-1], item)
        # A non-rendered slide XML name survives native PPTX edits and lets
        # polished decks be joined to the story by ID rather than position.
        presentation.slides[-1].element.cSld.set(
            "name", f"sheet-to-report:{item['slide_id']}"
        )
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build professional PPTX from report_model v2.")
    parser.add_argument("--model", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    model = json.loads(args.model.read_text(encoding="utf-8"))
    build_pptx(model, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
