#!/usr/bin/env python3
"""Extract text + embedded images from a source product PPTX.

Usage:
    python scripts/extract_source.py "<path/to/source.pptx>" [KEYWORD]

- Dumps all slide text to resources/extracted/source_text.txt
- Saves every picture to resources/images/slideNN_imgNN.<ext>
- Prints slides whose text mentions KEYWORD (case-insensitive; default "sinus")

Run this from the project working directory (the folder that contains
source_pack/, build/, output/, resources/).
"""
import os
import sys
from pptx import Presentation


def shape_text(sh):
    t = []
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            s = "".join(r.text for r in p.runs)
            if not s:
                s = p.text
            if s:
                t.append(s)
    if sh.has_table:
        for row in sh.table.rows:
            t.append(" | ".join(c.text for c in row.cells))
    return t


def main():
    if len(sys.argv) < 2:
        print("usage: python scripts/extract_source.py <source.pptx> [KEYWORD]")
        sys.exit(2)
    src = sys.argv[1]
    keyword = (sys.argv[2] if len(sys.argv) > 2 else "sinus").lower()

    out_txt = os.path.join("resources", "extracted", "source_text.txt")
    img_dir = os.path.join("resources", "images")
    os.makedirs(os.path.dirname(out_txt), exist_ok=True)
    os.makedirs(img_dir, exist_ok=True)

    prs = Presentation(src)
    lines = []
    hits = []
    img_count = 0

    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n===== SLIDE {i} =====")
        slide_txt = []
        for sh in slide.shapes:
            if sh.shape_type == 13:  # PICTURE
                try:
                    img = sh.image
                    ext = img.ext
                    fn = f"slide{i:02d}_img{img_count:02d}.{ext}"
                    with open(os.path.join(img_dir, fn), "wb") as f:
                        f.write(img.blob)
                    img_count += 1
                    lines.append(f"[IMAGE] {fn} ({sh.width}x{sh.height})")
                    continue
                except Exception as e:
                    lines.append(f"[IMAGE-ERR] {e}")
            txt = shape_text(sh)
            if txt:
                for line in txt:
                    lines.append(line)
                    slide_txt.append(line)
        blob = "\n".join(slide_txt)
        if keyword in blob.lower():
            hits.append((i, blob))

    with open(out_txt, "w") as f:
        f.write("\n".join(lines))

    print(f"SOURCE: {src}")
    print(f"TOTAL SLIDES: {len(prs.slides)}")
    print(f"IMAGES EXTRACTED: {img_count}")
    print(f"KEYWORD '{keyword}' FOUND ON SLIDES: {[h[0] for h in hits]}")
    print(f"\nTEXT DUMP: {out_txt}")
    if hits:
        print("\n----- KEYWORD SLIDE CONTENT (first 3000 chars each) -----")
        for idx, blob in hits:
            print(f"\n### SLIDE {idx} ###")
            print(blob[:3000])


if __name__ == "__main__":
    main()
