#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
西药房差错分析报告生成器
读取结构化 JSON 分析结果，生成符合三线表规范的 .docx 报告。

用法:
    python generate_report.py --input analysis.json --output 报告.docx

依赖: python-docx (pip install python-docx)
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    sys.stderr.write(
        "错误: 缺少 python-docx 依赖。\n"
        "请在隔离 venv 中安装:\n"
        "  pip install python-docx\n"
    )
    sys.exit(1)


# ---------- 字体与样式工具 ----------

CN_FONT = "Microsoft YaHei"
EN_FONT = "Calibri"


def set_run_font(run, size=10.5, bold=False, color=None):
    """统一设置 run 的中英文字体。"""
    run.font.name = EN_FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), CN_FONT)
    rFonts.set(qn("w:ascii"), EN_FONT)
    rFonts.set(qn("w:hAnsi"), EN_FONT)


def add_paragraph(doc, text, size=10.5, bold=False, align=None,
                  color=None, space_after=6, space_before=0, level=None):
    """添加段落并设置字体。"""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.5
    if level is not None:
        p.style = doc.styles[f"Heading {level}"]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def set_heading_font(doc):
    """让各级标题使用中文字体。"""
    for i in range(1, 5):
        try:
            style = doc.styles[f"Heading {i}"]
            style.font.name = EN_FONT
            rPr = style.element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.append(rFonts)
            rFonts.set(qn("w:eastAsia"), CN_FONT)
        except KeyError:
            pass


# ---------- 三线表工具 ----------

def _set_cell_border(cell, top=None, bottom=None, left=None, right=None,
                     inside_h=None, inside_v=None):
    """直接设置单元格各边边框。None 表示不修改;{'sz':..,'val':..,'color':..} 表示设置。"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    edges = {
        "top": top, "bottom": bottom, "left": left, "right": right,
        "insideH": inside_h, "insideV": inside_v,
    }
    for name, spec in edges.items():
        if spec is None:
            continue
        el = tcBorders.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            tcBorders.append(el)
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), str(spec.get("sz", 4)))
        el.set(qn("w:color"), spec.get("color", "000000"))


def make_three_line_table(doc, headers, rows, col_widths=None):
    """创建三线表: 顶线粗、表头下线中、底线粗, 无竖线与内部横线。"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    THICK = {"val": "single", "sz": 12, "color": "000000"}
    MID = {"val": "single", "sz": 6, "color": "000000"}
    NIL = {"val": "nil", "sz": 0, "color": "FFFFFF"}

    # 表头行
    hdr_cells = table.rows[0].cells
    for j, text in enumerate(headers):
        cell = hdr_cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        set_run_font(run, size=10, bold=True)
        # 顶线粗, 表头下线中, 无竖线
        _set_cell_border(cell, top=THICK, bottom=MID, left=NIL, right=NIL)

    # 数据行
    for i, row in enumerate(rows):
        cells = table.rows[i + 1].cells
        is_last = (i == len(rows) - 1)
        for j, val in enumerate(row):
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run("" if val is None else str(val))
            set_run_font(run, size=10)
            # 底线仅在最后一行加粗
            bottom = THICK if is_last else NIL
            _set_cell_border(cell, top=NIL, bottom=bottom, left=NIL, right=NIL)

    # 列宽
    if col_widths:
        for j, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[j].width = Cm(w)
    return table


# ---------- 报告组装 ----------

def build_report(data, output_path):
    doc = Document()
    set_heading_font(doc)

    # 默认正文样式字体
    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(10.5)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), CN_FONT)

    # ---- 标题 ----
    title = data.get("report_title", "西药房差错分析报告")
    add_paragraph(doc, title, size=18, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, level=1)

    meta = []
    if data.get("report_date"):
        meta.append(f"报告日期: {data['report_date']}")
    if data.get("prepared_by"):
        meta.append(f"编制: {data['prepared_by']}")
    if data.get("department"):
        meta.append(f"科室: {data['department']}")
    if meta:
        add_paragraph(doc, "    ".join(meta), size=10,
                      align=WD_ALIGN_PARAGRAPH.CENTER, color=(89, 89, 89),
                      space_after=12)

    # ---- 一、差错事件概述 ----
    add_paragraph(doc, "一、差错事件概述", size=14, bold=True, level=2, space_before=8)
    event = data.get("event", {})
    event_rows = [
        ["事件编号", event.get("event_id", "—")],
        ["发生时间", event.get("occur_date", "—")],
        ["发生地点", event.get("location", "—")],
        ["差错类型", event.get("error_type", "—")],
        ["危害等级", event.get("harm_level", "—")],
        ["发现经过", event.get("discovery_process", "—")],
        ["事件描述", event.get("event_description", "—")],
        ["涉及药品", event.get("drug_involved", "—")],
        ["涉及人员", event.get("personnel_involved", "—")],
    ]
    make_three_line_table(doc, ["项目", "内容"], event_rows, col_widths=[3.5, 12.5])

    # ---- 二、根因分析 ----
    rc = data.get("root_cause", {})
    add_paragraph(doc, "二、根因分析", size=14, bold=True, level=2, space_before=12)

    # 2.1 事件经过还原
    add_paragraph(doc, "2.1 事件经过还原", size=12, bold=True, level=3, space_before=6)
    add_paragraph(doc, rc.get("event_reconstruction", "（待补充事件经过时间线还原）"),
                  size=10.5, space_after=6)

    # 2.2 瑞士奶酪模型分析
    add_paragraph(doc, "2.2 瑞士奶酪模型——防线失效分析", size=12, bold=True,
                  level=3, space_before=6)
    sc = rc.get("swiss_cheese", [])
    if sc:
        sc_rows = [[item.get("layer", "—"),
                    item.get("status", "—"),
                    item.get("analysis", "—")] for item in sc]
        make_three_line_table(doc, ["防线层级", "状态", "失效分析"], sc_rows,
                              col_widths=[3.5, 2.5, 10.0])
    else:
        add_paragraph(doc, "（待补充各防线层级失效分析）", size=10.5)

    # 2.3 5Why 深度追溯
    add_paragraph(doc, "2.3 5Why——深度追溯", size=12, bold=True, level=3, space_before=6)
    fw = rc.get("five_whys", [])
    if fw:
        fw_rows = [[f"Why {i+1}", item.get("question", "—"),
                    item.get("answer", "—")] for i, item in enumerate(fw)]
        make_three_line_table(doc, ["层级", "追问", "回答"], fw_rows,
                              col_widths=[1.8, 5.5, 8.7])
    else:
        add_paragraph(doc, "（待补充 5Why 追溯）", size=10.5)

    # 2.4 鱼骨图——要因归集
    add_paragraph(doc, "2.4 鱼骨图——要因归集（5M1E）", size=12, bold=True,
                  level=3, space_before=6)
    fb = rc.get("fishbone", {})
    if fb and any(fb.values()):
        fb_rows = []
        for dim in ["人", "机", "料", "法", "环", "测"]:
            factors = fb.get(dim, [])
            fb_rows.append([dim, "、".join(factors) if factors else "—"])
        make_three_line_table(doc, ["维度", "可能要因"], fb_rows,
                              col_widths=[2.0, 14.0])
    else:
        add_paragraph(doc, "（待补充鱼骨图要因归集）", size=10.5)

    # 2.5 直接原因与根本原因
    add_paragraph(doc, "2.5 原因界定", size=12, bold=True, level=3, space_before=6)
    cause_rows = [
        ["直接原因", rc.get("direct_cause", "—")],
        ["潜在原因", rc.get("contributing_cause", "—")],
        ["根本原因（系统原因）", rc.get("root_cause", "—")],
    ]
    make_three_line_table(doc, ["原因层级", "描述"], cause_rows, col_widths=[4.0, 12.0])

    # ---- 三、改进建议 ----
    add_paragraph(doc, "三、改进建议", size=14, bold=True, level=2, space_before=12)
    improvements = data.get("improvements", [])
    if improvements:
        imp_rows = [[imp.get("dimension", "—"),
                     imp.get("measure", "—"),
                     imp.get("difficulty", "—"),
                     imp.get("owner", "—"),
                     imp.get("deadline", "—")] for imp in improvements]
        make_three_line_table(
            doc,
            ["维度(5M1E)", "改进措施", "难度", "责任人", "完成时限"],
            imp_rows, col_widths=[2.0, 7.0, 1.5, 2.5, 3.0],
        )
    else:
        add_paragraph(doc, "（待补充改进建议）", size=10.5)

    # 法规依据
    refs = data.get("regulatory_basis", [])
    if refs:
        add_paragraph(doc, "改进建议所依据的法规条款:", size=10.5, bold=True,
                      space_before=6, space_after=2)
        for r in refs:
            add_paragraph(doc, f"• {r}", size=10, space_after=2)

    # ---- 四、结论与跟进 ----
    add_paragraph(doc, "四、结论与跟进计划", size=14, bold=True, level=2, space_before=12)
    conclusion = data.get("conclusion", {})
    add_paragraph(doc, "4.1 结论", size=12, bold=True, level=3, space_before=6)
    add_paragraph(doc, conclusion.get("summary", "—"), size=10.5, space_after=6)

    add_paragraph(doc, "4.2 跟进计划", size=12, bold=True, level=3, space_before=6)
    follow_up = conclusion.get("follow_up", [])
    if follow_up:
        fu_rows = [[f"{i+1}", item] for i, item in enumerate(follow_up)]
        make_three_line_table(doc, ["序号", "跟进事项"], fu_rows,
                              col_widths=[1.5, 14.5])
    else:
        add_paragraph(doc, "（待补充跟进计划）", size=10.5)

    # ---- 页脚 ----
    sec = doc.sections[0]
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("本报告依据《医疗机构药事管理规定》《用药错误管理指导原则》编制  ·  仅供内部质量改进使用")
    set_run_font(fr, size=8, color=(128, 128, 128))

    output_path = Path(output_path)
    doc.save(str(output_path))
    return str(output_path)


def main():
    parser = argparse.ArgumentParser(description="生成西药房差错分析 Word 报告")
    parser.add_argument("--input", "-i", required=True,
                        help="结构化分析结果 JSON 文件路径")
    parser.add_argument("--output", "-o", required=True,
                        help="输出 .docx 文件路径")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    path = build_report(data, args.output)
    print(f"报告已生成: {path}")


if __name__ == "__main__":
    main()
