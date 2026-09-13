# -*- coding: utf-8 -*-
"""提取 PPTX 全部文本内容（标题、正文、表格、备注），供评分阅读使用。

用法:
    python extract_pptx.py <pptx路径> [输出md路径]

若不指定输出路径，结果打印到 stdout。
依赖: python-pptx
"""
import sys

from pptx import Presentation
from pptx.util import Emu


def shape_text(shape):
    """递归提取单个形状（含组合形状、表格）的文本。"""
    parts = []
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            parts.extend(shape_text(sub))
        return parts
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if text:
                parts.append(text)
    if shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            parts.append(" | ".join(cells))
    return parts


def main() -> None:
    if len(sys.argv) < 2:
        print("用法: python extract_pptx.py <pptx路径> [输出md路径]")
        sys.exit(1)

    pptx_path = sys.argv[1]
    prs = Presentation(pptx_path)

    slide_w = Emu(prs.slide_width).inches
    slide_h = Emu(prs.slide_height).inches

    lines = [
        f"# PPT 全文提取: {pptx_path}",
        f"幻灯片尺寸: {slide_w:.1f} x {slide_h:.1f} 英寸; 总页数: {len(prs.slides)}",
        "",
    ]

    for idx, slide in enumerate(prs.slides, 1):
        lines.append(f"\n## 第 {idx} 页")
        texts = []
        pics = 0
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                pics += 1
                continue
            texts.extend(shape_text(shape))
        if texts:
            for t in texts:
                lines.append(f"- {t}")
        else:
            lines.append("- (本页无文字)")
        if pics:
            lines.append(f"- [图片 x{pics}]")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"- [备注] {notes}")

    result = "\n".join(lines)
    if len(sys.argv) >= 3:
        with open(sys.argv[2], "w", encoding="utf-8") as f:
            f.write(result)
        print(f"已写入: {sys.argv[2]}")
    else:
        print(result)


if __name__ == "__main__":
    main()
