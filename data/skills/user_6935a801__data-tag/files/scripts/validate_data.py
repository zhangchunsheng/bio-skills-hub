#!/usr/bin/env python3
"""
Data Tag - 数据标注校验脚本

功能：
1. 读取 Excel 文件并输出 JSON 格式数据
2. 分析数据特征（标签分布、文本特征、数据质量）
3. 将校验结果写入新 Excel 文件，在结果字段 cell 上添加批注和颜色标注：
   - 错误的 cell 标红色 + 添加批注
   - 待确认的 cell 标黄色 + 添加批注
   - 不新增任何列
   - 输出文件命名：原文件名-标注版.xlsx

用法：
  python3 validate_data.py read <excel_path> [--sheet <sheet_name>] [--limit <n>]
  python3 validate_data.py analyze <excel_path> --data-fields <fields> --result-fields <fields> [--sheet <sheet_name>]
  python3 validate_data.py write <excel_path> <output_path> --result-field <field> --results <json_results>
  python3 validate_data.py write-batch <excel_path> <output_path> --all-results <json_results>
"""

import argparse
import json
import sys
import os

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from openpyxl.comments import Comment
except ImportError:
    print("ERROR: openpyxl is not installed. Install it with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

try:
    import pandas as pd
except ImportError:
    print("ERROR: pandas is not installed. Install it with: pip install pandas", file=sys.stderr)
    sys.exit(1)

# 颜色定义
RED_FILL = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")
RED_FONT = Font(color="FFFFFF", bold=True)
YELLOW_FILL = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
YELLOW_FONT = Font(bold=True)


def read_excel(excel_path, sheet_name=None, limit=None):
    """读取 Excel 文件并输出 JSON 格式数据"""
    if not os.path.exists(excel_path):
        print(json.dumps({"error": f"File not found: {excel_path}"}, ensure_ascii=False))
        sys.exit(1)

    try:
        if sheet_name:
            df = pd.read_excel(excel_path, sheet_name=sheet_name)
        else:
            df = pd.read_excel(excel_path)

        if limit:
            df = df.head(limit)

        # 将 NaN 替换为 None，便于 JSON 序列化
        result = {
            "columns": list(df.columns),
            "row_count": len(df),
            "data": df.where(df.notna(), None).to_dict(orient="records")
        }
        print(json.dumps(result, ensure_ascii=False, default=str))
    except Exception as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        sys.exit(1)


def analyze_data(excel_path, data_fields_str, result_fields_str, sheet_name=None):
    """
    分析数据特征，帮助 AI 在校验前充分理解数据结构和标签体系。

    输出包括：
    - 总行数
    - 每个结果字段的标签分布（值 + 频次 + 占比）
    - 每个数据字段的特征（类型、平均长度、空值率、示例值）
    - 数据质量信号（是否可能是语音转写、文本长度分布等）
    """
    if not os.path.exists(excel_path):
        print(json.dumps({"error": f"File not found: {excel_path}"}, ensure_ascii=False))
        sys.exit(1)

    try:
        if sheet_name:
            df = pd.read_excel(excel_path, sheet_name=sheet_name)
        else:
            df = pd.read_excel(excel_path)

        data_fields = [f.strip() for f in data_fields_str.split(",")]
        result_fields = [f.strip() for f in result_fields_str.split(",")]

        # 验证字段存在性
        missing = [f for f in data_fields + result_fields if f not in df.columns]
        if missing:
            print(json.dumps({"error": f"Fields not found: {missing}, available: {list(df.columns)}"}, ensure_ascii=False))
            sys.exit(1)

        analysis = {
            "total_rows": len(df),
            "columns": list(df.columns),
            "result_fields_analysis": {},
            "data_fields_analysis": {},
            "data_quality_signals": {}
        }

        # 分析结果字段（标签分布）
        for field in result_fields:
            col = df[field]
            value_counts = col.value_counts(dropna=False)
            label_dist = {}
            for val, count in value_counts.items():
                key = str(val) if pd.notna(val) else "<空值>"
                label_dist[key] = {
                    "count": int(count),
                    "ratio": round(count / len(df) * 100, 1)
                }
            analysis["result_fields_analysis"][field] = {
                "unique_values": int(col.nunique(dropna=False)),
                "null_count": int(col.isna().sum()),
                "null_ratio": round(col.isna().sum() / len(df) * 100, 1),
                "label_distribution": label_dist
            }

        # 分析数据字段（内容特征）
        for field in data_fields:
            col = df[field]
            field_info = {
                "dtype": str(col.dtype),
                "null_count": int(col.isna().sum()),
                "null_ratio": round(col.isna().sum() / len(df) * 100, 1),
            }

            if col.dtype == "object":
                # 文本类型字段的详细分析
                text_col = col.dropna().astype(str)
                if len(text_col) > 0:
                    lengths = text_col.str.len()
                    field_info["text_stats"] = {
                        "avg_length": round(lengths.mean(), 1),
                        "min_length": int(lengths.min()),
                        "max_length": int(lengths.max()),
                        "median_length": round(lengths.median(), 1),
                    }
                    # 检测是否可能是语音转写文本
                    asr_signals = 0
                    sample_texts = text_col.head(20)
                    for text in sample_texts:
                        # 口语化标志
                        if any(w in text for w in ["嗯", "那个", "就是说", "啊", "哦", "喂"]):
                            asr_signals += 1
                        # 缺少标点
                        if len(text) > 50 and text.count("，") + text.count("。") + text.count(",") + text.count(".") < 2:
                            asr_signals += 1
                    field_info["possible_asr_text"] = asr_signals > len(sample_texts) * 0.3

                    # 样本值（取前3条非空值，截断到200字符）
                    samples = text_col.head(3).tolist()
                    field_info["sample_values"] = [s[:200] + ("..." if len(s) > 200 else "") for s in samples]
            else:
                # 数值类型字段
                num_col = col.dropna()
                if len(num_col) > 0:
                    field_info["numeric_stats"] = {
                        "min": float(num_col.min()),
                        "max": float(num_col.max()),
                        "mean": round(float(num_col.mean()), 2),
                        "median": round(float(num_col.median()), 2),
                    }
                    field_info["sample_values"] = num_col.head(3).tolist()

            analysis["data_fields_analysis"][field] = field_info

        # 整体数据质量信号
        has_asr = any(
            analysis["data_fields_analysis"].get(f, {}).get("possible_asr_text", False)
            for f in data_fields
        )
        analysis["data_quality_signals"] = {
            "possible_asr_data": has_asr,
            "suggestion": "数据可能为语音转写文本，校验时需考虑转写噪声（同音字、漏字、口语化等），不要因字面错误而误判标注" if has_asr else "数据为常规文本/数值，按标准流程校验"
        }

        print(json.dumps(analysis, ensure_ascii=False, default=str))
    except Exception as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        sys.exit(1)


def write_results(excel_path, output_path, result_field, results_json):
    """
    将校验结果写入新 Excel 文件，在结果字段 cell 上添加批注和颜色标注，不新增任何列。

    Args:
        excel_path: 原始 Excel 文件路径
        output_path: 输出 Excel 文件路径（xx-标注版.xlsx）
        result_field: 结果字段名（需要标色和添加批注的列）
        results_json: 校验结果 JSON 字符串，格式为：
            [{"校验结果": "准确/错误/待确认", "批示": "原因说明"}, ...]
            - 错误和待确认的行会在结果字段 cell 上添加批注
            - 准确的行不做任何标记
    """
    if not os.path.exists(excel_path):
        print(json.dumps({"error": f"File not found: {excel_path}"}, ensure_ascii=False))
        sys.exit(1)

    try:
        results = json.loads(results_json)

        # 用 openpyxl 加载工作簿以支持批注和颜色标注
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        # 获取表头行
        headers = [cell.value for cell in ws[1]]

        # 找到结果字段的列索引（1-based）
        result_col_idx = None
        for idx, header in enumerate(headers, start=1):
            if header == result_field:
                result_col_idx = idx
                break

        if result_col_idx is None:
            print(json.dumps({"error": f"Result field '{result_field}' not found in columns: {headers}"}, ensure_ascii=False))
            sys.exit(1)

        # 校验结果数量必须与数据行数一致
        data_row_count = ws.max_row - 1
        if len(results) != data_row_count:
            print(json.dumps({
                "error": f"Results count ({len(results)}) does not match data rows ({data_row_count})"
            }, ensure_ascii=False))
            sys.exit(1)

        # 逐行处理：为错误/待确认的结果字段 cell 添加批注和颜色
        accurate_count = 0
        error_count = 0
        pending_count = 0

        for i, r in enumerate(results, start=2):  # 从第2行开始（跳过表头）
            status = r.get("校验结果", "待确认")
            remark = r.get("批示", "")

            result_cell = ws.cell(row=i, column=result_col_idx)

            if status == "错误":
                # 红色背景 + 白色加粗字体
                result_cell.fill = RED_FILL
                result_cell.font = RED_FONT
                # 添加批注
                if remark:
                    result_cell.comment = Comment(remark, "data-tag", width=300, height=100)
                error_count += 1
            elif status == "待确认":
                # 黄色背景 + 加粗字体
                result_cell.fill = YELLOW_FILL
                result_cell.font = YELLOW_FONT
                # 添加批注
                if remark:
                    result_cell.comment = Comment(remark, "data-tag", width=300, height=100)
                pending_count += 1
            else:
                accurate_count += 1

        # 保存到新文件
        wb.save(output_path)

        # 统计摘要
        summary = {
            "total": len(results),
            "accurate": accurate_count,
            "error": error_count,
            "pending": pending_count,
            "output_path": output_path
        }
        print(json.dumps(summary, ensure_ascii=False))
    except json.JSONDecodeError as e:
        print(json.dumps({"error": f"Invalid JSON in results: {e}"}, ensure_ascii=False))
        sys.exit(1)
    except Exception as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        sys.exit(1)


def write_batch_results(excel_path, output_path, all_results_json):
    """
    一次性写入所有结果字段的校验结果，避免多次加载导致标注覆盖。

    Args:
        excel_path: 原始 Excel 文件路径
        output_path: 输出 Excel 文件路径（xx-标注版.xlsx）
        all_results_json: 所有字段的校验结果 JSON 字符串，格式为：
            {
                "字段名1": [{"校验结果": "准确/错误/待确认", "批示": "原因说明"}, ...],
                "字段名2": [{"校验结果": "准确/错误/待确认", "批示": "原因说明"}, ...],
                ...
            }
    """
    if not os.path.exists(excel_path):
        print(json.dumps({"error": f"File not found: {excel_path}"}, ensure_ascii=False))
        sys.exit(1)

    try:
        all_results = json.loads(all_results_json)

        # 用 openpyxl 加载工作簿（只加载一次）
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        # 获取表头行，建立字段名到列索引的映射
        headers = [cell.value for cell in ws[1]]
        header_to_col = {}
        for idx, header in enumerate(headers, start=1):
            if header is not None:
                header_to_col[header] = idx

        data_row_count = ws.max_row - 1

        # 统计
        total_errors = 0
        total_pending = 0
        total_accurate = 0
        field_stats = {}

        for field_name, results in all_results.items():
            # 检查字段是否存在
            if field_name not in header_to_col:
                print(json.dumps({"error": f"Field '{field_name}' not found in columns: {headers}"}, ensure_ascii=False))
                sys.exit(1)

            # 校验结果数量必须与数据行数一致
            if len(results) != data_row_count:
                print(json.dumps({
                    "error": f"Results count for '{field_name}' ({len(results)}) does not match data rows ({data_row_count})"
                }, ensure_ascii=False))
                sys.exit(1)

            col_idx = header_to_col[field_name]
            accurate_count = 0
            error_count = 0
            pending_count = 0

            for i, r in enumerate(results, start=2):  # 从第2行开始（跳过表头）
                status = r.get("校验结果", "待确认")
                remark = r.get("批示", "")

                result_cell = ws.cell(row=i, column=col_idx)

                if status == "错误":
                    result_cell.fill = RED_FILL
                    result_cell.font = RED_FONT
                    if remark:
                        result_cell.comment = Comment(remark, "data-tag", width=300, height=100)
                    error_count += 1
                elif status == "待确认":
                    result_cell.fill = YELLOW_FILL
                    result_cell.font = YELLOW_FONT
                    if remark:
                        result_cell.comment = Comment(remark, "data-tag", width=300, height=100)
                    pending_count += 1
                else:
                    accurate_count += 1

            total_errors += error_count
            total_pending += pending_count
            total_accurate += accurate_count
            field_stats[field_name] = {
                "accurate": accurate_count,
                "error": error_count,
                "pending": pending_count
            }

        # 保存到新文件（只保存一次）
        wb.save(output_path)

        # 统计摘要
        total_cells = sum(len(v) for v in all_results.values())
        summary = {
            "total_fields": len(all_results),
            "total_cells": total_cells,
            "accurate": total_accurate,
            "error": total_errors,
            "pending": total_pending,
            "field_stats": field_stats,
            "output_path": output_path
        }
        print(json.dumps(summary, ensure_ascii=False))
    except json.JSONDecodeError as e:
        print(json.dumps({"error": f"Invalid JSON in results: {e}"}, ensure_ascii=False))
        sys.exit(1)
    except Exception as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        sys.exit(1)


def main():
    parser = argparse.ArgumentParser(description="Data Tag - 数据标注校验工具")
    subparsers = parser.add_subparsers(dest="command", help="可用命令")

    # read 命令
    read_parser = subparsers.add_parser("read", help="读取 Excel 文件")
    read_parser.add_argument("excel_path", help="Excel 文件路径")
    read_parser.add_argument("--sheet", help="工作表名称（默认读取第一个）")
    read_parser.add_argument("--limit", type=int, help="限制读取行数")

    # analyze 命令（数据特征分析）
    analyze_parser = subparsers.add_parser("analyze", help="分析数据特征和标签分布")
    analyze_parser.add_argument("excel_path", help="Excel 文件路径")
    analyze_parser.add_argument("--data-fields", required=True, help="数据字段名，多个用逗号分隔")
    analyze_parser.add_argument("--result-fields", required=True, help="结果字段名，多个用逗号分隔")
    analyze_parser.add_argument("--sheet", help="工作表名称（默认读取第一个）")

    # write 命令（单字段，保留向后兼容）
    write_parser = subparsers.add_parser("write", help="写入单个字段的校验结果")
    write_parser.add_argument("excel_path", help="原始 Excel 文件路径")
    write_parser.add_argument("output_path", help="输出 Excel 文件路径（xx-标注版.xlsx）")
    write_parser.add_argument("--result-field", required=True, help="结果字段名（需要标色和添加批注的列）")
    write_parser.add_argument("--results", required=True, help="校验结果 JSON 字符串")

    # write-batch 命令（多字段一次性写入）
    batch_parser = subparsers.add_parser("write-batch", help="一次性写入所有字段的校验结果")
    batch_parser.add_argument("excel_path", help="原始 Excel 文件路径")
    batch_parser.add_argument("output_path", help="输出 Excel 文件路径（xx-标注版.xlsx）")
    batch_parser.add_argument("--all-results", required=True, help="所有字段的校验结果 JSON 字符串")

    args = parser.parse_args()

    if args.command == "read":
        read_excel(args.excel_path, sheet_name=args.sheet, limit=args.limit)
    elif args.command == "analyze":
        analyze_data(
            args.excel_path,
            args.data_fields,
            args.result_fields,
            sheet_name=args.sheet
        )
    elif args.command == "write":
        write_results(
            args.excel_path,
            args.output_path,
            args.result_field,
            args.results
        )
    elif args.command == "write-batch":
        write_batch_results(
            args.excel_path,
            args.output_path,
            args.all_results
        )
    else:
        parser.print_help()
        sys.exit(1)


if __name__ == "__main__":
    main()
