#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 Markdown 格式的行业调查报告转换为 Word（.docx）格式
用法：python md_to_docx.py input.md output.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import sys
import os

def add_markdown_to_docx(doc, md_text):
    """将 Markdown 文本添加到 Word 文档中"""
    lines = md_text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # 空行
        if not line.strip():
            i += 1
            continue
        
        # 标题
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            p = doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            p = doc.add_heading(line[5:], level=4)
        
        # 引用块
        elif line.startswith('> '):
            p = doc.add_paragraph(line[2:])
            p.style = 'Intense Quote'
        
        # 分隔线
        elif line.strip() == '---':
            doc.add_paragraph('─' * 50)
        
        # 表格
        elif line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            i -= 1  # 回退一行
            
            # 解析表格
            table_data = []
            for tl in table_lines:
                # 分割单元格，去掉首尾的 |
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                table_data.append(cells)
            
            # 跳过分隔行（| --- | --- |）
            if len(table_data) >= 2:
                # 检查第二行是否全是分隔符
                if all(re.match(r'^-+$', c.strip()) for c in table_data[1] if c.strip()):
                    table_data = [table_data[0]] + table_data[2:]
            
            if table_data and len(table_data[0]) > 0:
                try:
                    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                    table.style = 'Light Grid Accent 1'
                    
                    for row_idx, row_data in enumerate(table_data):
                        for col_idx, cell_data in enumerate(row_data):
                            if col_idx < len(table.rows[row_idx].cells):
                                # 去掉 Markdown 加粗标记
                                cell_text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_data)
                                table.rows[row_idx].cells[col_idx].text = cell_text
                                
                                # 第一行加粗
                                if row_idx == 0:
                                    for paragraph in table.rows[row_idx].cells[col_idx].paragraphs:
                                        for run in paragraph.runs:
                                            run.bold = True
                except Exception as e:
                    print(f"表格解析失败：{e}")
                    doc.add_paragraph(f"表格解析失败：{table_lines}")
        
        # 无序列表
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)  # 去掉加粗
            doc.add_paragraph(text, style='List Bullet')
        
        # 有序列表
        elif re.match(r'^\d+\. ', line.strip()):
            text = re.sub(r'^\d+\. ', '', line.strip())
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)  # 去掉加粗
            doc.add_paragraph(text, style='List Number')
        
        # 普通段落
        else:
            # 去掉 Markdown 加粗标记
            text = line
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            # 去掉 Markdown 链接标记
            text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)
            doc.add_paragraph(text)
        
        i += 1

def convert_md_to_docx(md_file, docx_file):
    """将 Markdown 文件转换为 Word 文档"""
    # 创建新文档
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(10.5)
    
    # 读取 Markdown 文件
    with open(md_file, 'r', encoding='utf-8') as f:
        md_text = f.read()
    
    # 转换
    add_markdown_to_docx(doc, md_text)
    
    # 保存
    doc.save(docx_file)
    print(f"已生成 Word 文档：{docx_file}")

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("用法：python md_to_docx.py input.md output.docx")
        sys.exit(1)
    
    md_file = sys.argv[1]
    docx_file = sys.argv[2]
    
    if not os.path.exists(md_file):
        print(f"文件不存在：{md_file}")
        sys.exit(1)
    
    convert_md_to_docx(md_file, docx_file)
