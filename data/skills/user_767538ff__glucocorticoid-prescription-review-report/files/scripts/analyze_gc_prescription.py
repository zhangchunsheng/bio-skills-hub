#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
糖皮质激素处方/医嘱点评数据通用分析脚本
==========================================
功能：读取一份或多份糖皮质激素专项点评 xlsx 表格，自动识别表类型（住院医嘱 / 门急诊处方），
      完成基础统计与不合理明细提取，输出结构化文本结果，供后续撰写六步分析报告使用。

依赖：openpyxl（pip install openpyxl）

用法：
    python analyze_gc_prescription.py <xlsx路径1> [xlsx路径2 ...] [--out 输出结果文件路径]

表类型自动识别规则：
    - 含"出院科室"或"存在问题"列  → 住院医嘱表（inpatient）
    - 含"就诊/入院科室"或"处方号"列 → 门急诊处方表（outpatient）

输出内容：
    1. 各表基础信息（表头、总记录数、审核结果分布、时间范围）
    2. 科室分布
    3. 涉及糖皮质激素药品清单（去重）与通用名使用频次
    4. 不合理记录逐条明细（编号/处方号、科室、年龄、性别、诊断、激素药品、存在问题原文）
    5. 不合理记录科室分布
    6. 多表合并汇总（总条数、不合理条数、不合理率）
"""

import sys
import argparse
import openpyxl
from collections import Counter

# ===== 糖皮质激素通用名识别关键词（覆盖注射/口服/吸入/鼻喷/外用/滴眼） =====
GC_KEYS = [
    "甲泼尼龙", "甲强龙", "甲基强的松龙",
    "泼尼松", "强的松",
    "地塞米松", "氟美松",
    "布地奈德",
    "倍氯米松", "丙酸倍氯米松",
    "氟替卡松", "丙酸氟替卡松", "糠酸氟替卡松",
    "莫米松", "糠酸莫米松",
    "氢化可的松",
    "曲安奈德", "曲安西龙",
    "倍他米松",
    "可的松",
    "氟轻松",
    "泼尼松龙", "氢化泼尼松",
    "环索奈德",
]


def is_gc(name):
    """判断药品名是否为糖皮质激素，返回匹配关键词或 None"""
    if not name:
        return None
    for k in GC_KEYS:
        if k in name:
            return k
    return None


def extract_gc_drugs(drug_field):
    """从药品字段（逗号分隔）中提取糖皮质激素药品全名列表"""
    if not drug_field:
        return []
    drugs = [d.strip() for d in str(drug_field).split(",") if d.strip()]
    return [d for d in drugs if is_gc(d)]


def gc_generic_list(drug_field):
    """提取药品字段中出现的糖皮质激素通用名（去重，保持顺序）"""
    if not drug_field:
        return []
    found = []
    for k in GC_KEYS:
        if k in str(drug_field) and k not in found:
            found.append(k)
    return found


def detect_table_type(header):
    """根据表头列名自动识别表类型"""
    cols = set(str(c) for c in header)
    if "出院科室" in cols or "存在问题" in cols or "出院诊断" in cols:
        return "inpatient"
    if "就诊/入院科室" in cols or "处方号" in cols or "费别" in cols:
        return "outpatient"
    return "unknown"


def safe_str(v):
    return "" if v is None else str(v)


def analyze_file(path, out_lines):
    """分析单个 xlsx 文件，结果追加到 out_lines"""
    wb = openpyxl.load_workbook(path, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [safe_str(c) for c in rows[0]]
        data = rows[1:]
        ttype = detect_table_type(header)
        out_lines.append("=" * 80)
        out_lines.append(f"【文件】{path}")
        out_lines.append(f"【工作表】{sn} | 表类型={ttype} | 数据行数={len(data)}")
        out_lines.append("【表头】" + " | ".join(header))

        idx = {name: i for i, name in enumerate(header)}

        # —— 确定关键列（兼容两种表）——
        res_col = idx.get("审核结果")
        dept_col = idx.get("出院科室") or idx.get("就诊/入院科室")
        age_col = idx.get("年龄")
        sex_col = idx.get("性别")
        diag_col = idx.get("诊断")
        outdiag_col = idx.get("出院诊断")
        drug_col = idx.get("药品")
        prob_col = idx.get("存在问题")
        create_col = idx.get("创建时间")
        visit_col = idx.get("就诊/入院时间") or idx.get("就诊/入院时间")
        admit_col = idx.get("就诊/入院时间")
        id_col = idx.get("点评编号") or idx.get("处方号")
        fee_col = idx.get("费别")

        # 审核结果分布
        res_counter = Counter()
        for r in data:
            res_counter[safe_str(r[res_col]) if res_col is not None else ""] += 1
        out_lines.append("【审核结果分布】" + dict(res_counter).__str__())

        # 时间范围
        if create_col is not None:
            ts = [r[create_col] for r in data if r[create_col]]
            if ts:
                out_lines.append(f"【创建时间范围】{min(ts)} ~ {max(ts)}")
        if admit_col is not None:
            ts = [r[admit_col] for r in data if r[admit_col]]
            if ts:
                out_lines.append(f"【就诊/入院时间范围】{min(ts)} ~ {max(ts)}")

        # 科室分布
        if dept_col is not None:
            dept_counter = Counter(safe_str(r[dept_col]) for r in data)
            out_lines.append("【科室分布(全部)】")
            for d, c in dept_counter.most_common():
                out_lines.append(f"  {d}: {c}")

        # 激素药品清单
        gc_set = set()
        gc_generic_counter = Counter()
        if drug_col is not None:
            for r in data:
                gf = r[drug_col]
                for d in extract_gc_drugs(gf):
                    gc_set.add(d)
                for g in gc_generic_list(gf):
                    gc_generic_counter[g] += 1
        out_lines.append("【涉及糖皮质激素药品(去重)】")
        for d in sorted(gc_set):
            out_lines.append(f"  {d}")
        out_lines.append("【激素通用名使用频次(按条数)】")
        for g, c in gc_generic_counter.most_common():
            out_lines.append(f"  {g}: {c}")

        # 不合理明细
        unreasonable = []
        if res_col is not None:
            unreasonable = [r for r in data if safe_str(r[res_col]) == "不合理"]
        out_lines.append(f"【不合理记录明细 ({len(unreasonable)}条)】")
        for i, r in enumerate(unreasonable, 1):
            gc_list = extract_gc_drugs(r[drug_col]) if drug_col is not None else []
            gc_str = " ; ".join(gc_list) if gc_list else "(未识别到激素,请核对)"
            out_lines.append(f"[{i}] ID={safe_str(r[id_col]) if id_col is not None else ''} | 科室={safe_str(r[dept_col]) if dept_col is not None else ''} | 年龄={safe_str(r[age_col]) if age_col is not None else ''} | 性别={safe_str(r[sex_col]) if sex_col is not None else ''}")
            if diag_col is not None:
                out_lines.append(f"    诊断: {safe_str(r[diag_col])}")
            if outdiag_col is not None and r[outdiag_col]:
                out_lines.append(f"    出院诊断: {safe_str(r[outdiag_col])}")
            out_lines.append(f"    激素药品: {gc_str}")
            if fee_col is not None and r[fee_col]:
                out_lines.append(f"    费别: {safe_str(r[fee_col])}")
            if prob_col is not None and r[prob_col]:
                out_lines.append(f"    存在问题(原文): {safe_str(r[prob_col])}")
            else:
                out_lines.append("    存在问题(原文): (原表未填写,需依据指导原则补充判定)")
            out_lines.append("")

        # 不合理科室分布
        if dept_col is not None and unreasonable:
            un_dept = Counter(safe_str(r[dept_col]) for r in unreasonable)
            out_lines.append("【不合理记录科室分布】")
            for d, c in un_dept.most_common():
                out_lines.append(f"  {d}: {c}")

        # 返回统计用于汇总
        return {
            "type": ttype,
            "total": len(data),
            "unreasonable": len(unreasonable),
            "res_counter": dict(res_counter),
        }
    return None


def main():
    parser = argparse.ArgumentParser(description="糖皮质激素点评数据通用分析")
    parser.add_argument("files", nargs="+", help="一个或多个 xlsx 文件路径")
    parser.add_argument("--out", help="结果输出文件路径（默认输出到 stdout）")
    args = parser.parse_args()

    out_lines = []
    out_lines.append("糖皮质激素专项点评数据分析结果")
    out_lines.append("=" * 80)

    summaries = []
    for f in args.files:
        try:
            s = analyze_file(f, out_lines)
            if s:
                summaries.append(s)
        except Exception as e:
            out_lines.append(f"[错误] 分析 {f} 失败: {e}")

    # 合并汇总
    out_lines.append("")
    out_lines.append("=" * 80)
    out_lines.append("【多表合并汇总】")
    total = sum(s["total"] for s in summaries)
    total_un = sum(s["unreasonable"] for s in summaries)
    out_lines.append(f"抽查总条数: {total}")
    out_lines.append(f"不合理条数: {total_un}")
    if total > 0:
        out_lines.append(f"总体不合理率: {total_un/total*100:.2f}%")
        out_lines.append(f"总体合理率: {(total-total_un)/total*100:.2f}%")
    for s in summaries:
        label = {"inpatient": "住院", "outpatient": "门急诊", "unknown": "未知"}.get(s["type"], s["type"])
        rate = f"{s['unreasonable']/s['total']*100:.2f}%" if s["total"] > 0 else "N/A"
        out_lines.append(f"  {label}表: {s['unreasonable']}/{s['total']} (不合理率 {rate})")

    result = "\n".join(out_lines)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fp:
            fp.write(result)
        print(f"分析结果已写入: {args.out}")
    else:
        print(result)


if __name__ == "__main__":
    main()
