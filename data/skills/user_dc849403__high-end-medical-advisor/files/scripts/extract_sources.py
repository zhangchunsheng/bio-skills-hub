import os, re, glob
from docx import Document
from pptx import Presentation

WX = r"C:\Users\PC\xwechat_files\Ymondhuang_bbeb\msg\file\2026-09"
SRC_FILES = [
    os.path.join(WX, "副本理赔案例话术.docx"),
    os.path.join(WX, "中高医展业工具（2026年8月版上）.pptx"),
    os.path.join(WX, "国家政策体检销售闭环讲座.pptx"),
    os.path.join(WX, "中高医展业工具（2026年8月版下）.pptx"),
]
D_DIR = r"D:\中高端医疗险\中高医话术"
for f in sorted(glob.glob(os.path.join(D_DIR, "*.docx"))):
    SRC_FILES.append(f)

OUT_DIR = r"C:\Users\PC\.workbuddy\skills\high-end-medical-advisor\references\source"
os.makedirs(OUT_DIR, exist_ok=True)


def safe_name(p):
    base = os.path.splitext(os.path.basename(p))[0]
    for ch in '\\/:*?"<>|':
        base = base.replace(ch, "_")
    return base + ".md"


def extract_docx(path):
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        style = (p.style.name or "") if p.style else ""
        text = p.text.strip()
        if not text:
            continue
        if style.startswith("Heading") or style.startswith("标题"):
            level = 1
            m = re.search(r"\d+", style)
            if m:
                level = min(int(m.group()), 6)
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
    for ti, table in enumerate(doc.tables):
        lines.append(f"\n**表格 {ti + 1}**")
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
        lines.append("")
    return "\n".join(lines)


def extract_pptx(path):
    prs = Presentation(path)
    lines = []
    for si, slide in enumerate(prs.slides, 1):
        lines.append(f"\n## 幻灯片 {si}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"_备注: {notes}_")
    return "\n".join(lines)


for path in SRC_FILES:
    if not os.path.exists(path):
        print("MISSING:", path)
        continue
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".docx":
            content = extract_docx(path)
        elif ext == ".pptx":
            content = extract_pptx(path)
        else:
            print("SKIP:", path)
            continue
        out = os.path.join(OUT_DIR, safe_name(path))
        with open(out, "w", encoding="utf-8") as f:
            f.write(f"# {os.path.basename(path)}\n\n> 来源文件：{path}\n\n")
            f.write(content)
        print("OK ->", os.path.basename(out), "chars:", len(content))
    except Exception as e:
        print("ERROR", path, e)
