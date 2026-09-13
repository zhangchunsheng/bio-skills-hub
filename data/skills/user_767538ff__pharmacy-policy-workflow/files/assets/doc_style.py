# -*- coding: utf-8 -*-
"""药剂科文档统一排版辅助（深蓝标题+宋体+1.5倍行距）。供 pharmacy-policy-workflow skill 复用。"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)

def set_cjk(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color

def style_doc(doc):
    st = doc.styles['Normal']
    st.font.name = '宋体'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

def h1(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); set_cjk(r, "黑体", 14, True, NAVY); return p

def h2(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); set_cjk(r, "黑体", 12, True, NAVY); return p

def para(doc, text, size=10.5, bold=False, color=None, indent=False):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    if indent: p.paragraph_format.left_indent = Cm(0.75)
    r = p.add_run(text); set_cjk(r, "宋体", size, bold, color); return p

def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text); set_cjk(r, "宋体", size); return p

def title(doc, text, sub=None):
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(text); set_cjk(r, "黑体", 16, True, NAVY)
    if sub:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cjk(p.add_run(sub), "宋体", 10, color=GREY)

def table(doc, headers, rows, widths=None, fontsize=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h); set_cjk(run, "黑体", fontsize, True, RGBColor(0xFF,0xFF,0xFF))
        shd = hdr[i]._tc.get_or_add_tcPr().makeelement(qn('w:shd'), {qn('w:val'):'clear',qn('w:fill'):'1F4E79'})
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            set_cjk(cells[i].paragraphs[0].add_run(str(v)), "宋体", fontsize)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    return t
