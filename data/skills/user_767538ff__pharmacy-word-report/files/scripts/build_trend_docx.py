# -*- coding: utf-8 -*-
"""月度趋势分析 Word 报表生成器：合并 202606 与 202607 两月药库出库数据。
读取：out_compare/comparison_results.json, out_compare/dept_dive.json,
      out_202606/analysis_results.json, out_202607/analysis_results.json
输出：全院药品使用月度趋势分析_202606-202607.docx
"""
import json, os, sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CN_FONT, H_FONT = "宋体", "黑体"
try:
    font_manager.fontManager.addfont(r"C:\Windows\Fonts\simsun.ttc")
    font_manager.fontManager.addfont(r"C:\Windows\Fonts\simhei.ttf")
    plt.rcParams["font.sans-serif"] = ["SimHei"]
except Exception:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

COMPARE = sys.argv[1] if len(sys.argv) > 1 else "out_compare/comparison_results.json"
DIVE = sys.argv[2] if len(sys.argv) > 2 else "out_compare/dept_dive.json"
J606 = sys.argv[3] if len(sys.argv) > 3 else "out_202606/analysis_results.json"
J607 = sys.argv[4] if len(sys.argv) > 4 else "out_202607/analysis_results.json"
DOCX = sys.argv[5] if len(sys.argv) > 5 else "全院药品使用月度趋势分析_202606-202607.docx"

C = json.load(open(COMPARE, encoding="utf-8"))
D = json.load(open(DIVE, encoding="utf-8"))
R6 = json.load(open(J606, encoding="utf-8"))
R7 = json.load(open(J607, encoding="utf-8"))
FIG = "out_compare/figs_trend"
os.makedirs(FIG, exist_ok=True)

WAN = lambda x: f"{x/1e4:,.2f}" if x else "0.00"


# ---------------- helpers ----------------
def set_cn(run, font=CN_FONT, size=9, bold=False):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold


def set_cell(cell, text, size=8.5, bold=False, align="CENTER", font=CN_FONT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"CENTER": WD_ALIGN_PARAGRAPH.CENTER, "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
                   "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(str(text))
    set_cn(r, font, size, bold)


def three_line_table(t):
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_cn(r, CN_FONT, 8.5)
            tcPr = cell._tc.get_or_add_tcPr()
            borders = tcPr.find(qn("w:tcBorders"))
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tcPr.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                e = borders.find(qn(f"w:{edge}"))
                if e is None:
                    e = OxmlElement(f"w:{edge}")
                    borders.append(e)
                e.set(qn("w:val"), "nil")
    n = len(t.rows)
    def put_border(row_idx, edge, val="single", sz="6"):
        for cell in t.rows[row_idx].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = tcPr.find(qn("w:tcBorders"))
            if b is None:
                b = OxmlElement("w:tcBorders"); tcPr.append(b)
            e = b.find(qn(f"w:{edge}"))
            if e is None:
                e = OxmlElement(f"w:{edge}"); b.append(e)
            e.set(qn("w:val"), val); e.set(qn("w:sz"), sz)
    put_border(0, "top"); put_border(0, "bottom")        # 栏目线（表头上下）
    put_border(n - 1, "bottom")                            # 底线
    # 顶线（表头上方那条粗线）
    for cell in t.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders"); tcPr.append(b)
        e = b.find(qn("w:top"))
        if e is None:
            e = OxmlElement("w:top"); b.append(e)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "12")


def style_header(t, size=8.5):
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                set_cn(r, H_FONT, size, bold=True)


def H(doc, text, level=1):
    p = doc.add_heading("", level=level)
    r = p.add_run(text); set_cn(r, H_FONT, 13 if level == 1 else 11, bold=True)
    return p


def para(doc, text, size=10, bold=False, align="LEFT", font=CN_FONT):
    p = doc.add_paragraph()
    p.alignment = {"LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
                   "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(text); set_cn(r, font, size, bold)
    return p


def concl(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("【结论】 " + text); set_cn(r, CN_FONT, 10, bold=True)
    return p


def fill_table(t, headers, rows, widths=None, size=8.5):
    while len(t.rows) < len(rows) + 1:
        t.add_row()
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, size, bold=True, font=H_FONT)
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row):
            set_cell(t.rows[ri].cells[ci], v, size)
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)


# ---------------- figures ----------------
def fig_total():
    fig, ax = plt.subplots(figsize=(5.2, 3.6))
    labels = ["202606", "202607"]
    vals = [C["totals"]["m606"]["total_amount_wan"], C["totals"]["m607"]["total_amount_wan"]]
    bars = ax.bar(labels, vals, color=["#5B8FF9", "#F6BD16"])
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width()/2, v, f"{v/1e4:,.2f}亿", ha="center", va="bottom", fontsize=9)
    ax.set_ylabel("金额(万元)"); ax.set_title("全院药品出库总金额月度对比")
    fig.tight_layout(); p = os.path.join(FIG, "total.png"); fig.savefig(p, dpi=130); plt.close(fig); return p


def fig_top():
    top = C["top_drug_mom"][:10]
    names = [t["generic"] for t in top]
    a6 = [t["amt_606"]/1e4 for t in top]; a7 = [t["amt_607"]/1e4 for t in top]
    import numpy as np
    x = np.arange(len(names)); w = 0.4
    fig, ax = plt.subplots(figsize=(8, 4.2))
    ax.bar(x - w/2, a6, w, label="202606", color="#5B8FF9")
    ax.bar(x + w/2, a7, w, label="202607", color="#F6BD16")
    ax.set_xticks(x); ax.set_xticklabels(names, rotation=35, ha="right", fontsize=7)
    ax.set_ylabel("金额(万元)"); ax.set_title("Top10 药品两月金额对比"); ax.legend()
    fig.tight_layout(); p = os.path.join(FIG, "top.png"); fig.savefig(p, dpi=130); plt.close(fig); return p


def fig_atc():
    a1 = C["atc1_mom"][:10]
    names = [a["atc1"] for a in a1]
    s6 = [a["share_606"] for a in a1]; s7 = [a["share_607"] for a in a1]
    import numpy as np
    x = np.arange(len(names)); w = 0.4
    fig, ax = plt.subplots(figsize=(8, 4.4))
    ax.bar(x - w/2, s6, w, label="202606", color="#5B8FF9")
    ax.bar(x + w/2, s7, w, label="202607", color="#F6BD16")
    ax.set_xticks(x); ax.set_xticklabels(names, rotation=30, ha="right", fontsize=7)
    ax.set_ylabel("占比(%)"); ax.set_title("ATC1级分类金额占比月度对比"); ax.legend()
    fig.tight_layout(); p = os.path.join(FIG, "atc.png"); fig.savefig(p, dpi=130); plt.close(fig); return p


p_total, p_top, p_atc = fig_total(), fig_top(), fig_atc()

# ---------------- document ----------------
doc = Document()
sec = doc.sections[0]
sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5); sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
st = doc.styles["Normal"]; st.font.name = CN_FONT; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

# 标题
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = tp.add_run("全院药品使用月度趋势分析"); set_cn(tr, H_FONT, 18, bold=True)
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sp.add_run("（数据来源：药库出库　统计周期：2026年6月—7月　统计粒度：全院各科室）"); set_cn(sr, CN_FONT, 10)

# 一、总体环比
H(doc, "一、总体指标月度环比")
t = C["totals"]; d = t["delta"]
concl(doc, f"全院药品出库总金额由 202606 的 {t['m606']['total_amount_wan']/1e4:,.2f} 亿元升至 202607 的 "
         f"{t['m607']['total_amount_wan']/1e4:,.2f} 亿元，环比 {d['amount_pct']:+.2f}%（+{d['amount_wan']:,.0f}万元）；"
         f"覆盖科室由 {t['m606']['depts']} 个增至 {t['m607']['depts']} 个（+{d['depts']}），"
         f"药品通用名由 {t['m606']['drugs']} 个变为 {t['m607']['drugs']} 个（{d['drugs']:+d}），总体规模小幅扩张。")
tb = doc.add_table(rows=1, cols=6)
rows = [
    ["出库记录数", f"{t['m606']['rows']:,}", f"{t['m607']['rows']:,}", f"{d['rows']:+,.0f}", ""],
    ["覆盖科室数", str(t['m606']['depts']), str(t['m607']['depts']), f"{d['depts']:+d}", ""],
    ["药品通用名数", str(t['m606']['drugs']), str(t['m607']['drugs']), f"{d['drugs']:+d}", ""],
    ["出库总金额(万元)", f"{t['m606']['total_amount_wan']:,.0f}", f"{t['m607']['total_amount_wan']:,.0f}",
     f"{d['amount_wan']:+,.0f}", f"{d['amount_pct']:+.2f}%"],
    ["Top50品种合计占比(%)", f"{t['m606']['top50_share']:.2f}", f"{t['m607']['top50_share']:.2f}",
     f"{d['top50_share_delta']:+.2f}", "集中度"],
]
fill_table(tb, ["指标", "202606", "202607", "环比变化", "说明"], rows, widths=[3.6, 2.8, 2.8, 2.6, 1.6])
three_line_table(tb); style_header(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
doc.add_picture(p_total, width=Cm(9))
para(doc, "图0　全院药品出库总金额月度对比", size=8.5, align="CENTER")
para(doc, "说明：以上金额均直接取自原始出库记录、按科室/通用名如实汇总，未作任何缺失值填补或数据臆造；"
         "202606 有 54 行 ATC 分类缺失（布洛芬注射剂、水合氯醛口服溶液剂），已按缺失如实处理。", size=8.5)

# 二、Top药品环比
H(doc, "二、重点药品金额与排名环比")
concl(doc, f"两月均在 Top50 的品种共 {C['both_top50_count']} 个；Top 品种金额总体稳定，头部品种（抗肿瘤/高值专科药）"
         "月度波动主要受临床收治量影响。下表为两月金额合计前 20 名品种的金额与排名变化。")
doc.add_picture(p_top, width=Cm(15))
para(doc, "图1　Top10 药品两月金额对比", size=8.5, align="CENTER")
top_rows = []
for r in C["top_drug_mom"][:20]:
    pct_s = f"{r['pct']:+.1f}%" if r["pct"] is not None else "新增"
    top_rows.append([r["generic"], f"{r['amt_606']/1e4:,.1f}", f"{r['amt_607']/1e4:,.1f}",
                     pct_s, str(r["rank_606"] or "-"), str(r["rank_607"] or "-")])
tt = doc.add_table(rows=1, cols=6)
fill_table(tt, ["通用名", "202606(万元)", "202607(万元)", "环比", "排名606", "排名607"], top_rows,
           widths=[4.0, 2.6, 2.6, 2.2, 1.8, 1.8])
three_line_table(tt); style_header(tt); tt.alignment = WD_TABLE_ALIGNMENT.CENTER

# 三、ATC分类环比
H(doc, "三、ATC 分类金额占比月度环比")
concl(doc, "抗肿瘤药及免疫调节剂持续居首且占比稳定；全身抗感染药、血液系统药次之。ATC 结构两月基本一致，"
         "未见大类用药方向性偏移。")
doc.add_picture(p_atc, width=Cm(15))
para(doc, "图2　ATC1级分类金额占比月度对比", size=8.5, align="CENTER")
atc_rows = []
for a in C["atc1_mom"][:12]:
    atc_rows.append([a["atc1"], f"{a['amt_606_wan']:,.0f}", f"{a['amt_607_wan']:,.0f}",
                     f"{a['share_606']:.1f}", f"{a['share_607']:.1f}", f"{a['share_delta']:+.1f}"])
ta = doc.add_table(rows=1, cols=6)
fill_table(ta, ["ATC1级分类", "202606(万元)", "202607(万元)", "占比606(%)", "占比607(%)", "占比变化"], atc_rows,
           widths=[4.2, 2.4, 2.4, 2.0, 2.0, 1.8])
three_line_table(ta); style_header(ta); ta.alignment = WD_TABLE_ALIGNMENT.CENTER

# 四、存疑科室深挖
H(doc, "四、专科-主导药品匹配存疑科室单科室明细深挖")
concl(doc, f"两月合计检出 {len(D)} 个科室存在'专科性质与主导药品（抗肿瘤药及免疫调节剂）明显不匹配'的情形，"
         "需药师逐科核实药品归属/代发录入。值得注意的是，部分科室两月主导药发生明显切换"
         "（如中医针灸门诊由伏美替尼转为比卡鲁胺、中医内科病房由贝伐珠单抗转为伊沙佐米），"
         "更高度提示非真实本科室消耗，而是跨科代发或录入归科错误。")
# 存疑科室两月对照表
sm = doc.add_table(rows=1, cols=6)
sm_rows = []
for r in C["susp_mom"]:
    sm_rows.append([r["科室"], "是" if r["in_606"] else "—", "是" if r["in_607"] else "—",
                    f"{r['主导药品_606']}({r['主导占比_606']:.1f}%)" if r["in_606"] else "—",
                    f"{r['主导药品_607']}({r['主导占比_607']:.1f}%)" if r["in_607"] else "—",
                    ("本月切换" if (r["in_606"] and r["in_607"] and r["主导药品_606"] != r["主导药品_607"]) else "持续")])
fill_table(sm, ["科室", "606存疑", "607存疑", "606主导药(占比)", "607主导药(占比)", "月度变化"], sm_rows,
           widths=[3.0, 1.6, 1.6, 3.2, 3.2, 1.6])
three_line_table(sm); style_header(sm); sm.alignment = WD_TABLE_ALIGNMENT.CENTER
para(doc, "表4　存疑科室两月对照（主导药括号内为占本科室出库金额比）", size=8.5, align="CENTER")

# 逐科室明细
for r in D:
    dept = r["科室"]
    # 结论行
    months = r["出现月份"]
    notes = []
    for m in r["明细"]:
        if m["匹配存疑"]:
            notes.append(f"{m['period']} 主导{m['主导药品']}占{m['主导占比']:.1f}%")
    switch = len({m["主导药品"] for m in r["明细"]}) > 1
    line = f"{dept}：{'；'.join(notes)}"
    if switch:
        line += "　→ 两月主导药切换，强烈提示代发/归科录入问题，建议优先核查。"
    else:
        line += "　→ 持续存疑，建议核查。"
    concl(doc, line)
    # 明细表：合并两月 Top药品
    # 取两月 Top药品并集（最多6个）
    drug_map = {}
    for m in r["明细"]:
        for tr in m["top"]:
            g = tr["generic"]
            if g not in drug_map:
                drug_map[g] = {"atc1": tr["atc1"], "606": None, "607": None}
            drug_map[g][m["period"]] = tr
    ordered = sorted(drug_map.items(), key=lambda kv: (kv[1].get("607") or kv[1].get("606") or {"金额": 0})["金额"], reverse=True)[:6]
    drows = []
    for g, info in ordered:
        a6 = info.get("606"); a7 = info.get("607")
        drows.append([g,
                      f"{a6['金额_万元']:,.1f}({a6['占比']:.1f}%)" if a6 else "—",
                      f"{a7['金额_万元']:,.1f}({a7['占比']:.1f}%)" if a7 else "—",
                      info["atc1"] or "缺失"])
    dt = doc.add_table(rows=1, cols=4)
    fill_table(dt, [dept + "　逐月药品明细（金额万元/占比）", "202606", "202607", "ATC1"], drows,
               widths=[4.2, 3.4, 3.4, 2.2])
    three_line_table(dt); style_header(dt); dt.alignment = WD_TABLE_ALIGNMENT.CENTER

# 五、结论与建议
H(doc, "五、综合结论与药学干预建议")
concl(doc, f"1）规模：202607 全院出库 {t['m607']['total_amount_wan']/1e4:,.2f} 亿元，环比 {d['amount_pct']:+.2f}%，"
         "结构以抗肿瘤及高值专科药为主导，符合大型三甲医院专科用药特征。")
para(doc, "2）趋势：Top50 品种集中度两月稳定在约 50%，头部品种波动主要来自收治量；ATC 大类结构稳定。")
para(doc, "3）风险：共 11 个专科科室主导药为抗肿瘤药及免疫调节剂且专科性质明显不符，其中多个科室两月主导药切换，"
         "高度提示'跨科代发/药品归科录入错误'。建议：① 对康复理疗科病房（达雷妥尤单抗 64.3%）、"
         "中医针灸门诊、中医内科病房等优先复核 HIS 药品归属与请领科室；② 对抗肿瘤单抗类药品建立科室级消耗监测；"
         "③ 推动药库出库按'实际使用科室'而非'请领/代发科室'归集，从根源消除归科偏差。")
para(doc, "4）数据治理：202606 仍有 54 行 ATC 未分类，建议药学信息部门补全 ATC 映射，保障后续自动化分析完整性。")

# 附录
H(doc, "附录　方法学声明", level=2)
para(doc, "本报告由 pharmacy-word-report 技能脚本自动生成：analyze.py 负责画像/聚合/出图，compare.py 负责两月环比，"
         "dept_dive.py 负责存疑科室明细，build_trend_docx.py 负责成稿。所有金额均直接取自原始出库记录，"
         "按通用名/ATC/科室如实汇总，未对缺失值做任何填补或臆造。金额单位万元为主、亿元用于量级表述；"
         "占比为占全院出库金额比。专科-主导药品'匹配存疑'判定口径：科室名含康复/中医/皮肤/眼科/口腔/耳鼻喉/生殖/"
         "妇产/保健/营养/心理/精神/针灸/风湿等关键词，但其金额第一药品属'抗肿瘤药及免疫调节剂'，列为需药师核实，"
         "对阿达木单抗（风湿/眼科）、环孢素（生殖）等临床合理情形不视为不合理。")

doc.save(DOCX)
print("Saved:", DOCX, "| tables:", len(doc.tables))
