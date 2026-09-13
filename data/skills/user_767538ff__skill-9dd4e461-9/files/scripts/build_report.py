#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
医嘱点评助手 - Word 报告生成器（JSON 驱动）
把结构化点评数据渲染为标准化 Word 报告。

用法:
    python build_report.py --json review_spec.json --out 医嘱点评报告.docx

JSON 规格见 references/json_schema.md。
"""
import json
import argparse
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

CN = "Microsoft YaHei"
C_DARK = RGBColor(0x1F, 0x4E, 0x79)
C_BLUE = RGBColor(0x2E, 0x74, 0xB5)


def set_cn(run, name=CN):
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)


def setup_normal(doc):
    style = doc.styles["Normal"]
    style.font.name = CN
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN)


def h1(doc, text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    set_cn(r)
    r.font.size = Pt(15)
    r.font.color.rgb = C_DARK
    return p


def h2(doc, text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    set_cn(r)
    r.font.size = Pt(12)
    r.font.color.rgb = C_BLUE
    return p


def para(doc, text, bold=False, indent=False, size=10.5, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn(r)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if indent:
        p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_after = Pt(3)
    return p


def bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r1 = p.add_run(bold_lead)
        set_cn(r1)
        r1.bold = True
        r2 = p.add_run(text)
        set_cn(r2)
    else:
        r = p.add_run(text)
        set_cn(r)
    p.paragraph_format.space_after = Pt(2)
    return p


def build(doc, spec):
    setup_normal(doc)

    # 封面/总论
    title = doc.add_heading(level=0)
    tr = title.add_run(spec.get("title", "医嘱点评报告（用药合理性评价）"))
    set_cn(tr)
    tr.font.size = Pt(20)
    tr.font.color.rgb = C_DARK

    sub = doc.add_paragraph()
    obj = spec.get("review_objects", "（见各病例）")
    sr = sub.add_run("点评对象：" + obj + "\n")
    set_cn(sr)
    sr.font.size = Pt(10.5)
    basis = spec.get("review_basis", "")
    if basis:
        sr2 = sub.add_run("点评依据：" + basis)
        set_cn(sr2)
        sr2.font.size = Pt(10.5)

    # 模板学习要点（可选）
    tn = spec.get("template_note")
    if tn:
        para(doc, "一、模板学习要点", bold=True)
        para(doc, tn, indent=True)

    doc.add_page_break()

    # 逐例
    cases = spec.get("cases", [])
    for idx, c in enumerate(cases, start=1):
        h1(doc, c.get("heading", "病例 %d" % idx))

        h2(doc, "病例摘要")
        para(doc, c.get("summary", "（病历未提供）"))

        h2(doc, "主要用药情况")
        meds = c.get("medications", [])
        if meds:
            for m in meds:
                bullet(doc, m)
        else:
            para(doc, "（病历未提供用药明细）", indent=True)

        h2(doc, "用药合理性点评")
        for r in c.get("reasonable", []):
            bullet(doc, r, bold_lead="（一）合理：")
        concerns = c.get("concerns", [])
        if concerns:
            for cc in concerns:
                bullet(doc, cc, bold_lead="（二）需关注/需纠正：")

        h2(doc, "点评结论")
        para(doc, c.get("conclusion", "（略）"))

        doc.add_page_break()

    # 汇总表
    rows = spec.get("summary_rows")
    if rows:
        h1(doc, "用药合理性点评汇总")
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Light Grid Accent 1"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = tbl.rows[0].cells
        for i, t in enumerate(["病例", "主要诊断", "点评结论", "主要改进/关注点"]):
            hdr[i].text = ""
            rp = hdr[i].paragraphs[0].add_run(t)
            set_cn(rp)
            rp.bold = True
            rp.font.size = Pt(10)
        for r in rows:
            cells = tbl.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""
                rp = cells[i].paragraphs[0].add_run(v)
                set_cn(rp)
                rp.font.size = Pt(9.5)
        doc.add_paragraph()

    note = spec.get("note")
    if note:
        np_ = doc.add_paragraph()
        nr = np_.add_run(note)
        set_cn(nr)
        nr.font.size = Pt(9)
        nr.italic = True


def main():
    ap = argparse.ArgumentParser(description="医嘱点评助手 Word 生成器")
    ap.add_argument("--json", required=True, help="点评规格 JSON 文件")
    ap.add_argument("--out", required=True, help="输出 .docx 路径")
    args = ap.parse_args()

    with open(args.json, "r", encoding="utf-8") as fh:
        spec = json.load(fh)

    doc = Document()
    build(doc, spec)
    doc.save(args.out)
    print("saved:", args.out)


if __name__ == "__main__":
    main()
