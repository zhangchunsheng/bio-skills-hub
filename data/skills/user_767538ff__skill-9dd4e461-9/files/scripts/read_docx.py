#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
医嘱点评助手 - docx 内容提取器
把一份或多份 .docx 的段落与表格文本导出，供点评分析使用。

用法:
    python read_docx.py "病例1.docx" "病例2.docx" [--out extracted.txt]
若不指定 --out，则输出到 stdout。
"""
import docx
import os
import sys
import argparse


def dump_file(path):
    lines = []
    lines.append("=" * 80)
    lines.append("FILE: " + os.path.basename(path))
    lines.append("=" * 80)
    try:
        d = docx.Document(path)
    except Exception as e:
        lines.append("[无法打开文档: %s]" % e)
        lines.append("")
        return lines

    # 段落
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if t:
            # 保留样式提示（标题级别），便于识别章节
            style = p.style.name if p.style else ""
            if style.startswith("Heading") or style in ("Title",):
                lines.append("[%s] %s" % (style, t))
            else:
                lines.append(t)

    # 表格
    if d.tables:
        lines.append("")
        lines.append("--- TABLES ---")
        for ti, tbl in enumerate(d.tables):
            lines.append("")
            lines.append("[Table %d] (rows=%d, cols=%d)" % (ti + 1, len(tbl.rows), len(tbl.columns)))
            for row in tbl.rows:
                cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
                lines.append(" | ".join(cells))
    lines.append("")
    return lines


def main():
    ap = argparse.ArgumentParser(description="医嘱点评助手 docx 提取器")
    ap.add_argument("files", nargs="+", help="一个或多个 .docx 文件路径")
    ap.add_argument("--out", default=None, help="输出文本文件，默认 stdout")
    args = ap.parse_args()

    all_lines = []
    for f in args.files:
        if not os.path.exists(f):
            all_lines.append("[文件不存在: %s]" % f)
            continue
        all_lines.extend(dump_file(f))

    text = "\n".join(all_lines)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text)
        print("已导出 %d 份文档内容 -> %s" % (len(args.files), args.out))
    else:
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
