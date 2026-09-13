# -*- coding: utf-8 -*-
"""
临床药学服务相关政策汇总表生成脚本模板
基于 assets/临床药学服务相关政策汇总表模板.xlsx 生成输出文件
用法：修改 DATA 列表后运行 python gen_policy_excel.py
"""
import shutil, os
from datetime import datetime
from copy import copy
import openpyxl
from openpyxl.styles import Alignment

# === 配置 ===
# 模板路径（相对于skill根目录）
TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        'assets', '临床药学服务相关政策汇总表模板.xlsx')
# 输出路径（按需修改）
OUTPUT = './临床药学服务相关政策汇总表-输出.xlsx'
# 汇总标题（修改日期）
TITLE = '临床药学服务相关政策汇总表——汇总日期：____年__月__日'

# === 数据（9列：序号, 发布机构, 发布日期, 标题, 主要内容, 是否有更新, 更新日期, 更新后的标题, 备注）===
# 发布日期：能确定用 datetime(Y,M,D)，无法确定用字符串如 "2025年5月"
DATA = [
    # 示例数据，替换为实际整理的政策
    (1, '国家卫生健康委办公厅', datetime(2021, 10, 9),
     '国家卫生健康委办公厅关于印发医疗机构药学门诊服务规范等5项规范的通知 国卫办医函〔2021〕520号',
     '制定并印发《医疗机构药学门诊服务规范》《药物重整服务规范》《用药教育服务规范》《药学监护服务规范》《居家药学服务规范》5项规范，规范发展药学服务',
     '否', None, '', ''),
    # ... 继续添加政策数据
]


def csv_field(v):
    """转义CSV字段（含逗号/引号时用双引号包裹）"""
    s = str(v) if v is not None else ''
    if ',' in s or '"' in s or '\n' in s:
        s = s.replace('"', '""')
        s = '"' + s + '"'
    return s


def generate():
    # 1. 复制模板
    shutil.copyfile(TEMPLATE, OUTPUT)
    wb = openpyxl.load_workbook(OUTPUT)
    ws = wb.active

    # 2. 修改标题
    ws.cell(row=1, column=1).value = TITLE

    # 3. 清空第3行起旧数据（保留样式）
    for r in range(3, ws.max_row + 1):
        for c in range(1, 10):
            ws.cell(row=r, column=c).value = None

    # 4. 填入数据（复制第3行样式）
    style_src_row = 3
    start_row = 3
    for i, row_data in enumerate(DATA):
        r = start_row + i
        for c, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            src = ws.cell(row=style_src_row, column=c)
            if src.has_style:
                cell.font = copy(src.font)
                cell.fill = copy(src.fill)
                cell.border = copy(src.border)
                cell.alignment = copy(src.alignment)
                cell.number_format = src.number_format
            # 日期列设置格式
            if c == 3 and isinstance(val, datetime):
                cell.number_format = 'yyyy"年"m"d"日"'
            # 主要内容、备注列自动换行
            if c in (5, 9):
                cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
            # 其他文本列顶部对齐
            if c in (1, 2, 3, 4, 6, 7, 8):
                cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
        ws.row_dimensions[r].height = 75

    # 5. 保存
    wb.save(OUTPUT)
    print(f'文件已生成: {OUTPUT}')
    print(f'共 {len(DATA)} 条政策')


if __name__ == '__main__':
    generate()
