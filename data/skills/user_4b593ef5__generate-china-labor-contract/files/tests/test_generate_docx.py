import unittest
import json
import subprocess
import sys
import zipfile
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from scripts.generate_docx import extract_contract_text, write_docx


class ExtractContractTextTest(unittest.TestCase):
    def test_excludes_post_contract_guidance_sections(self):
        source = """# 策略判断卡
推荐合同类型：固定期限劳动合同

# 劳动合同

甲方（用人单位）名称：，

第一条 合同期限
合同期限自____年__月__日起至____年__月__日止。

【已重点处理以下用工风险相关条款】
- 已补足薪酬结构、支付周期和支付日期。

待补充字段
- 甲方名称
"""

        extracted = extract_contract_text(source)

        self.assertIn("# 劳动合同", extracted)
        self.assertIn("甲方（用人单位）名称：__________，", extracted)
        self.assertIn("第一条 合同期限", extracted)
        self.assertNotIn("策略判断卡", extracted)
        self.assertNotIn("已重点处理", extracted)
        self.assertNotIn("待补充字段", extracted)

    def test_recognizes_bold_markdown_contract_title_and_heading_markers(self):
        source = """# 策略判断卡
推荐合同类型：固定期限劳动合同

## **劳动合同**

第一条 合同期限
合同期限自____年__月__日起至____年__月__日止。

## 【已重点处理以下用工风险相关条款】
- 已补足薪酬结构。
"""

        extracted = extract_contract_text(source)

        self.assertTrue(extracted.startswith("## **劳动合同**"))
        self.assertNotIn("策略判断卡", extracted)
        self.assertNotIn("已重点处理", extracted)

    def test_requires_contract_title_unless_raw_mode_is_explicit(self):
        source = """# 策略判断卡
推荐合同类型：固定期限劳动合同

甲方（用人单位）名称：__________
"""

        self.assertEqual("", extract_contract_text(source))
        self.assertIn("策略判断卡", extract_contract_text(source, raw=True))

    def test_clean_docx_uses_a4_and_has_no_highlight(self):
        text = """# 劳动合同

第一条 劳动报酬
乙方工资为人民币____元/月。
"""

        with TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "contract.docx"
            write_docx(text, output_path, mode="clean")

            document = Document(output_path)
            section = document.sections[0]
            self.assertAlmostEqual(section.page_width / 36000, 210, delta=1)
            self.assertAlmostEqual(section.page_height / 36000, 297, delta=1)
            with zipfile.ZipFile(output_path) as docx_zip:
                document_xml = docx_zip.read("word/document.xml")
            self.assertNotIn(b"w:highlight", document_xml)

    def test_clean_docx_uses_table_signature_area(self):
        text = """# 劳动合同

甲方（盖章）：__________
乙方（签字）：__________
日期：____年__月__日
"""

        with TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "contract.docx"
            write_docx(text, output_path, mode="clean")

            document = Document(output_path)
            self.assertGreaterEqual(len(document.tables), 1)
            cells_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("甲方（盖章）：", cells_text)
            self.assertIn("乙方（签字）：", cells_text)

    def test_clean_docx_uses_table_for_standard_signature_block(self):
        text = """# 劳动合同

甲方（盖章）：__________
法定代表人或授权代表（签字）：__________
乙方（签字）：__________
日期：____年__月__日
"""

        with TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "contract.docx"
            write_docx(text, output_path, mode="clean")

            document = Document(output_path)
            self.assertGreaterEqual(len(document.tables), 1)
            self.assertEqual(len(document.paragraphs), 1)
            cells_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("甲方（盖章）：", cells_text)
            self.assertIn("法定代表人或授权代表（签字）：", cells_text)
            self.assertIn("乙方（签字）：", cells_text)

    def test_cli_accepts_structured_facts_file(self):
        text = """# 劳动合同

甲方（用人单位）名称：__________
乙方（劳动者）姓名：__________
合同期限：固定期限，自____年__月__日起至____年__月__日止。
合同期限一年。试用期二个月，试用期工资为转正工资的80%。
工作内容：__________。工作地点：__________。
工作时间和休息休假：标准工时制，依法休息休假。
劳动报酬：工资人民币____元/月。加班工资按工作日150%、休息日不能补休则200%、法定节假日300%执行。
社会保险：甲方依法缴纳社会保险。
劳动保护和劳动条件：甲方依法提供劳动保护和劳动条件。
解除终止：双方依法解除或终止。
劳动争议：依法申请劳动仲裁。
甲方（盖章）：__________
乙方（签字）：__________
"""

        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_path = tmp_path / "contract.md"
            facts_path = tmp_path / "facts.json"
            output_path = tmp_path / "contract.docx"
            input_path.write_text(text, encoding="utf-8")
            contract_text = extract_contract_text(text)
            facts_path.write_text(
                json.dumps(
                    {
                        "schema_version": 1,
                        "document_mode": "final",
                        "term_type": "fixed",
                        "term_months": 12,
                        "probation_months": 2,
                        "probation_wage_ratio": 0.8,
                        "prior_probation_used": False,
                        "noncompete_mode": "reference_only",
                        "noncompete_person_eligible": False,
                        "secret_access_confirmed": False,
                        "noncompete_months": None,
                        "noncompete_compensation_confirmed": False,
                        "social_insurance": "statutory",
                        "contract_sha256": sha256(contract_text.encode("utf-8")).hexdigest(),
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    "scripts/generate_docx.py",
                    "--input",
                    str(input_path),
                    "--facts",
                    str(facts_path),
                    "--output",
                    str(output_path),
                    "--mode",
                    "clean",
                ],
                cwd=Path(__file__).resolve().parents[1],
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(output_path.exists())

    def test_cli_requires_facts_for_export(self):
        text = """# 劳动合同

甲方（用人单位）名称：__________
乙方（劳动者）姓名：__________
合同期限：固定期限，自____年__月__日起至____年__月__日止。
工作内容：__________。工作地点：__________。
工作时间和休息休假：标准工时制，依法休息休假。
劳动报酬：工资人民币____元/月。加班工资按工作日150%、休息日不能补休则200%、法定节假日300%执行。
社会保险：甲方依法缴纳社会保险。
劳动保护和劳动条件：甲方依法提供劳动保护和劳动条件。
解除终止：双方依法解除或终止。
劳动争议：依法申请劳动仲裁。
甲方（盖章）：__________
乙方（签字）：__________
"""

        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_path = tmp_path / "contract.md"
            output_path = tmp_path / "contract.docx"
            input_path.write_text(text, encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    "scripts/generate_docx.py",
                    "--input",
                    str(input_path),
                    "--output",
                    str(output_path),
                    "--mode",
                    "clean",
                ],
                cwd=Path(__file__).resolve().parents[1],
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("FACTS_REQUIRED", result.stderr)

    def test_validate_cli_reports_bad_json_without_traceback(self):
        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_path = tmp_path / "contract.md"
            facts_path = tmp_path / "facts.json"
            input_path.write_text("# 劳动合同\n", encoding="utf-8")
            facts_path.write_text("{bad json", encoding="utf-8")

            result = subprocess.run(
                [
                    sys.executable,
                    "scripts/validate_contract.py",
                    "--input",
                    str(input_path),
                    "--facts",
                    str(facts_path),
                ],
                cwd=Path(__file__).resolve().parents[1],
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("ERROR", result.stderr)
            self.assertNotIn("Traceback", result.stderr)

    def test_generate_cli_reports_missing_input_without_traceback(self):
        with TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "contract.docx"

            result = subprocess.run(
                [
                    sys.executable,
                    "scripts/generate_docx.py",
                    "--input",
                    str(Path(tmp) / "missing.md"),
                    "--output",
                    str(output_path),
                    "--mode",
                    "clean",
                ],
                cwd=Path(__file__).resolve().parents[1],
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("ERROR", result.stderr)
            self.assertNotIn("Traceback", result.stderr)


if __name__ == "__main__":
    unittest.main()
