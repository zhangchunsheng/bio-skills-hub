#!/usr/bin/env python3
"""Generate a DOCX file from finalized labor contract text.

Usage:
  python3 scripts/generate_docx.py --input contract.md --facts facts.json --output output/劳动合同.docx --mode clean
  cat contract.md | python3 scripts/generate_docx.py --facts facts.json --output output/劳动合同.docx --mode draft
  python3 scripts/generate_docx.py --input raw-contract.md --facts facts.json --output output/劳动合同.docx --raw

The script intentionally depends only on the finalized contract text. It does
not read or require the local development templates directory. When fed a
broader skill response, it keeps only the formal contract body and excludes
compliance notes, pending-field checklists, and DOCX action menus. By default,
the CLI validates hard legal rules before export.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from contract_text import is_contract_title as shared_is_contract_title
except ImportError:  # pragma: no cover - package import during tests
    from scripts.contract_text import is_contract_title as shared_is_contract_title

try:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_COLOR_INDEX
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm
    from docx.shared import Pt
except ImportError as exc:  # pragma: no cover - exercised only without dependency
    raise SystemExit(
        "Missing dependency: python-docx. Install it with: "
        "python3 -m pip install python-docx"
    ) from exc


TITLE_RE = re.compile(r"^劳动合同(?:（.*?）)?$")
MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
CHAPTER_RE = re.compile(r"^(第[一二三四五六七八九十百]+[章节条]|[一二三四五六七八九十]+、).+")
ARTICLE_RE = re.compile(r"^第[一二三四五六七八九十百]+条\s*.+")
MARKDOWN_BOLD_RE = re.compile(r"\*\*(.*?)\*\*")
MARKDOWN_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)")
MARKDOWN_CODE_RE = re.compile(r"`([^`]*)`")
MARKDOWN_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")
PLACEHOLDER_RE = re.compile(
    r"(人民币_{2,}元/月|第_{2,}种|_{2,}年_{2,}月_{2,}日|（\s*）|_{2,})"
)
KEY_CLAUSE_KEYWORDS = (
    "试用期",
    "劳动报酬",
    "保密",
    "知识产权",
    "竞业限制",
    "离职交接",
    "文书送达",
)
EXCLUDED_SECTION_MARKERS = (
    "【已重点处理以下用工风险相关条款】",
    "【合规说明（简版）】",
    "待补充字段",
    "可选条款确认清单",
    "确认清单",
    "请选择导出方式",
    "请选择下一步",
    "DOCX 选择",
    "生成 DOCX",
    "编号式 DOCX 生成选择",
)
EMPTY_FIELD_RE = re.compile(r"(?P<label>[^\s：:]{1,30}[为是]：)\s*(?P<punct>[，。；])")
EMPTY_VALUE_AFTER_COLON_RE = re.compile(r"(?P<label>[^\s：:\n]{1,30}：)\s*(?P<punct>[，。；])")
LABEL_PREFIXES = (
    "甲方（用人单位）",
    "乙方（劳动者）",
    "名称",
    "统一社会信用代码",
    "注册地址",
    "法定代表人",
    "法定代表人或主要负责人",
    "联系电话",
    "姓名",
    "身份证号码",
    "身份证号码/护照号码/其他有效身份证件号码",
    "户籍地址",
    "有效通讯地址",
    "电子邮箱",
    "紧急联系人",
    "紧急联系人关系",
    "紧急联系人电话",
)


def set_run_font(run, size: int = 11, bold: bool = False, highlight=None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    if highlight is not None:
        run.font.highlight_color = highlight
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 6, line: float = 1.25) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_page_number(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9)
    page_run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    page_run._r.append(fld_char_1)
    page_run._r.append(instr_text)
    page_run._r.append(fld_char_2)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9)


def set_document_defaults(document: Document, *, mode: str = "draft") -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(25.4)
    section.left_margin = Mm(25.4)
    section.right_margin = Mm(25.4)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.bold = True

    if mode == "clean":
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(header, after=0, line=1)
        run = header.add_run("劳动合同")
        set_run_font(run, size=9)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(footer, after=0, line=1)
        add_page_number(footer)


def add_text_paragraph(document: Document, text: str, *, bold: bool = False, size: int = 11) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_heading(document: Document, text: str, level: int, *, mode: str = "draft") -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=8, after=6, line=1.2)
    paragraph.style = document.styles[f"Heading {min(level, 3)}"]
    run = paragraph.add_run(text)
    highlight = WD_COLOR_INDEX.BRIGHT_GREEN if mode == "draft" and is_key_clause_heading(text) else None
    set_run_font(run, size=14 if level == 1 else 12, bold=True, highlight=highlight)


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=0, after=14, line=1.2)
    run = paragraph.add_run(text)
    set_run_font(run, size=18, bold=True)


def add_signature_line(document: Document, text: str, *, mode: str = "draft") -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=2, after=4, line=1.2)
    add_segmented_runs(paragraph, text, default_size=11, mode=mode)


def hide_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def add_signature_table(document: Document, left_lines: list[str], right_lines: list[str], *, mode: str = "clean") -> None:
    rows = max(len(left_lines), len(right_lines), 1)
    table = document.add_table(rows=rows, cols=2)
    table.autofit = False
    hide_table_borders(table)

    for cell in table.columns[0].cells:
        cell.width = Mm(78)
    for cell in table.columns[1].cells:
        cell.width = Mm(78)

    for row_index in range(rows):
        for col_index, lines_for_col in enumerate((left_lines, right_lines)):
            text = lines_for_col[row_index] if row_index < len(lines_for_col) else ""
            paragraph = table.cell(row_index, col_index).paragraphs[0]
            set_paragraph_spacing(paragraph, before=4 if row_index else 8, after=12, line=1.2)
            if text:
                add_segmented_runs(paragraph, clean_markdown_inline(text), mode=mode)


def is_signature_line(text: str) -> bool:
    return clean_markdown_inline(text.strip()).startswith(
        (
            "甲方（盖章）",
            "乙方（签字）",
            "法定代表人",
            "授权代表",
            "日期：",
            "签署日期",
        )
    )


def split_signature_block(lines: list[str]) -> tuple[list[str], list[str]]:
    left: list[str] = []
    right: list[str] = []
    date_lines: list[str] = []
    for raw_line in lines:
        line = clean_markdown_inline(raw_line.strip())
        if line.startswith("甲方（盖章）") or line.startswith(("法定代表人", "授权代表")):
            left.append(line)
        elif line.startswith("乙方（签字）"):
            right.append(line)
        elif line.startswith(("日期：", "签署日期")):
            date_lines.append(line)
    if date_lines:
        left.append(date_lines[0])
        right.append(date_lines[-1])
    return left, right


def clean_markdown_inline(text: str) -> str:
    """Remove common inline Markdown markers before writing text to Word."""
    cleaned = text.strip()
    previous = None
    while cleaned != previous:
        previous = cleaned
        cleaned = MARKDOWN_LINK_RE.sub(r"\1", cleaned)
        cleaned = MARKDOWN_CODE_RE.sub(r"\1", cleaned)
        cleaned = MARKDOWN_BOLD_RE.sub(r"\1", cleaned)
        cleaned = MARKDOWN_ITALIC_RE.sub(r"\1", cleaned)
    return cleaned


def is_key_clause_heading(text: str) -> bool:
    return any(keyword in text for keyword in KEY_CLAUSE_KEYWORDS)


def split_label_and_value(text: str) -> tuple[str, str] | None:
    if "：" not in text:
        return None
    label, value = text.split("：", 1)
    label = label.strip()
    value = value.lstrip()
    if not any(label.startswith(prefix) for prefix in LABEL_PREFIXES):
        return None
    return label + "：", value


def add_placeholder_aware_run(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    size: int = 11,
    mode: str = "draft",
) -> None:
    start = 0
    for match in PLACEHOLDER_RE.finditer(text):
        if match.start() > start:
            run = paragraph.add_run(text[start:match.start()])
            set_run_font(run, size=size, bold=bold)
        run = paragraph.add_run(match.group(0))
        highlight = WD_COLOR_INDEX.YELLOW if mode == "draft" else None
        set_run_font(run, size=size, bold=bold, highlight=highlight)
        start = match.end()
    if start < len(text):
        run = paragraph.add_run(text[start:])
        set_run_font(run, size=size, bold=bold)


def add_segmented_runs(paragraph, text: str, *, default_size: int = 11, mode: str = "draft") -> None:
    label_split = split_label_and_value(text)
    if label_split:
        label, value = label_split
        run = paragraph.add_run(label)
        set_run_font(run, size=default_size, bold=True)
        if value:
            add_placeholder_aware_run(paragraph, value, size=default_size, mode=mode)
        return
    add_placeholder_aware_run(paragraph, text, size=default_size, mode=mode)


def normalize_input(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    return [line.rstrip() for line in normalized.split("\n")]


def repair_empty_placeholders(text: str) -> str:
    """Patch common empty-field output into printable placeholders."""

    def replace_empty_field(match: re.Match[str]) -> str:
        return f"{match.group('label')}__________{match.group('punct')}"

    repaired = EMPTY_FIELD_RE.sub(replace_empty_field, text)
    repaired = EMPTY_VALUE_AFTER_COLON_RE.sub(replace_empty_field, repaired)
    return repaired


def normalize_control_line(text: str) -> str:
    line = text.strip()
    markdown_heading = MARKDOWN_HEADING_RE.match(line)
    if markdown_heading:
        line = markdown_heading.group(2)
    return clean_markdown_inline(line)


def extract_contract_text(text: str, *, raw: bool = False) -> str:
    """Extract only the formal contract body from a broader skill response.

    This keeps the finalized contract text while excluding strategy cards,
    compliance notes, pending-field checklists, and DOCX action menus.
    """
    lines = normalize_input(text)
    if not lines:
        return ""
    if raw:
        return repair_empty_placeholders("\n".join(lines).strip())

    start_idx = None
    for idx, raw_line in enumerate(lines):
        if shared_is_contract_title(raw_line):
            start_idx = idx
            break
    if start_idx is None:
        return ""

    trimmed = lines[start_idx:]
    end_idx = len(trimmed)
    for idx, raw_line in enumerate(trimmed):
        line = normalize_control_line(raw_line)
        if any(line.startswith(marker) for marker in EXCLUDED_SECTION_MARKERS):
            end_idx = idx
            break

    contract_lines = trimmed[:end_idx]
    while contract_lines and not contract_lines[-1].strip():
        contract_lines.pop()
    return repair_empty_placeholders("\n".join(contract_lines).strip())


def write_docx(text: str, output_path: Path, *, mode: str = "draft") -> None:
    if mode not in {"draft", "clean"}:
        raise ValueError("mode must be 'draft' or 'clean'")

    document = Document()
    set_document_defaults(document, mode=mode)

    lines = normalize_input(text)
    previous_blank = True

    index = 0
    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.strip()

        if not line:
            previous_blank = True
            index += 1
            continue

        markdown_heading = MARKDOWN_HEADING_RE.match(line)
        if markdown_heading:
            heading_level = len(markdown_heading.group(1))
            heading_text = clean_markdown_inline(markdown_heading.group(2))
            if TITLE_RE.match(heading_text):
                add_title(document, heading_text)
            else:
                add_heading(document, heading_text, min(heading_level, 3), mode=mode)
            previous_blank = False
            index += 1
            continue

        line = clean_markdown_inline(line)

        if TITLE_RE.match(line) and previous_blank:
            add_title(document, line)
        elif CHAPTER_RE.match(line) or ARTICLE_RE.match(line):
            add_heading(document, line, 2, mode=mode)
        elif mode == "clean" and line.startswith("甲方（盖章）"):
            signature_lines = [line]
            scan_index = index + 1
            while scan_index < len(lines):
                candidate = lines[scan_index].strip()
                if not candidate:
                    scan_index += 1
                    continue
                if not is_signature_line(candidate):
                    break
                signature_lines.append(candidate)
                scan_index += 1
            left_lines, right_lines = split_signature_block(signature_lines)
            if left_lines and right_lines:
                add_signature_table(document, left_lines, right_lines, mode=mode)
                index = scan_index - 1
            else:
                add_signature_line(document, line, mode=mode)
        elif line.startswith(("甲方（盖章）", "乙方（签字）", "法定代表人", "日期：")):
            add_signature_line(document, line, mode=mode)
        else:
            paragraph = document.add_paragraph()
            set_paragraph_spacing(paragraph)
            add_segmented_runs(paragraph, line, mode=mode)

        previous_blank = False
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate DOCX from finalized labor contract text.")
    parser.add_argument("--input", "-i", type=Path, help="Path to a UTF-8 text or Markdown contract file.")
    parser.add_argument("--output", "-o", type=Path, required=True, help="Output .docx path.")
    parser.add_argument("--mode", choices=("draft", "clean"), default="draft", help="DOCX output mode.")
    parser.add_argument("--facts", type=Path, help="Path to structured contract facts JSON for validation.")
    parser.add_argument("--raw", action="store_true", help="Treat input as contract body even without a contract title.")
    return parser.parse_args()


def read_utf8_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except OSError as exc:
        raise SystemExit(f"ERROR cannot read input file: {exc}") from None
    except UnicodeDecodeError as exc:
        raise SystemExit(f"ERROR input file must be UTF-8: {exc}") from None


def read_json_file(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise SystemExit(f"ERROR cannot read facts file: {exc}") from None
    except UnicodeDecodeError as exc:
        raise SystemExit(f"ERROR facts file must be UTF-8: {exc}") from None
    except json.JSONDecodeError as exc:
        raise SystemExit(f"ERROR invalid facts JSON: {exc.msg} at line {exc.lineno} column {exc.colno}") from None


def main() -> None:
    args = parse_args()

    if args.input:
        text = read_utf8_text(args.input)
    else:
        text = sys.stdin.read()

    if not text.strip():
        raise SystemExit("No contract text provided. Use --input or pipe text through stdin.")

    contract_text = extract_contract_text(text, raw=args.raw)
    if not contract_text:
        raise SystemExit("No contract body detected after filtering. Provide finalized contract text.")

    from validate_contract import validate_contract

    facts = read_json_file(args.facts) if args.facts else None
    findings = validate_contract(contract_text, facts=facts)
    errors = [finding for finding in findings if finding.severity == "error"]
    if errors:
        for finding in errors:
            print(f"ERROR {finding.code}: {finding.message}", file=sys.stderr)
        raise SystemExit("Contract validation failed; DOCX export blocked.")
    for finding in findings:
        print(f"WARNING {finding.code}: {finding.message}", file=sys.stderr)

    write_docx(contract_text, args.output, mode=args.mode)
    print(f"DOCX generated: {args.output}")


if __name__ == "__main__":
    main()
