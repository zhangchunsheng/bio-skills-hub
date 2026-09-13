# -*- coding: utf-8 -*-
"""
病案首页数据 哈希化去标识工具 v1.5（Skill/CLI 版）
==================================================
从 v1.4 可视化版移植：核心算法与输出规则完全一致，去掉图形界面，改为命令行，
便于接入 AI Agent 技能与脚本自动化。

与 v1.4 GUI 版的差异（发布前安全强化）：
  1. 强制 HMAC 密钥模式：取消「无密钥时退回 SHA256(公开盐+病案号)」的回退逻辑。
     病案号通常是有规律可枚举的（如按年份顺序编号），公开盐值的无密钥哈希
     可以被字典攻击还原——那只是假名化，不是去标识化。
  2. 对照表默认不保留姓名（--keep-name 显式开启，仅病案科留存场景）。
  3. 自检样例中的科室名改为通用名称。

命令：
  python deid_cli.py convert <输入文件或目录> --out-dir <导出目录> [选项]
  python deid_cli.py scan    <输入文件或目录> --id-col 病案号 [--name-col 姓名]
  python deid_cli.py selftest                  # 27 项核心逻辑自检
  python deid_cli.py --doctor                  # 依赖体检

输出规则（与 v1.4 一致）：
  xxx.csv      -> xxx.csv（病案号已替换为研究ID）+ xxx对照表.csv（病案科留存）
  xxx.xlsx     -> xxx.xlsx + xxx对照表.xlsx
  另附 key.bin（密钥，病案科留存）、批量转换报告.txt、总对照表.csv
"""
from __future__ import annotations

import argparse
import csv
import datetime
import glob
import hashlib
import hmac
import json
import os
import re
import sys
import traceback

# ---------- 核心逻辑（纯标准库 + openpyxl） ----------


def normalize_id(raw, strip_zero=True):
    """病案号归一化——保证同一患者在两套系统/两种格式下得到同一研究ID。
    处理：Excel 公式包装 ="0841651"、首尾/全角空格、纯数字前导零、末尾 .0。
    返回归一化后的字符串；空值返回空串"""
    if raw is None:
        return ""
    s = str(raw).strip().strip("\ufeff").replace("\u3000", "").strip()
    if s.startswith('="') and s.endswith('"') and len(s) >= 4:
        s = s[2:-1].strip()
    if s.startswith("'"):
        s = s[1:].strip()
    s = s.strip('"').strip()
    if not s:
        return ""
    if re.match(r"^\d+\.0+$", s):
        s = s.split(".")[0]
    if strip_zero and s.isdigit():
        s2 = s.lstrip("0")
        s = s2 if s2 else "0"
    return s


def is_header_like(value, col_name):
    """识别「伪数据行」：导出文件带标题行导致表头被当成一条记录。"""
    if value is None:
        return False
    v = str(value).strip()
    if not v:
        return False
    return v == str(col_name).strip() or v in ("病案号", "住院号", "患者ID", "研究ID")


def make_pid(raw, key):
    """病案号 -> 研究ID（HMAC-SHA256 截断 16 位十六进制，不可逆；空值返回空串）。
    key 必须提供：无密钥的哈希在病案号可枚举的场景下可被字典攻击还原。"""
    if raw is None:
        return ""
    raw = str(raw).strip()
    if not raw:
        return ""
    if not key:
        raise ValueError("未提供 HMAC 密钥：拒绝在无密钥模式下计算研究ID（可被字典攻击还原）")
    digest = hmac.new(key, raw.encode("utf-8"), hashlib.sha256).hexdigest()[:16]
    return "PID" + digest


def load_or_create_key(out_dir):
    """读取或创建 HMAC 密钥（同一导出目录复用同一密钥，保证跨批次映射稳定）"""
    key_path = os.path.join(out_dir, "key.bin")
    if os.path.exists(key_path):
        with open(key_path, "rb") as f:
            return f.read()
    key = os.urandom(32)
    with open(key_path, "wb") as f:
        f.write(key)
    return key


def prepare_key(out_dir, mode="reuse"):
    """批量转换前的密钥决策。
    mode='reuse' 沿用导出目录已有 key.bin（不存在则新建）→ 增量转换；
    mode='new'   强制新建密钥（旧密钥自动备份为 key.bin.时间戳.bak）。
    返回 (key, reused)"""
    key_path = os.path.join(out_dir, "key.bin")
    existed = os.path.exists(key_path)
    if existed and mode == "new":
        stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
        os.replace(key_path, key_path + "." + stamp + ".bak")
        existed = False
    return load_or_create_key(out_dir), existed


def merge_master_map(out_dir):
    """把导出目录内所有 xxx对照表.csv/.xlsx 合并为 总对照表.csv
    （累计历次批次的唯一患者清单：病案号<->研究ID[，姓名]）。
    返回 (总对照表路径, 累计唯一患者数)"""
    entries = {}
    name_used = False
    files = [p for p in glob.glob(os.path.join(out_dir, "*对照表.*"))
             if not os.path.basename(p).startswith("总对照表")]
    for path in sorted(files):
        try:
            if get_format(path) == "xlsx":
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    rows = wb.active.iter_rows(values_only=True)
                    next(rows, None)
                    for row in rows:
                        if not row or len(row) < 2 or row[0] is None or row[1] is None:
                            continue
                        raw = str(row[0]).strip()
                        nm = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
                        if raw:
                            e = entries.setdefault(raw, [str(row[1]).strip(), ""])
                            if nm:
                                e[1] = nm
                                name_used = True
                finally:
                    wb.close()
            else:
                enc = detect_encoding(path)
                with open(path, "r", encoding=enc, newline="") as f:
                    rd = csv.reader(f)
                    next(rd, None)
                    for row in rd:
                        if len(row) < 2 or not row[0].strip() or not row[1].strip():
                            continue
                        raw = row[0].strip()
                        nm = row[2].strip() if len(row) > 2 else ""
                        e = entries.setdefault(raw, [row[1].strip(), ""])
                        if nm:
                            e[1] = nm
                            name_used = True
        except Exception:
            pass
    out_path = os.path.join(out_dir, "总对照表.csv")
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["病案号", "研究ID"] + (["姓名"] if name_used else []))
        for raw in sorted(entries):
            pid, nm = entries[raw]
            w.writerow([raw, pid] + ([nm] if name_used else []))
    return out_path, len(entries)


def detect_encoding(path):
    """自动检测 CSV 编码：gb18030 > utf-8-sig > gbk > latin-1"""
    for enc in ("gb18030", "utf-8-sig", "gbk", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                f.read(65536)
            return enc
        except (UnicodeDecodeError, OSError):
            continue
    return "gb18030"


def get_format(path):
    """判断文件格式：csv / xlsx（.xls 老格式不支持，会提示转存）"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        return "csv"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    return ext.lstrip(".") or "csv"


def _cell_to_str(v):
    """把 Excel 单元格值统一转成字符串（数字/日期/时间/布尔/None）"""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    return str(v)


def iter_rows(path):
    """统一行迭代器：CSV / XLSX 逐行 yield 字符串列表（第一行为表头）。
    XLSX 用 openpyxl 只读模式流式读取，百万行不占内存。"""
    if get_format(path) == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for row in wb.active.iter_rows(values_only=True):
                yield [c.strip() for c in (_cell_to_str(x) for x in row)]
        finally:
            wb.close()
    else:
        enc = detect_encoding(path)
        with open(path, "r", encoding=enc, newline="") as f:
            for row in csv.reader(f):
                yield [c.strip() for c in row]


def count_rows(path):
    """预估总数据行数（用于进度统计）：XLSX 用工作表维度，CSV 返回 None"""
    if get_format(path) != "xlsx":
        return None
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        mx = wb.active.max_row
        return max(mx - 1, 0) if mx else None
    finally:
        wb.close()


def read_header(path):
    """读取表头，返回列名列表（CSV / XLSX 通用）"""
    for row in iter_rows(path):
        return [h.strip() for h in row]
    return []


def match_column(header, name):
    """按列名在表头中匹配（精确优先，其次包含）"""
    if not name:
        return None
    cleaned = [h.strip() for h in header]
    if name in cleaned:
        return cleaned.index(name)
    for i, h in enumerate(cleaned):
        if name in h:
            return i
    return None


def pick_column(header, candidates):
    """从候选列名中挑出第一个命中的列（精确优先，其次包含）"""
    cleaned = [h.strip() for h in header]
    for cand in candidates:
        for i, name in enumerate(cleaned):
            if name == cand:
                return i
    for cand in candidates:
        for i, name in enumerate(cleaned):
            if cand in name:
                return i
    return None


def output_paths(path, out_dir):
    """按规则生成输出路径：
    研究者版 = 原名.扩展名（若与源文件同路径则自动加「（去标识）」后缀防止覆盖）
    对照表   = 原名对照表.扩展名"""
    base = os.path.splitext(os.path.basename(path))[0]
    ext = os.path.splitext(path)[1].lower()
    if ext not in (".csv", ".xlsx", ".xlsm"):
        ext = ".csv"
    res_path = os.path.join(out_dir, base + ext)
    map_path = os.path.join(out_dir, base + "对照表" + ext)
    if os.path.abspath(res_path) == os.path.abspath(path):
        res_path = os.path.join(out_dir, base + "（去标识）" + ext)
    return res_path, map_path


def resolve_input(args_input):
    """输入可以是文件、目录或通配符：解析为数据文件列表"""
    files = []
    targets = glob.glob(args_input) or [args_input]
    for t in targets:
        if os.path.isdir(t):
            for name in sorted(os.listdir(t)):
                if name.lower().endswith((".csv", ".xlsx", ".xlsm")):
                    files.append(os.path.join(t, name))
        elif os.path.isfile(t):
            fmt = get_format(t)
            if fmt in ("csv", "xlsx"):
                files.append(t)
    # 去重保序
    seen, uniq = set(), []
    for p in files:
        k = os.path.abspath(p)
        if k not in seen:
            seen.add(k)
            uniq.append(p)
    return uniq


NAME_COL_CANDIDATES = ["姓名", "患者姓名", "病人姓名"]


def locate_name_column(header, name_col):
    """定位姓名列：显式指定优先；未指定时自动识别（研究者版必须删除姓名列）。
    返回列索引或 None"""
    if name_col:
        return match_column(header, name_col)
    return pick_column(header, NAME_COL_CANDIDATES)


def batch_scan(files, id_col, name_col=None):
    """批量干跑：所有文件合并统计（不产生任何文件）。
    返回 dict(total, blank, uniq, repeat_dist, same_diff, skipped)"""
    total = 0
    blank = 0
    pid_count = {}
    name_by_pid = {}
    same_diff = 0
    skipped = []
    for path in files:
        try:
            header = read_header(path)
            id_idx = match_column(header, id_col) if id_col else None
            name_idx = locate_name_column(header, name_col)
            if id_idx is None:
                skipped.append({"file": os.path.basename(path), "reason": "未找到病案号列「%s」" % (id_col or "未指定")})
                continue
            first = True
            for row in iter_rows(path):
                if first:
                    first = False
                    continue
                total += 1
                raw = normalize_id(row[id_idx] if id_idx < len(row) and row[id_idx] else "")
                if raw and is_header_like(row[id_idx] if id_idx < len(row) else "", id_col or "病案号"):
                    total -= 1
                    continue
                if not raw:
                    blank += 1
                    continue
                pid_count[raw] = pid_count.get(raw, 0) + 1
                if name_idx is not None and name_idx < len(row):
                    nm = row[name_idx].strip()
                    if nm:
                        prev = name_by_pid.get(raw)
                        if prev is not None and prev != nm:
                            same_diff += 1
                        name_by_pid[raw] = nm
        except Exception as e:
            skipped.append({"file": os.path.basename(path), "reason": str(e)})
    uniq = len(pid_count)
    repeat_dist = {"ge2": 0, "eq2": 0, "eq3": 0, "ge4": 0}
    for c in pid_count.values():
        if c >= 2:
            repeat_dist["ge2"] += 1
        if c == 2:
            repeat_dist["eq2"] += 1
        elif c == 3:
            repeat_dist["eq3"] += 1
        elif c >= 4:
            repeat_dist["ge4"] += 1
    return {"total": total, "blank": blank, "uniq": uniq, "repeat_dist": repeat_dist,
            "same_diff": same_diff, "skipped": skipped}


def convert_one(path, out_dir, id_col, name_col=None, keep_name=False, key=None):
    """单文件转换：输出 原名.ext（研究者版）+ 原名对照表.ext（病案科留存）。
    key 必须由 batch_convert 统一传入（或自行读取导出目录密钥）。
    返回结果 dict；若该文件找不到病案号列则返回 {"skipped": 原因}"""
    os.makedirs(out_dir, exist_ok=True)
    if key is None:
        key = load_or_create_key(out_dir)

    header = read_header(path)
    if not header:
        return {"path": path, "skipped": "文件为空或没有表头"}
    id_idx = match_column(header, id_col) if id_col else None
    if id_idx is None:
        return {"path": path, "skipped": "未找到病案号列「%s」" % (id_col or "未指定")}
    name_idx = locate_name_column(header, name_col)

    id_name = header[id_idx]
    name_name = header[name_idx] if name_idx is not None else None

    res_fields = []
    for i, c in enumerate(header):
        if i == id_idx:
            res_fields.append("研究ID")
        elif i == name_idx:
            continue
        else:
            res_fields.append(c)

    res_path, map_path = output_paths(path, out_dir)
    total_est = count_rows(path)

    total = 0
    blank = 0
    uniq = {}
    fmt = get_format(path)
    is_xlsx = (fmt == "xlsx")

    if is_xlsx:
        from openpyxl import Workbook
        wb_res = Workbook(write_only=True)
        ws_res = wb_res.create_sheet()
        ws_res.append(res_fields)
        wb_map = Workbook(write_only=True)
        ws_map = wb_map.create_sheet()
        ws_map.append(["病案号", "研究ID"] + ([name_name] if keep_name and name_name else []))
    else:
        enc = detect_encoding(path)
        fres = open(res_path, "w", encoding=enc, newline="")
        fmap = open(map_path, "w", encoding=enc, newline="")
        w_res = csv.writer(fres)
        w_map = csv.writer(fmap)
        w_res.writerow(res_fields)
        w_map.writerow(["病案号", "研究ID"] + ([name_name] if keep_name and name_name else []))

    first = True
    header_skips = 0
    for row in iter_rows(path):
        if first:
            first = False
            continue
        total += 1
        raw = row[id_idx].strip() if id_idx < len(row) and row[id_idx] else ""
        raw = normalize_id(raw)
        if is_header_like(row[id_idx] if id_idx < len(row) else "", id_name):
            header_skips += 1
            total -= 1
            continue
        pid = make_pid(raw, key)
        if not pid:
            blank += 1
        else:
            uniq[pid] = uniq.get(pid, 0) + 1
        out = []
        for i, c in enumerate(header):
            val = row[i] if i < len(row) else ""
            if i == id_idx:
                out.append(pid)
            elif i == name_idx:
                continue
            else:
                out.append(val)
        if is_xlsx:
            ws_res.append(out)
        else:
            w_res.writerow(out)
        if pid:
            map_row = [raw, pid]
            if keep_name and name_name and name_idx < len(row):
                map_row.append(row[name_idx].strip())
            if is_xlsx:
                ws_map.append(map_row)
            else:
                w_map.writerow(map_row)

    if is_xlsx:
        wb_res.save(res_path)
        wb_map.save(map_path)
    else:
        fres.close()
        fmap.close()

    ge2 = sum(1 for c in uniq.values() if c >= 2)
    return {"path": path, "ok": True,
            "total": total, "blank": blank, "uniq": len(uniq), "ge2": ge2,
            "res": res_path, "map": map_path, "header_skips": header_skips}


def batch_convert(files, out_dir, id_col, name_col=None, keep_name=False, key_mode="reuse"):
    """批量转换：所有文件共用同一密钥，跨文件同一患者研究ID一致。
    key_mode='reuse'（默认）沿用 key.bin → 增量转换；'new' 新建（旧密钥备份）。
    完成后合并总对照表、写批量转换报告.txt。返回汇总 dict。"""
    os.makedirs(out_dir, exist_ok=True)
    key, key_reused = prepare_key(out_dir, key_mode)
    results = []
    for path in files:
        results.append(convert_one(path, out_dir, id_col, name_col, keep_name, key=key))

    ok_list = [r for r in results if r.get("ok")]
    skip_list = [r for r in results if not r.get("ok")]
    t_total = sum(r["total"] for r in ok_list)
    t_blank = sum(r["blank"] for r in ok_list)
    t_uniq_all = {}
    for r in ok_list:
        try:
            if get_format(r["map"]) == "xlsx":
                from openpyxl import load_workbook
                wb = load_workbook(r["map"], read_only=True, data_only=True)
                try:
                    rows = wb.active.iter_rows(values_only=True)
                    next(rows, None)
                    for row in rows:
                        if row and row[1]:
                            pid = str(row[1])
                            t_uniq_all[pid] = t_uniq_all.get(pid, 0) + 1
                finally:
                    wb.close()
            else:
                enc = detect_encoding(r["map"])
                with open(r["map"], "r", encoding=enc, newline="") as f:
                    rd = csv.reader(f)
                    next(rd, None)
                    for row in rd:
                        if len(row) >= 2 and row[1]:
                            pid = row[1].strip()
                            t_uniq_all[pid] = t_uniq_all.get(pid, 0) + 1
        except Exception:
            pass
    t_uniq = len(t_uniq_all)
    t_ge2 = sum(1 for c in t_uniq_all.values() if c >= 2)
    header_skips_total = sum(r.get("header_skips", 0) for r in ok_list)

    try:
        master_path, master_n = merge_master_map(out_dir)
    except Exception:
        master_path, master_n = None, 0

    lines = []
    A = lines.append
    A("=" * 52)
    A("  病案首页数据 哈希化去标识 批量转换报告（v1.5 CLI）")
    A("=" * 52)
    A("处理时间    : " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    A("导出文件夹  : " + out_dir)
    if key_reused:
        A("密钥        : 沿用已有 key.bin ← 增量转换：老患者研究ID与之前批次一致")
    else:
        A("密钥        : 本次新生成 key.bin（本批所有文件共用）")
    A("病案号归一化: 已开启（去Excel公式包装=\"\"、空格、前导零、数字.0尾巴）")
    A("病案号列    : " + (id_col or "未指定"))
    A("姓名列      : " + (name_col or "（无）"))
    A("对照表保留姓名: " + ("是（仅病案科留存）" if keep_name else "否"))
    A("-" * 52)
    A("文件处理明细：")
    for i, r in enumerate(ok_list, 1):
        A("  [%d] 成功  %s" % (i, os.path.basename(r["path"])))
        A("       总记录 %d | 唯一患者 %d | 住院>=2次 %d | 输出 %s + %s"
          % (r["total"], r["uniq"], r["ge2"], os.path.basename(r["res"]), os.path.basename(r["map"])))
    for i, r in enumerate(skip_list, 1):
        A("  [跳过] %s（%s）" % (os.path.basename(r.get("path", "?")), r.get("skipped", "")))
    A("-" * 52)
    A("汇总统计（跨文件合并）：")
    A("  文件数    : %d（成功 %d / 跳过 %d）" % (len(files), len(ok_list), len(skip_list)))
    A("  总记录数  : %d" % t_total)
    A("  空病案号行: %d" % t_blank)
    A("  唯一患者数: %d" % t_uniq)
    A("  住院>=2次 : %d  ← 这个数越大，纵向研究越可行" % t_ge2)
    A("-" * 52)
    A("重要提示：")
    A("1. 对照表、总对照表与 key.bin 必须由病案科保管，严禁随数据外发。")
    A("2. 各文件共用同一密钥，同一患者在不同文件中的研究ID相同，")
    A("   可直接拼接为纵向队列（再入院/年急性加重频次等研究）。")
    A("3. 以后新增批次的数据：在同一导出目录、密钥沿用（默认）再跑一次即可，")
    A("   老患者ID不变，新患者自动分配新ID，无需重跑旧文件。")
    A("4. 若某文件被「跳过」，说明其中没有匹配到病案号列，请检查后单独处理。")
    A("=" * 52)
    rep_path = os.path.join(out_dir, "批量转换报告.txt")
    with open(rep_path, "w", encoding="utf-8-sig") as f:
        f.write("\n".join(lines))
    return {"results": results, "rep": rep_path, "ok_count": len(ok_list),
            "skip_count": len(skip_list), "total": t_total, "blank": t_blank,
            "uniq": t_uniq, "ge2": t_ge2, "key_reused": key_reused,
            "master": master_path, "master_n": master_n,
            "header_skips": header_skips_total}


# ---------- 自检（27 项，与 v1.4 GUI 版同源，科室名已通用化） ----------

def self_test():
    import tempfile
    lines = []
    ok = True
    try:
        tmp = tempfile.mkdtemp()
        key = b"test-key-123"
        a = make_pid("123456", key)
        b = make_pid("123456", key)
        c = make_pid("789012", key)
        lines.append("[1] 同一病案号->同一ID: %s" % ("OK" if a == b else "FAIL"))
        lines.append("[2] 不同病案号->不同ID: %s" % ("OK" if a != c else "FAIL"))
        lines.append("[3] 空值->空ID: %s" % ("OK" if make_pid("", key) == "" else "FAIL"))
        lines.append("[4] ID前缀PID: %s (%s)" % ("OK" if a.startswith("PID") else "FAIL", a))
        try:
            make_pid("123456", None)
            lines.append("[4b] 无密钥拒绝计算: FAIL")
        except ValueError:
            lines.append("[4b] 无密钥拒绝计算: OK")
        # 5. 转换流程（模拟2个患者3次住院 + 1条空）
        sample = os.path.join(tmp, "sample.csv")
        with open(sample, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["病案号", "姓名", "出院科室", "总费用"])
            w.writerow(["10001", "张一", "呼吸内科一病区", "8530"])
            w.writerow(["10001", "张一", "呼吸内科一病区", "11200"])
            w.writerow(["10002", "李二", "呼吸内科二病区", "5600"])
            w.writerow(["", "缺失", "呼吸内科一病区", "3000"])
        enc = detect_encoding(sample)
        header = read_header(sample)
        id_idx = pick_column(header, ["病案号", "住院号", "患者ID"])
        name_idx = pick_column(header, ["姓名"])
        lines.append("[5] 列识别 病案号=%s 姓名=%s: %s" % (
            header[id_idx] if id_idx is not None else None,
            header[name_idx] if name_idx is not None else None,
            "OK" if id_idx == 0 and name_idx == 1 else "FAIL"))
        out_dir = os.path.join(tmp, "out")
        res = convert_one(sample, out_dir, "病案号", "姓名", keep_name=True, key=key)
        with open(res["res"], encoding="gb18030") as f:
            rows = list(csv.DictReader(f))
        ids = [r["研究ID"] for r in rows]
        lines.append("[6] 研究ID列写入: %s" % ("OK" if ids[0] and ids[0].startswith("PID") else "FAIL"))
        lines.append("[7] 同一患者跨次住院可关联: %s" % ("OK" if ids[0] == ids[1] else "FAIL"))
        lines.append("[8] 不同患者ID不同: %s" % ("OK" if ids[0] != ids[2] else "FAIL"))
        lines.append("[9] 姓名列已删除: %s" % ("OK" if "姓名" not in rows[0] else "FAIL"))
        lines.append("[10] 空病案号行ID为空: %s" % ("OK" if ids[3] == "" else "FAIL"))
        lines.append("[11] 映射表生成: %s" % ("OK" if os.path.exists(res["map"]) else "FAIL"))
        # key.bin 由批量模式的 prepare_key 统一落盘（单文件直传密钥时不落盘）
        kd = os.path.join(tmp, "outk")
        batch_convert([sample], kd, "病案号", "姓名", keep_name=True, key_mode="new")
        lines.append("[12] 密钥文件存在(批量模式): %s" % ("OK" if os.path.exists(os.path.join(kd, "key.bin")) else "FAIL"))
        lines.append("[13] 文件名规则 研究者版=%s 对照表=%s: %s"
                     % (os.path.basename(res["res"]), os.path.basename(res["map"]),
                        "OK" if os.path.basename(res["res"]) == "sample.csv"
                              and os.path.basename(res["map"]) == "sample对照表.csv" else "FAIL"))
        st = batch_scan([sample], "病案号", "姓名")
        lines.append("[14] 干跑统计 total=%d uniq=%d ge2=%d: %s" % (
            st["total"], st["uniq"], st["repeat_dist"]["ge2"],
            "OK" if st["total"] == 4 and st["uniq"] == 2 and st["repeat_dist"]["ge2"] == 1 else "FAIL"))
        # 15-16. XLSX 支持
        import openpyxl
        from datetime import date as _date
        xlsx_path = os.path.join(tmp, "sample.xlsx")
        wbx = openpyxl.Workbook()
        wsx = wbx.active
        wsx.append(["病案号", "姓名", "出院科室", "总费用", "出院日期"])
        wsx.append([10001, "张一", "呼吸内科一病区", 8530, _date(2022, 3, 8)])
        wsx.append([10001, "张一", "呼吸内科一病区", 11200, _date(2022, 7, 22)])
        wsx.append([20002, "李二", "呼吸内科二病区", 5600.5, _date(2022, 4, 10)])
        wsx.append([None, "缺失", "呼吸内科一病区", 3000, _date(2022, 6, 2)])
        wbx.save(xlsx_path)
        hx = read_header(xlsx_path)
        ix = pick_column(hx, ["病案号"])
        nx = pick_column(hx, ["姓名"])
        sx = batch_scan([xlsx_path], "病案号", "姓名")
        rowsx_raw = list(iter_rows(xlsx_path))
        ok15 = (hx[0] == "病案号" and sx["total"] == 4 and sx["uniq"] == 2
                and sx["repeat_dist"]["ge2"] == 1 and rowsx_raw[1][0] == "10001"
                and rowsx_raw[1][4] == "2022-03-08")
        lines.append("[15] XLSX 读取 表头/数字病案号/日期转换/统计 total=%d uniq=%d: %s"
                     % (sx["total"], sx["uniq"], "OK" if ok15 else "FAIL"))
        ox = os.path.join(tmp, "outx")
        rx = convert_one(xlsx_path, ox, "病案号", "姓名", keep_name=True, key=key)
        lines.append("[15b] XLSX 输出仍为xlsx 研究=%s 对照=%s: %s"
                     % (os.path.basename(rx["res"]), os.path.basename(rx["map"]),
                        "OK" if os.path.basename(rx["res"]) == "sample.xlsx"
                              and os.path.basename(rx["map"]) == "sample对照表.xlsx" else "FAIL"))
        with open(rx["res"], "rb") as f:
            head = f.read(4)
        import zipfile
        ok16 = zipfile.is_zipfile(rx["res"]) and head[:2] == b"PK"
        lines.append("[16] XLSX 转换 输出为合法xlsx: %s" % ("OK" if ok16 else "FAIL"))
        # 17. 批量：两个文件（不同列顺序），同一患者跨文件ID一致
        batch_dir = os.path.join(tmp, "batch")
        os.makedirs(batch_dir)
        f1 = os.path.join(batch_dir, "2022.1.csv")
        with open(f1, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["病案号", "姓名", "出院科室"])
            w.writerow(["10001", "张一", "呼吸内科一病区"])
            w.writerow(["10003", "王五", "呼吸内科一病区"])
        f2 = os.path.join(batch_dir, "2022.2.csv")
        with open(f2, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["姓名", "出院科室", "病案号"])
            w.writerow(["张一", "呼吸内科二病区", "10001"])
            w.writerow(["李二", "呼吸内科一病区", "10002"])
        outb = os.path.join(tmp, "outb")
        rb = batch_convert([f1, f2], outb, "病案号", "姓名", keep_name=True, key_mode="new")
        lines.append("[17] 批量 2文件成功数=%d 总记录=%d: %s"
                     % (rb["ok_count"], rb["total"], "OK" if rb["ok_count"] == 2 and rb["total"] == 4 else "FAIL"))
        map_files = sorted(p for p in glob.glob(os.path.join(outb, "*对照表.csv"))
                           if not os.path.basename(p).startswith("总对照表"))
        id_by_file = {}
        for p in map_files:
            with open(p, encoding="gb18030") as f:
                rd = csv.DictReader(f)
                for row in rd:
                    id_by_file.setdefault(os.path.basename(p), {})[row["病案号"]] = row["研究ID"]
        ok17b = (len(map_files) == 2
                 and id_by_file.get("2022.1对照表.csv", {}).get("10001")
                     == id_by_file.get("2022.2对照表.csv", {}).get("10001")
                 and id_by_file.get("2022.1对照表.csv", {}).get("10003")
                 and id_by_file.get("2022.2对照表.csv", {}).get("10002"))
        lines.append("[17b] 跨文件同一患者ID一致(按对照表验证): %s" % ("OK" if ok17b else "FAIL"))
        # 18. 防覆盖
        same_dir = os.path.join(tmp, "same")
        os.makedirs(same_dir)
        fs = os.path.join(same_dir, "demo.csv")
        with open(fs, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["病案号", "出院科室"])
            w.writerow(["10001", "呼吸内科一病区"])
        res_s = convert_one(fs, same_dir, "病案号", None, keep_name=False, key=key)
        ok18 = os.path.basename(res_s["res"]) == "demo（去标识）.csv" and os.path.exists(fs)
        lines.append("[18] 防覆盖 同目录自动加后缀且源文件保留: %s" % ("OK" if ok18 else "FAIL"))
        # 19. 列不匹配跳过
        f3 = os.path.join(tmp, "no_id.csv")
        with open(f3, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["患者姓名", "出院科室"])
            w.writerow(["张三", "呼吸内科一病区"])
        r3 = convert_one(f3, outb, "病案号", "姓名", keep_name=True, key=key)
        ok19 = r3.get("skipped") and not r3.get("ok")
        lines.append("[19] 缺病案号列文件被跳过并说明原因: %s" % ("OK" if ok19 else "FAIL"))
        lines.append("[20] 批量转换报告.txt 生成: %s" % ("OK" if os.path.exists(os.path.join(outb, "批量转换报告.txt")) else "FAIL"))
        # 21-22. 增量转换
        out_inc = os.path.join(tmp, "out_inc")
        f25 = os.path.join(tmp, "2025.1.csv")
        with open(f25, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["病案号", "姓名", "出院科室"])
            w.writerow(["10001", "张一", "呼吸内科一病区"])
            w.writerow(["10002", "李二", "呼吸内科二病区"])
        rb1 = batch_convert([f25], out_inc, "病案号", "姓名", keep_name=True, key_mode="new")
        f26 = os.path.join(tmp, "2026.1.csv")
        with open(f26, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["病案号", "姓名", "出院科室"])
            w.writerow(["10002", "李二", "呼吸内科一病区"])
            w.writerow(["20003", "王三", "呼吸内科二病区"])
        rb2 = batch_convert([f26], out_inc, "病案号", "姓名", keep_name=True, key_mode="reuse")
        id_map = {}
        for mf in ("2025.1对照表.csv", "2026.1对照表.csv"):
            with open(os.path.join(out_inc, mf), encoding="gb18030") as f:
                for row in csv.DictReader(f):
                    id_map[(mf, row["病案号"])] = row["研究ID"]
        ok21 = (rb1.get("key_reused") is False and rb2.get("key_reused") is True
                and id_map[("2025.1对照表.csv", "10002")] == id_map[("2026.1对照表.csv", "10002")]
                and id_map[("2025.1对照表.csv", "10001")] != id_map[("2026.1对照表.csv", "20003")])
        lines.append("[21] 增量转换 沿用密钥后老患者ID一致/新患者新ID/key_reused标志: %s"
                     % ("OK" if ok21 else "FAIL"))
        master_path = os.path.join(out_inc, "总对照表.csv")
        mid = {}
        with open(master_path, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                mid[row["病案号"]] = row["研究ID"]
        ok22 = (os.path.exists(master_path) and len(mid) == 3 and rb2["master_n"] == 3
                and mid["10002"] == id_map[("2025.1对照表.csv", "10002")])
        lines.append("[22] 总对照表 累计唯一病案号=%d 且ID与各批一致: %s" % (len(mid), "OK" if ok22 else "FAIL"))
        # 23. 强制新密钥
        old_key = open(os.path.join(out_inc, "key.bin"), "rb").read()
        new_key, reused = prepare_key(out_inc, "new")
        bak = glob.glob(os.path.join(out_inc, "key.bin.*.bak"))
        ok23 = (len(bak) >= 1 and reused is False and new_key != old_key
                and open(bak[0], "rb").read() == old_key)
        lines.append("[23] 强制新密钥 旧密钥备份.bak且新旧不同: %s" % ("OK" if ok23 else "FAIL"))
        # 24-26. 病案号归一化
        lines.append("[24] 归一化 去Excel包装: %s" % ("OK" if normalize_id('="0841651"') == "841651" else "FAIL"))
        lines.append("[25] 归一化 纯数字/前导零/.0/空格: %s" % (
            "OK" if normalize_id("1331855") == "1331855"
                 and normalize_id(1331855.0) == "1331855"
                 and normalize_id("  01331855 ") == "1331855" else "FAIL"))
        kk = b"normalize-test-key"
        id_old = make_pid(normalize_id('="0841651"'), kk)
        id_new = make_pid(normalize_id("841651"), kk)
        id_raw = make_pid('="0841651"', kk)
        lines.append("[26] 跨期关联 老格式(=\"0841651\")与新格式(841651)同ID: %s (未归一化时ID不同: %s)"
                     % ("OK" if id_old == id_new else "FAIL", "是" if id_raw != id_new else "否"))
        # 27. 伪数据行
        fhdr = os.path.join(tmp, "hdr.csv")
        with open(fhdr, "w", encoding="gb18030", newline="") as f:
            w = csv.writer(f)
            w.writerow(["病案号", "出院科室"])
            w.writerow(["病案号", "呼吸内科一病区"])
            w.writerow(["10001", "呼吸内科一病区"])
            w.writerow(["10002", "呼吸内科二病区"])
        rh = convert_one(fhdr, os.path.join(tmp, "outh"), "病案号", None, keep_name=False, key=key)
        lines.append("[27] 伪数据行自动跳过 有效记录=%d(应为2) 跳过=%d(应为1): %s"
                     % (rh["total"], rh.get("header_skips", 0),
                        "OK" if rh["total"] == 2 and rh.get("header_skips", 0) == 1 else "FAIL"))
    except Exception:
        ok = False
        lines.append("[异常] %s" % traceback.format_exc())
    result = "\n".join(lines)
    print(result)
    return 0 if ok and "FAIL" not in result else 1


# ---------- 命令行入口 ----------


def _emit(obj):
    print(json.dumps(obj, ensure_ascii=False))


def cmd_scan(args):
    files = resolve_input(args.input)
    if not files:
        _emit({"success": False, "error": {"code": "no_input_files",
               "message": "输入不是可处理的 csv/xlsx 文件或目录", "suggestion": "检查路径与文件扩展名"}})
        return 1
    r = batch_scan(files, args.id_col, args.name_col)
    ratio = round(r["repeat_dist"]["ge2"] * 100 / r["uniq"], 1) if r["uniq"] else 0.0
    _emit({"success": True, "mode": "scan", "files": len(files),
           "total": r["total"], "blank": r["blank"], "uniq": r["uniq"],
           "ge2": r["repeat_dist"]["ge2"], "repeat_dist": r["repeat_dist"],
           "same_diff": r["same_diff"], "ge2_ratio_pct": ratio,
           "skipped": r["skipped"],
           "note": "干跑不产生任何文件。住院≥2次患者占比 <5% 时，病案号可能=每次住院唯一（住院号），"
                   "建议改选「患者ID/身份证号」列重新干跑。"})
    return 0


def cmd_convert(args):
    files = resolve_input(args.input)
    if not files:
        _emit({"success": False, "error": {"code": "no_input_files",
               "message": "输入不是可处理的 csv/xlsx 文件或目录", "suggestion": "检查路径与文件扩展名"}})
        return 1
    if not args.out_dir:
        _emit({"success": False, "error": {"code": "no_out_dir",
               "message": "缺少 --out-dir 导出目录", "suggestion": "所有结果（含 key.bin）都会放进导出目录，必须显式指定"}})
        return 1
    rb = batch_convert(files, args.out_dir, args.id_col, args.name_col,
                       keep_name=args.keep_name, key_mode=args.key_mode)
    _emit({"success": True, "mode": "convert", "outdir": os.path.abspath(args.out_dir),
           "files": len(files), "ok_count": rb["ok_count"], "skip_count": rb["skip_count"],
           "total": rb["total"], "blank": rb["blank"], "uniq": rb["uniq"], "ge2": rb["ge2"],
           "key_reused": rb["key_reused"], "master": rb["master"], "master_n": rb["master_n"],
           "report": rb["rep"],
           "skipped": [{"file": os.path.basename(r.get("path", "?")), "reason": r.get("skipped", "")}
                       for r in rb["results"] if not r.get("ok")],
           "security": "对照表、总对照表与 key.bin 必须由病案科保管，严禁随数据外发"})
    return 0


def cmd_doctor():
    info = {"success": True, "mode": "doctor",
            "python": sys.version.split()[0], "platform": sys.platform}
    try:
        import openpyxl
        info["openpyxl"] = openpyxl.__version__
    except Exception:
        info["openpyxl"] = "缺失（处理 .xlsx 需要，pip install openpyxl）"
    _emit(info)
    return 0


def main():
    p = argparse.ArgumentParser(description="病案首页数据 哈希化去标识工具 v1.5（CLI 版）")
    p.add_argument("mode", nargs="?", default=None,
                   help="convert（批量转换）/ scan（干跑统计）/ selftest（自检）；--doctor 体检")
    p.add_argument("input", nargs="?", default=None, help="输入文件、目录或通配符")
    p.add_argument("--out-dir", default=None, help="导出目录（convert 必填，key.bin 与全部产物在此）")
    p.add_argument("--id-col", default="病案号", help="病案号列名（默认 病案号，支持精确/包含匹配）")
    p.add_argument("--name-col", default=None, help="姓名列名（默认自动识别；对照表默认不含姓名）")
    p.add_argument("--keep-name", action="store_true",
                   help="对照表中保留姓名（仅病案科留存场景；研究者版始终不含姓名）")
    p.add_argument("--key-mode", choices=["reuse", "new"], default="reuse",
                   help="reuse=沿用导出目录已有 key.bin（默认，增量转换）；new=新建密钥（旧密钥备份 .bak）")
    p.add_argument("--doctor", action="store_true", help="依赖体检")
    args = p.parse_args()

    if args.doctor or (args.mode is None and not args.input):
        return cmd_doctor()
    if args.mode == "selftest":
        return self_test()
    if args.mode == "scan":
        return cmd_scan(args)
    if args.mode == "convert":
        return cmd_convert(args)
    _emit({"success": False, "error": {"code": "unknown_mode",
           "message": "mode 应为 convert / scan / selftest / --doctor", "suggestion": "见脚本头部用法说明"}})
    return 1


if __name__ == "__main__":
    sys.exit(main())
