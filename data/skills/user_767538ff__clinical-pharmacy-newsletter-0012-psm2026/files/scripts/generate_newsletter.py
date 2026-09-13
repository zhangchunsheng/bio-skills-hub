#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
月度临床药讯Word文档生成器
输入：JSON内容文件（schema见 references/content_schema.md）
输出：带封面、目录、栏目分节、页眉页脚的.docx文件
"""
import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


# ===== 字体常量 =====
FONT_BODY_CN = "宋体"
FONT_TITLE_CN = "黑体"
FONT_EN = "Times New Roman"

# ===== 颜色 =====
COLOR_PRIMARY = RGBColor(0x1F, 0x4E, 0x79)  # 深蓝
COLOR_ACCENT = RGBColor(0xC0, 0x00, 0x00)   # 深红（警示）
COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)     # 正文深灰


def set_run_font(run, font_cn=FONT_BODY_CN, font_en=FONT_EN,
                 size=10.5, bold=False, color=None):
    """设置run字体"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font_en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_cn)
    if color is not None:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1, color=COLOR_PRIMARY):
    """添加自定义样式标题"""
    p = doc.add_paragraph()
    p.style = doc.styles[f"Heading {level}"]
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, font_cn=FONT_TITLE_CN, size=16, bold=True, color=color)
    elif level == 2:
        set_run_font(run, font_cn=FONT_TITLE_CN, size=14, bold=True, color=color)
    else:
        set_run_font(run, font_cn=FONT_TITLE_CN, size=12, bold=True, color=color)
    return p


def add_body_paragraph(doc, text, bold=False, color=None, size=10.5):
    """添加正文段落"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)  # 首行缩进2字符
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_label_value(doc, label, value, label_color=COLOR_PRIMARY):
    """添加'标签：值'格式段落"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    run_label = p.add_run(f"【{label}】")
    set_run_font(run_label, font_cn=FONT_TITLE_CN, size=10.5, bold=True,
                 color=label_color)
    run_value = p.add_run(value)
    set_run_font(run_value, size=10.5)
    return p


def setup_page(doc):
    """设置页面：A4，页边距"""
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)


def add_cover_page(doc, issue):
    """封面页"""
    # 空行撑开
    for _ in range(6):
        doc.add_paragraph()

    # 主标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("临床药讯")
    set_run_font(run, font_cn=FONT_TITLE_CN, size=36, bold=True,
                 color=COLOR_PRIMARY)

    # 期号
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{issue['year']}年第{issue['month']}期")
    set_run_font(run, font_cn=FONT_TITLE_CN, size=20, bold=True,
                 color=COLOR_TEXT)

    # 专题（可选）
    topic = issue.get("topic")
    if topic:
        for _ in range(2):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"本期专题：{topic}")
        set_run_font(run, font_cn=FONT_TITLE_CN, size=16, bold=True,
                     color=COLOR_ACCENT)

    # 底部单位
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("药学部  临床药学组")
    set_run_font(run, font_cn=FONT_TITLE_CN, size=14, bold=True,
                 color=COLOR_TEXT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{issue['year']}年{issue['month']:02d}月")
    set_run_font(run, font_cn=FONT_BODY_CN, size=12, color=COLOR_TEXT)

    doc.add_page_break()


def add_toc(doc, sections):
    """目录页（手动生成，非Word自动目录）"""
    add_heading_styled(doc, "本期目录", level=1)
    doc.add_paragraph()

    toc_items = [
        ("一、卷首语", "preface"),
        ("二、指南更新", "guideline_updates"),
        ("三、药物安全警示", "safety_alerts"),
        ("四、文献速递", "literature"),
        ("五、新药速览", "new_drugs"),
    ]
    for title, key in toc_items:
        if key in sections and sections[key]:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.8
            run = p.add_run(title)
            set_run_font(run, font_cn=FONT_TITLE_CN, size=12,
                         color=COLOR_TEXT)

    doc.add_page_break()


def add_header_footer(doc, issue):
    """页眉页脚"""
    section = doc.sections[0]
    # 页眉
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_p.add_run(f"临床药讯 {issue['year']}年第{issue['month']}期")
    set_run_font(run, font_cn=FONT_BODY_CN, size=9, color=COLOR_TEXT)
    # 页眉下划线
    pPr = header_p._element.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single", qn("w:sz"): "6",
        qn("w:space"): "1", qn("w:color"): "1F4E79"
    })
    pBdr.append(bottom)
    pPr.append(pBdr)

    # 页脚（页码）
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    set_run_font(run, size=9, color=COLOR_TEXT)
    # 插入页码字段
    fldChar1 = run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    run._element.append(fldChar1)
    instrText = run._element.makeelement(qn("w:instrText"), {})
    instrText.text = "PAGE"
    run._element.append(instrText)
    fldChar2 = run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    run._element.append(fldChar2)


# ===== 栏目渲染 =====

def render_preface(doc, text):
    """卷首语"""
    add_heading_styled(doc, "一、卷首语", level=1)
    add_body_paragraph(doc, text)


def render_guideline_updates(doc, items):
    """指南更新"""
    add_heading_styled(doc, "二、指南更新", level=1)
    for i, item in enumerate(items, 1):
        add_heading_styled(doc, f"{i}. {item['title']}", level=2)
        add_label_value(doc, "来源", f"{item.get('source', '—')}（{item.get('publish_date', '—')}）")
        add_label_value(doc, "关键变更", item.get("changes", "—"))
        add_label_value(doc, "临床影响", item.get("impact", "—"))
        add_label_value(doc, "证据等级", item.get("evidence_level", "—"))
        if item.get("translation_note"):
            add_label_value(doc, "翻译说明", item["translation_note"],
                            label_color=COLOR_ACCENT)
        doc.add_paragraph()


def render_safety_alerts(doc, items):
    """药物安全警示"""
    add_heading_styled(doc, "三、药物安全警示", level=1)
    for i, item in enumerate(items, 1):
        add_heading_styled(doc, f"{i}. {item['title']}", level=2)
        add_label_value(doc, "类型", item.get("type", "—"))
        add_label_value(doc, "涉及药品", item.get("drug", "—"))
        if item.get("case"):
            add_label_value(doc, "案例", item["case"])
        if item.get("national_alert"):
            add_label_value(doc, "监管通报", item["national_alert"])
        add_label_value(doc, "风险机制", item.get("risk", "—"))
        add_label_value(doc, "临床建议", item.get("suggestion", "—"))
        if item.get("report_deadline"):
            add_label_value(doc, "上报时限", item["report_deadline"],
                            label_color=COLOR_ACCENT)
        doc.add_paragraph()


def render_literature(doc, items):
    """文献速递"""
    add_heading_styled(doc, "四、文献速递", level=1)
    for i, item in enumerate(items, 1):
        add_heading_styled(doc, f"{i}. {item['title']}", level=2)
        add_label_value(doc, "期刊", f"{item.get('journal', '—')}（{item.get('publish_date', '—')}）")
        add_label_value(doc, "研究类型", item.get("study_type", "—"))
        add_label_value(doc, "标识", f"PMID: {item.get('pmid', '—')}  DOI: {item.get('doi', '—')}")
        add_label_value(doc, "摘要", item.get("summary", "—"))
        add_label_value(doc, "临床启示", item.get("clinical_implication", "—"))
        add_label_value(doc, "证据等级", item.get("evidence_level", "—"))
        doc.add_paragraph()


def render_new_drugs(doc, items):
    """新药速览"""
    add_heading_styled(doc, "五、新药速览", level=1)
    for i, item in enumerate(items, 1):
        add_heading_styled(doc, f"{i}. {item['name']}", level=2)
        add_label_value(doc, "分类", item.get("class", "—"))
        add_label_value(doc, "适应症", item.get("indication", "—"))
        add_label_value(doc, "用法用量", item.get("dosage", "—"))
        add_label_value(doc, "注意事项", item.get("cautions", "—"))
        add_label_value(doc, "医保属性", item.get("insurance", "—"))
        add_label_value(doc, "药事会状态", item.get("formulary_status", "—"))
        if item.get("key_interaction"):
            add_label_value(doc, "重要相互作用", item["key_interaction"],
                            label_color=COLOR_ACCENT)
        doc.add_paragraph()


def build_document(content):
    """构建完整文档"""
    doc = Document()

    # 设置默认样式
    style = doc.styles["Normal"]
    style.font.name = FONT_EN
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY_CN)

    setup_page(doc)

    issue = content["issue"]
    sections = content.get("sections", {})

    add_header_footer(doc, issue)
    add_cover_page(doc, issue)
    add_toc(doc, sections)

    # 渲染各栏目
    if sections.get("preface"):
        render_preface(doc, sections["preface"])
        doc.add_page_break()

    if sections.get("guideline_updates"):
        render_guideline_updates(doc, sections["guideline_updates"])
        doc.add_page_break()

    if sections.get("safety_alerts"):
        render_safety_alerts(doc, sections["safety_alerts"])
        doc.add_page_break()

    if sections.get("literature"):
        render_literature(doc, sections["literature"])
        doc.add_page_break()

    if sections.get("new_drugs"):
        render_new_drugs(doc, sections["new_drugs"])

    return doc


def main():
    parser = argparse.ArgumentParser(description="月度临床药讯Word文档生成器")
    parser.add_argument("--input", "-i", required=True, help="输入JSON内容文件路径")
    parser.add_argument("--output", "-o", required=True, help="输出.docx文件路径")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"错误：输入文件不存在 {input_path}", file=sys.stderr)
        sys.exit(1)

    try:
        with open(input_path, "r", encoding="utf-8") as f:
            content = json.load(f)
    except json.JSONDecodeError as e:
        print(f"错误：JSON解析失败 {e}", file=sys.stderr)
        sys.exit(1)

    if "issue" not in content:
        print("错误：缺少 issue 字段（year/month）", file=sys.stderr)
        sys.exit(1)

    doc = build_document(content)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"✅ 药讯已生成：{output_path}")


if __name__ == "__main__":
    main()
