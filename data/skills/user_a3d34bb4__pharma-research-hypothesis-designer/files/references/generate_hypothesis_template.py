# -*- coding: utf-8 -*-
"""
研究假设方案 Word 文档生成模板
================================
基于 python-docx 生成格式化的药学课题研究假设方案文档。

使用方法：
1. 复制本脚本到项目目录
2. 替换 # ====== 课题信息 ====== 部分的所有内容
3. 运行：python generate_hypothesis.py
4. 文档自动保存到脚本同目录

依赖：pip install python-docx
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os, datetime

# 中文智能引号（用Unicode转义避免编码问题）
LQ = "\u201C"
RQ = "\u201D"

# ====== 课题信息（按实际替换）======
COURSE_TITLE = "雷贝拉唑与骨折风险——研究假设方案"
COURSE_SUBTITLE = "课题方向：长期使用雷贝拉唑的老年患者新发骨折风险前瞻性队列研究"
OUTPUT_FILENAME = "雷贝拉唑与骨折风险_研究假设方案.docx"

# 研究空白背景（来自前期缺口分析）
GAP_PEOPLE = [
    "≥80岁高龄老人完全缺失（P5, C4提示老年人群但上限仅75岁）；",
    "中国社区老年队列空白（C4, C11为中国数据但样本小、设计旧）；",
    "长期用药（>2年）的老年亚组未单独分析（P8, P9提到长期暴露但未按年龄分层）。",
]
GAP_METHOD = [
    "雷贝拉唑特异性前瞻性队列缺失（P2, P7的Meta分析均为PPI整体，无单药证据）；",
    "cDDD量化暴露+骨折终点的设计未出现（C1, C2提到cDDD概念但未用于前瞻性研究）；",
    "中文文献无前瞻性设计（C1–C12全部为综述、Meta分析或回顾性）；",
    "竞争风险模型未应用（P9提到死亡竞争风险但未做正式竞争风险建模）。",
]
GAP_MECHANISM = [
    "SERCA2a-OPG通路仅为理论假设（P4, P6提到但未验证）；",
    "雷贝拉唑 vs 奥美拉唑的骨代谢差异无头对头比较（P1, P5均未区分）；",
    "炎症标志物（CRP）与PPI骨丢失的关联未被检验（P4提到炎症但无实证）；",
    "钙吸收抑制的剂量-效应关系未建立（C3, C6提到机制但无量化数据）。",
]

# PICO 四要素
PICO_P = {
    "title": "P（人群）",
    "content": [
        [("≥65岁、首次处方雷贝拉唑、连续用药≥90天的门诊老年患者。", False)],
        [("排除标准：", True),
         ("①既往骨折史；②已确诊骨质疏松症（T值≤-2.5）；③正在使用双膦酸盐/降钙素/特立帕肽；④甲旁腺功能异常；⑤慢性肾病eGFR<30；⑥胃大部切除术后。", False)],
    ],
}
PICO_I = {
    "title": "I（暴露）",
    "content": [
        [("长期使用雷贝拉唑，以cDDD（累计限定日剂量）量化暴露量：", False)],
        [("高暴露组 ", False), ("cDDD≥180", True, (0xB0, 0x30, 0x30)),
         ("（相当于20 mg/d持续≥180天）；中暴露组 cDDD 90–179；低暴露组 cDDD 30–89。", False)],
    ],
}
PICO_C = {
    "title": "C（对照）",
    "content": "同期门诊≥65岁、未使用PPI或cDDD<30的老年患者，按年龄±3岁/性别/BMI±2 kg/m²/MoCA评分±2分行1:1个体匹配。",
}
PICO_O = {
    "title": "O（结局）",
    "content": [
        [("主要终点：", True),
         ("12个月内影像学（X线/CT）确诊的新发骨折发生率（‰）。", False)],
        [("次要终点：", True),
         ("①腰椎L1-L4及股骨颈骨密度年变化率（g/cm²/年）；②血清校正钙水平12个月变化值（mmol/L）。", False)],
    ],
}

# 核心假设
HYPOTHESIS = (
    "雷贝拉唑cDDD≥180通过降低肠道钙吸收效率"
    "导致≥65岁老年患者12个月新发骨折发生率"
    "较未用药对照组升高30%以上"
)
HYPOTHESIS_DECOMPOSE = (
    "句式拆解：X（雷贝拉唑cDDD≥180）通过 Y（降低肠道钙吸收效率）"
    "影响 Z（≥65岁老年患者12个月新发骨折发生率升高30%以上）。"
)

# 研究设计
DESIGN_PRIMARY = "首选设计：前瞻性队列研究"
DESIGN_PRIMARY_REASONS = [
    "1. 暴露识别可执行：HIS处方数据库可直接提取雷贝拉唑处方记录，精确计算每例患者cDDD值，无需额外建库。",
    "2. 结局追踪可执行：门诊随访记录系统可追踪患者就诊时间、影像学检查申请及骨折诊断编码，骨折终点以影像学报告为准。",
    "3. 混杂控制可执行：已有部分患者MoCA评分数据，可识别认知障碍这一跌倒/骨折的关键混杂因素，实现分层匹配而非简单排除。",
    "4. 时序优势：前瞻性设计可在入组时统一采集基线骨密度和血清钙，避免回顾性研究的数据缺失偏倚。",
]
DESIGN_BACKUP_TITLE = "备选方案：巢式病例对照研究"
DESIGN_BACKUP = (
    "若12个月随访依从性低于70%或新发骨折事件少于30例，则在已建立队列基础上以骨折病例为index case，"
    "按1:4匹配非骨折对照（匹配变量：年龄±3岁/性别/BMI±2/MoCA±2），"
    "以条件Logistic回归估计cDDD与骨折的OR值及95%CI。"
)

# 五段式研究问题陈述
RQ_1 = ("1. 临床现象　", True, (0x2E, 0x5C, 0x8A)), (
    "我院门诊≥65岁老年患者中长期服用雷贝拉唑（连续处方≥90天）者在PPI长期用药人群中占主要部分；"
    "骨科住院记录显示老年低暴力骨折患者合并PPI用药史的比例高于非骨折同龄住院者。"
)
RQ_2 = ("2. 已有研究　", True, (0x2E, 0x5C, 0x8A)), (
    "P2（Robert 2025）和P7（Veettil 2022）的Meta分析共纳入超10万例患者，"
    "报告PPI使用者骨折风险pooled OR=1.30（95%CI: 1.18–1.43）；"
    "P8（Park 2022）韩国全国队列显示PPI使用≥1年者髋部骨折HR=1.42；"
    "C4（李敏 2019）和C11（袁志敏 2013）的中国数据亦支持此关联。"
)
RQ_3 = ("3. 研究空白（关键）　", True, (0xB0, 0x30, 0x30)), (
    "上述所有证据均将PPI作为整体暴露，未分离雷贝拉唑的独立效应（方法空白第1条）；"
    "无一研究采用cDDD进行剂量-效应建模（方法空白第2条）；"
    "中文文献12篇中无一篇采用前瞻性设计（方法空白第3条）；"
    "≥80岁高龄亚组在全部文献中完全缺失（人群空白第1条）。"
)
RQ_4 = ("4. 本研究做什么　", True, (0x2E, 0x5C, 0x8A)), (
    "建立≥65岁老年患者前瞻性队列，以cDDD三档量化雷贝拉唑暴露量，"
    "随访12个月以影像学确诊为新发骨折终点，"
    "采用Cox比例风险模型估计cDDD每增加90单位对应的骨折风险增量（HR及95%CI），"
    "并将血清校正钙变化值纳入Baron-Kenny中介效应检验。"
)
RQ_5 = [
    ("5. 预期回答什么　", True, (0x2E, 0x5C, 0x8A)),
    ("量化回答", False),
    (LQ + "cDDD≥180是否使12个月新发骨折发生率升高30%以上" + RQ, True),
    ("，并给出cDDD从90增至180每增加90单位对应的骨折风险增量百分比。", False),
]

# 自查清单
CHECKLIST = [
    {
        "dim": "方向性",
        "judge": ("强", True, (0x2E, 0x7D, 0x32)),
        "reason": "假设明确预测" + LQ + "升高" + RQ + "方向，非双向探索。",
    },
    {
        "dim": "可检验性",
        "judge": ("强", True, (0x2E, 0x7D, 0x32)),
        "reason": "cDDD可从HIS处方记录直接计算；骨折终点以影像学确诊为金标准；12个月随访期在我科门诊条件下可操作。",
    },
    {
        "dim": "机制性",
        "judge": ("中→强（已改）", True, (0xD4, 0xA0, 0x17)),
        "reason": [
            [("原假设提及" + LQ + "降低肠道钙吸收效率" + RQ + "但未直接测量吸收率。", False)],
            [("改法：", True, (0xB0, 0x30, 0x30)),
             ("将血清校正钙12个月变化值设为中介变量，纳入Baron-Kenny中介效应检验，使机制路径可量化。", False)],
        ],
    },
    {
        "dim": "具体性",
        "judge": ("强", True, (0x2E, 0x7D, 0x32)),
        "reason": "锁定了暴露阈值（cDDD≥180）、人群（≥65岁首处方）、随访时长（12个月）、效应量阈值（升高30%以上）、终点定义（影像学确诊新发骨折）。",
    },
]


# ===== Helper functions =====

def set_cell_shading(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def set_cell_borders(cell, color="999999", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for b in ['top', 'left', 'bottom', 'right']:
        elem = OxmlElement('w:' + b)
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), sz)
        elem.set(qn('w:color'), color)
        borders.append(elem)
    tcPr.append(borders)

def set_font(run, size=11, bold=False, color=None, cn="\u5b8b\u4f53", en="Times New Roman"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=16, bold=True, color=(0x1F, 0x3A, 0x5F))
    return p

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text)
    set_font(r, size=10, bold=False, color=(0x66, 0x66, 0x66))
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=14, bold=True, color=(0x1F, 0x3A, 0x5F))
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=12, bold=True, color=(0x2E, 0x5C, 0x8A))
    return p

def add_body(doc, content, size=11, indent=False, space_after=4, line_spacing=1.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if isinstance(content, str):
        content = [(content, False, None)]
    for seg in content:
        text = seg[0]
        bold = seg[1] if len(seg) > 1 else False
        color = seg[2] if len(seg) > 2 else None
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    return p

def fill_cell(cell, content, size=10, bold=False, align=None):
    if isinstance(content, str):
        lines = [[(content, bold)]]
    elif isinstance(content, tuple):
        lines = [content]
    elif isinstance(content, list):
        if content and isinstance(content[0], tuple):
            lines = [content]
        else:
            lines = content
    else:
        lines = [[(str(content), bold)]]
    for i, segments in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        if align:
            p.alignment = align
        for seg in segments:
            if isinstance(seg, str):
                seg = (seg, bold, None)
            text = seg[0]
            b = seg[1] if len(seg) > 1 else bold
            color = seg[2] if len(seg) > 2 else None
            r = p.add_run(text)
            set_font(r, size=size, bold=b, color=color)

def make_row(table, cells, header=False, header_dark=False):
    row = table.add_row()
    for i, cell_data in enumerate(cells):
        cell = row.cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        fill_cell(cell, cell_data, size=10, bold=header or header_dark)
        set_cell_borders(cell)
        if header_dark:
            set_cell_shading(cell, "1F3A5F")
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif header:
            set_cell_shading(cell, "D6E4F0")
        else:
            idx = len(table.rows)
            if idx % 2 == 0:
                set_cell_shading(cell, "F5F9FC")
            else:
                set_cell_shading(cell, "FFFFFF")
    return row


# ===== Build document =====

def build_document():
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    today = datetime.date.today().strftime("%Y\u5e74%m\u6708%d\u65e5")

    # ---- Title ----
    add_title(doc, COURSE_TITLE)
    add_subtitle(doc, COURSE_SUBTITLE + "\u3000|\u3000" + today)

    # ---- Background: research gaps ----
    add_h1(doc, "\u7814\u7a76\u7a7a\u767d\u80cc\u666f\uff08\u6765\u81ea\u524d\u671f\u7f3a\u53e3\u5206\u6790\uff09")

    add_body(doc, [
        ("\u3010\u4eba\u7fa4\u7a7a\u767d\u3011", True, (0xB0, 0x30, 0x30)),
    ] + [("\u2460 " + g, False) for g in GAP_PEOPLE], size=10, indent=True, space_after=3)

    add_body(doc, [
        ("\u3010\u65b9\u6cd5\u7a7a\u767d\u3011", True, (0xB0, 0x30, 0x30)),
    ] + [("\u2460 " + g, False) for g in GAP_METHOD], size=10, indent=True, space_after=3)

    add_body(doc, [
        ("\u3010\u673a\u5236\u7a7a\u767d\u3011", True, (0xB0, 0x30, 0x30)),
    ] + [("\u2460 " + g, False) for g in GAP_MECHANISM], size=10, indent=True, space_after=8)

    # ---- 一、PICO ----
    add_h1(doc, "\u4e00\u3001PICO\u56db\u8981\u7d20")

    pico_table = doc.add_table(rows=0, cols=2)
    pico_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pico_table.style = 'Table Grid'

    make_row(pico_table, ["\u8981\u7d20", "\u5177\u4f53\u5b9a\u4e49"], header_dark=True)

    for pico in [PICO_P, PICO_I, PICO_C, PICO_O]:
        make_row(pico_table, [[(pico["title"], True)], pico["content"]])

    for row in pico_table.rows:
        row.cells[0].width = Cm(3.2)
        row.cells[1].width = Cm(13.3)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ---- 二、核心假设 ----
    add_h1(doc, "\u4e8c\u3001\u6838\u5fc3\u5047\u8bbe")

    hyp_table = doc.add_table(rows=1, cols=1)
    hyp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hyp_cell = hyp_table.rows[0].cells[0]
    set_cell_shading(hyp_cell, "FFF8E1")
    tcPr = hyp_cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for b in ['top', 'left', 'bottom', 'right']:
        elem = OxmlElement('w:' + b)
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), '8')
        elem.set(qn('w:color'), 'D4A017')
        borders.append(elem)
    tcPr.append(borders)

    p = hyp_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.6
    r = p.add_run(HYPOTHESIS)
    set_font(r, size=13, bold=True, color=(0x1F, 0x3A, 0x5F))

    add_body(doc, [
        ("\u53e5\u5f0f\u62c6\u89e3\uff1a", False, (0x66, 0x66, 0x66)),
        (HYPOTHESIS_DECOMPOSE, False, (0x66, 0x66, 0x66)),
    ], size=9, space_after=10)

    # ---- 三、研究设计选型 ----
    add_h1(doc, "\u4e09\u3001\u7814\u7a76\u8bbe\u8ba1\u9009\u578b")

    add_h2(doc, DESIGN_PRIMARY)
    add_body(doc, [("\u7406\u7531\uff08\u7d27\u6263\u79d1\u5ba4\u6570\u636e\u6761\u4ef6\uff09\uff1a", True)], size=11, space_after=2)

    for r_text in DESIGN_PRIMARY_REASONS:
        add_body(doc, r_text, size=11, indent=True, space_after=2)

    add_h2(doc, DESIGN_BACKUP_TITLE)
    add_body(doc, DESIGN_BACKUP, size=11, indent=True)

    # ---- 四、研究问题陈述 ----
    add_h1(doc, "\u56db\u3001\u7814\u7a76\u95ee\u9898\u9648\u8ff0\uff08\u4e94\u6bb5\u5f0f\uff09")

    for rq in [RQ_1, RQ_2, RQ_3, RQ_4]:
        add_body(doc, [rq[0], (rq[1], False)], size=11, indent=True)

    add_body(doc, RQ_5, size=11, indent=True)

    # ---- 五、自查清单 ----
    add_h1(doc, "\u4e94\u3001\u81ea\u67e5\u6e05\u5355")

    check_table = doc.add_table(rows=0, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    check_table.style = 'Table Grid'

    make_row(check_table, ["\u7ef4\u5ea6", "\u5224\u65ad", "\u7406\u7531\u6216\u6539\u6cd5"], header_dark=True)

    for item in CHECKLIST:
        make_row(check_table, [
            [(item["dim"], True)],
            [item["judge"]],
            item["reason"],
        ])

    for row in check_table.rows:
        row.cells[0].width = Cm(2.8)
        row.cells[1].width = Cm(3.0)
        row.cells[2].width = Cm(10.7)

    # ---- Save ----
    output = os.path.join(os.path.dirname(os.path.abspath(__file__)), OUTPUT_FILENAME)
    doc.save(output)
    print(f"\u6587\u6863\u5df2\u4fdd\u5b58: {output}")
    return output


if __name__ == "__main__":
    build_document()
